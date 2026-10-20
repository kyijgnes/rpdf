"""창 하나: 문서 탭 · 리본 · 왼쪽 막대/패널 · 보기(뷰어 또는 구성 격자) · 오른쪽 속성 · 상태줄 · 시작 화면."""
from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.parse
import webbrowser

import fitz
from PySide6.QtCore import QPointF, QRectF, QSize, Qt, QTimer, Signal
from PySide6.QtGui import QAction, QCloseEvent, QColor, QFont, QIcon, QImage, QKeySequence, QPainter, QPixmap, QShortcut
from PySide6.QtPrintSupport import QPrinterInfo
from PySide6.QtWidgets import (QApplication, QButtonGroup, QComboBox, QFileDialog, QFrame, QHBoxLayout, QInputDialog,
                               QLabel, QLineEdit, QMainWindow, QMenu, QMessageBox, QSlider, QSplitter, QStackedWidget,
                               QTabBar, QToolButton, QVBoxLayout, QWidget)

from ..engine import annots as A
from ..engine import pages as P
from ..engine.document import Document, DocumentError, NeedsPassword
from . import dialogs as D
from . import icons, theme
from .organize import OrganizeView
from .props import PropsPanel
from .ribbon import LABELS, TOOL_KEYS, Ribbon
from .sidebar import AnnotPanel, AttachmentPanel, BookmarkPanel, FieldPanel, LayerPanel, SearchPanel, ThumbnailPanel
from .start import StartPage
from .update import UpdateUI
from .viewer import PageView

from ..appinfo import APP_ID, APP_NAME, ICON, OLD_APP_IDS  # noqa: E402

SETTINGS = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), APP_ID, "settings.json")
PDF_FILTER = "PDF 파일 (*.pdf);;모든 파일 (*)"
OPEN_FILTER = "PDF·이미지·문서 (*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff *.gif *.webp *.xps *.epub *.txt *.svg *.cbz);;PDF 파일 (*.pdf);;모든 파일 (*)"


def _engine(name: str):
    """변환·도구·OCR 엔진은 따로 짓고 있어서 없을 수도 있다."""
    try:
        import importlib
        return importlib.import_module(f"app.engine.{name}")
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"{name} 엔진을 아직 쓸 수 없습니다 ({e})") from e


def _copy_old_settings():
    """이름을 바꾸면 설정 폴더(%APPDATA% 아래 APP_ID 폴더)도 바뀐다. 새 자리가 비었으면 옛 이름 자리의 설정을 복사해 온다(옛 exe 도 그대로 쓰게 옮기지 않고 복사)."""
    if os.path.exists(SETTINGS):
        return
    root = os.path.dirname(os.path.dirname(SETTINGS))
    for old in OLD_APP_IDS:
        src = os.path.join(root, old, "settings.json")
        if old != APP_ID and os.path.isfile(src):
            try:
                os.makedirs(os.path.dirname(SETTINGS), exist_ok=True)
                shutil.copy2(src, SETTINGS)
            except OSError:
                pass
            return


def load_settings() -> dict:
    _copy_old_settings()
    try:
        with open(SETTINGS, encoding="utf-8") as f:
            return json.load(f)
    except Exception:  # noqa: BLE001
        return {}


def save_settings(s: dict):
    os.makedirs(os.path.dirname(SETTINGS), exist_ok=True)
    with open(SETTINGS, "w", encoding="utf-8") as f:
        json.dump(s, f, ensure_ascii=False, indent=1)


class DocTab(QWidget):
    """문서 하나의 화면 묶음."""
    def __init__(self, doc: Document, win: "MainWindow"):
        super().__init__()
        self.doc = doc
        self.win = win
        self.viewer = PageView(doc)
        self.organize = OrganizeView(doc)
        self.view_stack = QStackedWidget()
        self.view_stack.addWidget(self.viewer); self.view_stack.addWidget(self.organize)
        self.thumbs = ThumbnailPanel(doc)
        self.bookmarks = BookmarkPanel(doc, lambda: self.viewer.current)
        self.annots = AnnotPanel(doc)
        self.attachments = AttachmentPanel(doc)
        self.search = SearchPanel(doc)
        self.layers = LayerPanel(doc)
        self.fieldlist = FieldPanel(doc)
        self.side = QStackedWidget(); self.side.setObjectName("sidepanel"); self.side.setFixedWidth(240)
        self.panels = {"thumbs": self.thumbs, "bookmark": self.bookmarks, "comments": self.annots, "attach": self.attachments,
                       "search": self.search, "layers": self.layers, "fields": self.fieldlist}
        for p in self.panels.values():
            self.side.addWidget(p)
        self.side.hide()
        self.props = PropsPanel(self.viewer.props)
        self.props.hide()
        self.warn = QLabel(""); self.warn.setObjectName("warn"); self.warn.hide()
        root = QVBoxLayout(self); root.setContentsMargins(0, 0, 0, 0); root.setSpacing(0)
        root.addWidget(self.warn)
        row = QHBoxLayout(); row.setContentsMargins(0, 0, 0, 0); row.setSpacing(0)
        row.addWidget(self.side); row.addWidget(self.view_stack, 1); row.addWidget(self.props)
        root.addLayout(row)
        self.side_key = ""
        # 신호
        self.thumbs.pageClicked.connect(self.viewer.goto)
        self.thumbs.reordered.connect(self.reorder)
        self.thumbs.pagesAction.connect(self.win.pages_action)
        self.organize.reordered.connect(self.reorder)
        self.organize.pagesAction.connect(self.win.pages_action)
        self.organize.openPage.connect(lambda i: (self.win.ribbon.set_tab("home"), self.viewer.goto(i)))
        self.bookmarks.gotoPage.connect(self.viewer.goto)
        self.bookmarks.changed.connect(lambda: self.win.mark_modified())
        self.annots.gotoAnnot.connect(lambda pg, r: self.viewer.goto(pg, r.y0))
        self.annots.changed.connect(lambda: self.changed(None))
        self.annots.importExport.connect(self.win.annots_import_export)
        self.attachments.changed.connect(lambda: self.win.mark_modified())
        self.search.hitsChanged.connect(self.show_hits)
        self.search.replaced.connect(lambda: self.changed(None))
        self.layers.changed.connect(self.view_changed)
        self.fieldlist.gotoField.connect(lambda pg, r: self.viewer.goto(pg, r.y0))
        self.viewer.pageChanged.connect(self.thumbs.set_current)
        self.viewer.selectionChanged.connect(self.props.show_selection)
        self.props.applyToSel.connect(self.apply_props)
        self.props.deleteSel.connect(self.viewer.delete_selection)
        self.props.changed.connect(lambda d: None)
        if doc.fitz.is_encrypted or doc.password:
            self.warn.setText("암호 걸린 문서 · 저장해도 같은 열기 암호 유지"); self.warn.show()

    def show_side(self, key: str):
        if self.side_key == key and self.side.isVisible():
            self.side.hide(); self.side_key = ""
            return
        self.side_key = key
        self.side.setCurrentWidget(self.panels[key]); self.side.show()
        if key == "search":
            self.search.focus()

    def changed(self, pages: list[int] | None, structural: bool = False, mapping: list[int | None] | None = None):
        """문서가 바뀌었다: 뷰어·썸네일·목록을 새로 그린다.
        structural 일 때 mapping[새 쪽] = 옛 쪽 을 주면 안 바뀐 쪽 그림은 옮겨 쓰고 pages(새 쪽 번호)만 다시 그린다.
        mapping 없이 pages 만 주면 쪽 수는 그대로이고 그 쪽만 바뀐 것(회전)."""
        if structural and mapping is None and pages is not None:
            mapping = list(range(self.doc.page_count))
        dirty = set(pages or ())
        self.viewer.refresh_pages(pages, structural=structural, mapping=mapping)
        if structural:
            self.thumbs.reload(mapping, dirty); self.organize.reload(mapping, dirty)
        else:
            self.thumbs.refresh_pages(pages)
            self.organize.thumbs.forget(pages)   # 숨어 있으면 안 그리고, 다음에 보일 때 그린다
        self.annots.reload(); self.fieldlist.reload(); self.bookmarks.reload(); self.attachments.reload(); self.layers.reload()
        self.win.mark_modified()

    def view_changed(self):
        """보기 설정(레이어 켜고 끄기)만 바뀌었다: 다시 그리되 고친 문서로 치지 않는다."""
        self.viewer.refresh_pages(None)
        self.thumbs.refresh_pages(None)
        self.organize.thumbs.forget(None)

    def reorder(self, order: list[int]):
        P.reorder(self.doc, order)
        self.viewer.clear_selection()
        self.changed(None, structural=True, mapping=list(order))

    def show_hits(self, hits, current):
        self.viewer.search_hits = hits
        self.viewer.search_current = current
        if 0 <= current < len(hits):
            pg, rs = hits[current]
            r = rs[0] if isinstance(rs, list) else rs
            self.viewer.goto(pg, max(0, r.y0 - 100))
        self.viewer.viewport().update()

    def apply_props(self, d: dict):
        s = self.viewer.sel
        if s and s.kind == "annot":
            A.update(self.doc, s.page, s.xref, stroke=d["stroke"], fill=d["fill"], width=d["width"], opacity=d["opacity"],
                     content=d["content"])
            self.changed([s.page])

    def close(self):
        self.viewer.close(); self.thumbs.close(); self.organize.close(); self.search.close()
        self.doc.close()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = load_settings()
        self.settings.setdefault("recent", []); self.settings.setdefault("favorites", []); self.settings.setdefault("prefs", {})
        self.prefs = self.settings["prefs"]
        self.setWindowTitle(APP_NAME)
        self.resize(1360, 820)
        self.setWindowIcon(QIcon(ICON) if os.path.isfile(ICON) else icons.icon("thumbs", 32, theme.ACCENT))
        self.tabs: list[DocTab] = []
        self._receiver = None   # 다른 데서 켠 RPDF 가 파일을 넘겨 오는 곳(single.Receiver)
        self._build()
        self.apply_theme()
        self._shortcuts()
        self.show_start()

    def show_initial(self):
        """켤 때 창 크기: 환경 설정 「열 때 창 크기」(마지막으로 닫은 크기 · 화면 가득 · 보통 크기)."""
        how = self.prefs.get("window_size", "last")
        geo = self.settings.get("window_geometry")
        if how == "last" and geo:
            from PySide6.QtCore import QByteArray
            ok = self.restoreGeometry(QByteArray.fromBase64(geo.encode("ascii")))
            if ok and self._on_screen():
                if self.settings.get("window_max"):   # 최대화로 닫았으면 최대화로(풀면 닫기 전 크기)
                    self.showMaximized(); return
                self.show(); return
            how = "max" if self.settings.get("window_max") else "normal"   # 그 모니터가 없어졌다
        if how == "max":
            self.showMaximized(); return
        scr = self.screen().availableGeometry()
        self.resize(min(1360, scr.width() - 40), min(820, scr.height() - 40))
        self.move(scr.center() - self.rect().center())
        self.show()

    def _on_screen(self) -> bool:
        """창 제목줄이 어느 화면 안에 들어 있나(모니터를 빼서 화면 밖으로 나가 있지 않나)."""
        from PySide6.QtCore import QRect
        g = self.normalGeometry() if self.normalGeometry().isValid() else self.geometry()
        bar = QRect(g.left() + 40, g.top() - 25, max(g.width() - 80, 40), 25)   # 제목줄 가운데
        return any(scr.availableGeometry().intersects(bar) for scr in QApplication.screens())

    def reopen_last(self):
        """환경 설정 「시작할 때 지난번 문서 다시 열기」."""
        if self.prefs.get("restore"):
            for p in self.settings.get("last_open", []):
                if os.path.isfile(p):
                    self.open_path(p)

    def open_forwarded(self, paths: list[str]):
        """다른 데서 켠 RPDF 가 넘긴 파일을 이 창의 탭으로 연다. 파일이 없으면 창만 앞으로."""
        if self.isMinimized():   # 최대화였으면 최대화로 되돌아온다
            self.setWindowState((self.windowState() & ~Qt.WindowState.WindowMinimized) | Qt.WindowState.WindowActive)
        self.raise_(); self.activateWindow()
        for p in paths:
            if os.path.isfile(p):
                self.open_path(p)

    def _remember_window(self):
        self.settings["window_max"] = self.isMaximized() or self.isFullScreen()
        if not self.isFullScreen():
            self.settings["window_geometry"] = bytes(self.saveGeometry().toBase64()).decode("ascii")

    # ── 화면 짓기 ─────────────────────────────────────────────
    def _build(self):
        central = QWidget(); central.setObjectName("central"); self.setCentralWidget(central)
        root = QVBoxLayout(central); root.setContentsMargins(0, 0, 0, 0); root.setSpacing(0)
        # 문서 탭 줄
        tabrow = QWidget(); tabrow.setObjectName("titlebar"); tl = QHBoxLayout(tabrow); tl.setContentsMargins(6, 0, 6, 0); tl.setSpacing(4)
        home = QToolButton(); home.setObjectName("quick"); home.setIcon(icons.icon("thumbs", 18, theme.ACCENT)); home.setToolTip("시작 화면")
        home.clicked.connect(self.show_start); tl.addWidget(home)
        self.doctabs = QTabBar(); self.doctabs.setObjectName("doctabs"); self.doctabs.setTabsClosable(True); self.doctabs.setMovable(True)
        self.doctabs.setExpanding(False); self.doctabs.setDrawBase(False); self.doctabs.setElideMode(Qt.TextElideMode.ElideMiddle)
        self.doctabs.currentChanged.connect(self._tab_changed); self.doctabs.tabCloseRequested.connect(self.close_tab)
        self.doctabs.tabMoved.connect(self._tab_moved)
        tl.addWidget(self.doctabs)
        addb = QToolButton(); addb.setObjectName("quick"); addb.setIcon(icons.icon("add", 16)); addb.setToolTip("열기"); addb.clicked.connect(self.open_dialog)
        tl.addWidget(addb); tl.addStretch()
        self.updates = UpdateUI(self); tl.addWidget(self.updates.bar)  # [새 버전 x 설치]·「설치 끝」. 시작 화면에서도 보이는 줄
        root.addWidget(tabrow)
        # 리본
        self.ribbon = Ribbon()
        self.ribbon.action.connect(self.on_action); self.ribbon.tool.connect(self.on_tool); self.ribbon.toggled.connect(self.on_toggle)
        self.ribbon.tabChanged.connect(self.on_ribbon_tab); self.ribbon.searchRequested.connect(self.quick_search)
        self.ribbon.fileMenu.connect(self.show_file_menu)
        root.addWidget(self.ribbon)
        # 본문
        self.body = QStackedWidget()
        self.start = StartPage(self.settings["recent"], self.settings["favorites"])
        self.start.openFile.connect(self.open_dialog); self.start.openPath.connect(self.open_path)
        self.start.tool.connect(self.start_tool); self.start.newBlank.connect(lambda: self.on_action("create_blank"))
        self.start.recentChanged.connect(lambda: save_settings(self.settings))
        self.body.addWidget(self.start)
        docarea = QWidget(); dl = QHBoxLayout(docarea); dl.setContentsMargins(0, 0, 0, 0); dl.setSpacing(0)
        self.navstrip = QWidget(); self.navstrip.setObjectName("navstrip"); self.navstrip.setFixedWidth(34)
        nl = QVBoxLayout(self.navstrip); nl.setContentsMargins(0, 6, 0, 6); nl.setSpacing(2)
        self.navbtns: dict[str, QToolButton] = {}
        for key, name, tip in (("thumbs", "thumbs", "썸네일"), ("bookmark", "bookmark", "북마크"), ("comments", "comments", "주석"),
                               ("attach", "attach", "첨부 파일"), ("search", "search", "검색"), ("layers", "layers", "레이어")):
            b = QToolButton(); b.setObjectName("nav"); b.setIcon(icons.icon(name, 16)); b.setToolTip(tip); b.setCheckable(True)
            b.clicked.connect(lambda _, k=key: self.toggle_side(k)); nl.addWidget(b); self.navbtns[key] = b
        nl.addStretch()
        hb = QToolButton(); hb.setObjectName("nav"); hb.setIcon(icons.icon("help", 16)); hb.setToolTip("도움말"); hb.clicked.connect(self.show_help); nl.addWidget(hb)
        dl.addWidget(self.navstrip)
        self.docstack = QStackedWidget(); dl.addWidget(self.docstack, 1)
        self.rightstrip = QWidget(); self.rightstrip.setObjectName("rightstrip"); self.rightstrip.setFixedWidth(34)
        rl = QVBoxLayout(self.rightstrip); rl.setContentsMargins(0, 6, 0, 6); rl.setSpacing(2)
        self.rightbtns = {}
        for key, name, tip in (("props", "props", "속성"), ("fields", "fields", "양식 칸"), ("sign", "signature", "서명"), ("translate", "translate", "글자 번역·검색")):
            b = QToolButton(); b.setObjectName("nav"); b.setIcon(icons.icon(name, 16)); b.setToolTip(tip); b.setCheckable(key in ("props", "fields"))
            b.setStyleSheet("QToolButton#nav{border-left:none; border-right:2px solid transparent} QToolButton#nav:checked{border-right:2px solid %s}" % theme.ACCENT)
            b.clicked.connect(lambda _, k=key: self.right_action(k)); rl.addWidget(b); self.rightbtns[key] = b
        rl.addStretch()
        dl.addWidget(self.rightstrip)
        self.body.addWidget(docarea)
        root.addWidget(self.body, 1)
        # 상태줄
        self.status = QWidget(); self.status.setObjectName("statusbar"); self.status.setFixedHeight(28)
        sl = QHBoxLayout(self.status); sl.setContentsMargins(8, 0, 8, 0); sl.setSpacing(4)
        self.size_label = QLabel(""); sl.addWidget(self.size_label)
        self.msg_label = QLabel(""); self.msg_label.setObjectName("dim"); sl.addWidget(self.msg_label)
        self.msg_label.setMinimumWidth(1)  # 긴 알림이 창 최소 폭을 밀어 올리지 않게(넘치면 잘리고 툴팁에 원문)
        sl.addStretch()
        self.sb_hand = QToolButton(); self.sb_hand.setObjectName("quick"); self.sb_hand.setIcon(icons.icon("hand", 16)); self.sb_hand.setCheckable(True); self.sb_hand.setToolTip("손")
        self.sb_sel = QToolButton(); self.sb_sel.setObjectName("quick"); self.sb_sel.setIcon(icons.icon("select", 16)); self.sb_sel.setCheckable(True); self.sb_sel.setToolTip("선택")
        self.sb_hand.clicked.connect(lambda: self.on_tool("hand")); self.sb_sel.clicked.connect(lambda: self.on_tool("select"))
        sl.addWidget(self.sb_hand); sl.addWidget(self.sb_sel); sl.addSpacing(8)
        for key, name in (("first", "first"), ("prev", "prev")):
            b = QToolButton(); b.setObjectName("quick"); b.setIcon(icons.icon(name, 14)); b.clicked.connect(lambda _, k=key: self.on_action(k)); sl.addWidget(b)
        self.page_edit = QLineEdit(); self.page_edit.setFixedWidth(44); self.page_edit.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.page_edit.returnPressed.connect(self._goto_entered); sl.addWidget(self.page_edit)
        self.page_total = QLabel("/ 0"); sl.addWidget(self.page_total)
        for key, name in (("next", "next"), ("last", "last")):
            b = QToolButton(); b.setObjectName("quick"); b.setIcon(icons.icon(name, 14)); b.clicked.connect(lambda _, k=key: self.on_action(k)); sl.addWidget(b)
        sl.addSpacing(8)
        for key, name, tip in (("layout_single", "single", "한 페이지씩"), ("layout_continuous", "continuous", "연속"),
                               ("layout_facing", "facing", "두 페이지"), ("layout_facing_continuous", "tile", "두 페이지 연속")):
            b = QToolButton(); b.setObjectName("quick"); b.setIcon(icons.icon(name, 14)); b.setToolTip(tip)
            b.clicked.connect(lambda _, k=key: self.on_action(k)); sl.addWidget(b)
        sl.addStretch()
        zo = QToolButton(); zo.setObjectName("quick"); zo.setIcon(icons.icon("zoomout", 14)); zo.clicked.connect(lambda: self.on_action("zoomout")); sl.addWidget(zo)
        self.zoom_slider = QSlider(Qt.Orientation.Horizontal); self.zoom_slider.setRange(10, 800); self.zoom_slider.setFixedWidth(110)
        self.zoom_slider.sliderMoved.connect(self._zoom_slider); sl.addWidget(self.zoom_slider)
        zi = QToolButton(); zi.setObjectName("quick"); zi.setIcon(icons.icon("zoomin", 14)); zi.clicked.connect(lambda: self.on_action("zoomin")); sl.addWidget(zi)
        self.zoom_box = QComboBox(); self.zoom_box.setEditable(True); self.zoom_box.setFixedWidth(92)
        self.zoom_box.addItems(["50%", "75%", "100%", "125%", "150%", "200%", "300%", "400%", "창 너비", "창 크기", "실제 크기"])
        self.zoom_box.activated.connect(self._zoom_box); self.zoom_box.lineEdit().returnPressed.connect(self._zoom_box)
        sl.addWidget(self.zoom_box)
        fs = QToolButton(); fs.setObjectName("quick"); fs.setIcon(icons.icon("fullscreen", 14)); fs.setToolTip("전체 화면"); fs.clicked.connect(lambda: self.on_action("fullscreen")); sl.addWidget(fs)
        root.addWidget(self.status)
        self.setAcceptDrops(True)

    def _shortcuts(self):
        for seq, key in (("Ctrl+O", "open"), ("Ctrl+S", "save"), ("Ctrl+Shift+S", "saveas"), ("Ctrl+P", "print"), ("Ctrl+Z", "undo"),
                         ("Ctrl+Y", "redo"), ("Ctrl+Shift+Z", "redo"), ("Ctrl+F", "search"), ("Ctrl+W", "close"), ("Ctrl+Tab", "next_tab"),
                         ("Ctrl+Shift+Tab", "prev_tab"), ("Ctrl+=", "zoomin"), ("Ctrl++", "zoomin"), ("Ctrl+-", "zoomout"), ("Ctrl+0", "zoom_width"),
                         ("Ctrl+1", "zoom_actual"), ("Ctrl+2", "zoom_page"), ("Ctrl+L", "fullscreen"), ("F11", "fullscreen"), ("Ctrl+D", "properties"),
                         ("Ctrl+K", "preferences"), ("Ctrl+G", "goto"), ("Ctrl+N", "create_blank"), ("Ctrl+Shift+N", "create_files"),
                         ("Ctrl+Shift+T", "thumbs"), ("Ctrl+Shift+B", "bookmark"), ("Ctrl+Shift+A", "comments"), ("Ctrl+E", "edit_all_toggle"),
                         ("Ctrl+Shift+H", "tool_hand"), ("Ctrl+Shift+V", "tool_select"), ("Ctrl+Shift+P", "print_batch"),
                         ("Ctrl+Shift+C", "snapshot_tool"), ("Ctrl+B", "bookmark_add"), ("F1", "help")):
            QShortcut(QKeySequence(seq), self, activated=lambda k=key: self.on_action(k))

    # ── 테마 ──────────────────────────────────────────────────
    def apply_theme(self):
        dark = bool(self.prefs.get("dark"))
        QApplication.instance().setStyleSheet(theme.qss(dark))
        self.ribbon.set_toggle("dark", dark)
        for t in self.tabs:
            t.viewer.setBackgroundBrush(QColor(theme.DARK["CANVAS"] if dark else theme.CANVAS))

    # ── 탭·문서 ───────────────────────────────────────────────
    def cur(self) -> DocTab | None:
        i = self.doctabs.currentIndex()
        return self.tabs[i] if 0 <= i < len(self.tabs) else None

    def show_start(self):
        self.body.setCurrentWidget(self.start)
        self.start.reload()
        self.ribbon.set_enabled_all(False)
        self.status.hide()
        self.doctabs.blockSignals(True); self.doctabs.setCurrentIndex(-1); self.doctabs.blockSignals(False)

    def _tab_changed(self, i: int):
        if i < 0 or i >= len(self.tabs):
            return
        t = self.tabs[i]
        self.body.setCurrentWidget(self.body.widget(1))
        self.docstack.setCurrentWidget(t)
        self.ribbon.set_enabled_all(True)
        self.status.show()
        self.setWindowTitle(f"{t.doc.name} - {APP_NAME}")
        for k, b in self.navbtns.items():
            b.setChecked(t.side.isVisible() and t.side_key == k)
        self.rightbtns["props"].setChecked(t.props.isVisible())
        self._sync_status(t)
        self.ribbon.reflect_tool(t.viewer.tool)
        self.ribbon.set_toggle("edit_all", t.viewer.edit_mode); self.ribbon.set_toggle("form_edit", t.viewer.form_mode)
        self.ribbon.set_toggle("show_annots", t.viewer.show_annots)
        self.ribbon.set_enabled("undo", t.doc.can_undo()); self.ribbon.set_enabled("redo", t.doc.can_redo())

    def _tab_moved(self, a: int, b: int):
        self.tabs.insert(b, self.tabs.pop(a))

    def add_doc(self, doc: Document):
        t = DocTab(doc, self)
        t.viewer.pageChanged.connect(lambda i: self._sync_status(t))
        t.viewer.zoomChanged.connect(lambda z: (self._sync_status(t), self._remember_zoom(t)))
        t.viewer.status.connect(self.say)
        t.viewer.requestDialog.connect(lambda kind, data, tab=t: self.on_request(tab, kind, data))
        t.viewer.toolDone.connect(self._tool_done)
        t.viewer.docChanged.connect(lambda: (self.ribbon.set_enabled("undo", t.doc.can_undo()), self.ribbon.set_enabled("redo", t.doc.can_redo())))
        t.viewer.selectionChanged.connect(lambda s: self._on_selection(t, s))
        self.tabs.append(t)
        self.docstack.addWidget(t)
        i = self.doctabs.addTab(doc.name)
        self.doctabs.setTabToolTip(i, doc.path or "")
        pl = self.prefs.get("default_layout", "continuous")
        t.viewer.layout_mode = pl
        zm = self.prefs.get("default_zoom", "last")
        last = self.settings.get("last_zoom") or {}
        if zm == "last":   # 마지막으로 쓴 확대: 맞추기였으면 그 맞추기, 직접 고른 배율이면 그 배율
            zm = last.get("mode", "fit_width")
        t.viewer.zoom_mode = zm
        t.viewer.set_layout_mode(pl)
        t.viewer.props["author"] = self.prefs.get("author", "")
        t.viewer.props["measure_unit"] = self.prefs.get("unit", "cm")
        # apply_theme 은 그때 열린 탭만 칠한다. 어두운 화면에서 새로 연 문서도 문서 뒤 바탕을 맞춘다
        t.viewer.setBackgroundBrush(QColor(theme.DARK["CANVAS"] if self.prefs.get("dark") else theme.CANVAS))
        self.doctabs.setCurrentIndex(i)
        self._tab_changed(i)
        def first_view():
            if t.viewer.zoom_mode == "fixed":
                t.viewer.set_zoom(float(last.get("zoom", 1.0)))
            t.viewer._apply_fit(); t.viewer.goto(0)
        QTimer.singleShot(60, first_view)

    def open_path(self, path: str, password: str | None = None):
        path = os.path.abspath(path)
        for k, t in enumerate(self.tabs):
            if t.doc.path and os.path.abspath(t.doc.path) == path:
                self.doctabs.setCurrentIndex(k); return
        try:
            doc = Document(path, password)
        except NeedsPassword:
            pw = D.ask_password(self, os.path.basename(path))
            if pw is None:
                return
            try:
                doc = Document(path, pw)
            except NeedsPassword:
                QMessageBox.warning(self, "암호", "암호가 맞지 않습니다"); return
            except DocumentError as e:
                QMessageBox.critical(self, "열기", str(e)); return
        except DocumentError as e:
            QMessageBox.critical(self, "열기", str(e)); return
        self.add_doc(doc)
        self.add_recent(path)

    def add_recent(self, path: str):
        rec = self.settings["recent"]
        rec[:] = [r for r in rec if r["path"] != path]
        rec.insert(0, {"path": path, "time": time.time()})
        del rec[self.prefs.get("recent_n", 30):]
        save_settings(self.settings)

    def open_dialog(self):
        start_dir = os.path.dirname(self.settings["recent"][0]["path"]) if self.settings["recent"] else ""
        paths, _ = QFileDialog.getOpenFileNames(self, "열기", start_dir, OPEN_FILTER)
        for p in paths:
            self.open_path(p)

    def close_tab(self, i: int) -> bool:
        t = self.tabs[i]
        if t.doc.modified:
            r = QMessageBox.question(self, "닫기", f"「{t.doc.name}」 을 저장할까요?",
                                     QMessageBox.StandardButton.Save | QMessageBox.StandardButton.Discard | QMessageBox.StandardButton.Cancel)
            if r == QMessageBox.StandardButton.Cancel:
                return False
            if r == QMessageBox.StandardButton.Save and not self.save(t):
                return False
        self.tabs.pop(i)
        self.docstack.removeWidget(t)
        self.doctabs.removeTab(i)
        t.close(); t.deleteLater()
        if not self.tabs:
            self.show_start()
        return True

    def closeEvent(self, e: QCloseEvent):
        self._remember_window()
        self.settings["last_open"] = [t.doc.path for t in self.tabs if t.doc.path]
        if self._receiver:   # 닫는 동안 넘어오는 파일은 받지 않는다(닫히는 창에 열려 사라지지 않게)
            self._receiver.stop()
        while self.tabs:
            if not self.close_tab(len(self.tabs) - 1):
                if self._receiver:
                    self._receiver.listen()
                e.ignore(); return
        save_settings(self.settings)
        if self._receiver:   # 이 뒤에 켜지는 RPDF 는 넘기지 않고 제 창을 띄운다
            from .. import single
            single.release()
        e.accept()

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        for u in e.mimeData().urls():
            p = u.toLocalFile()
            if os.path.isfile(p):
                self.open_path(p)

    def mark_modified(self):
        t = self.cur()
        if not t:
            return
        i = self.tabs.index(t)
        self.doctabs.setTabText(i, ("● " if t.doc.modified else "") + t.doc.name)
        self.ribbon.set_enabled("undo", t.doc.can_undo()); self.ribbon.set_enabled("redo", t.doc.can_redo())

    def say(self, msg: str):
        self.msg_label.setText(msg); self.msg_label.setToolTip(msg)
        QTimer.singleShot(5000, lambda: self.msg_label.setText("") if self.msg_label.text() == msg else None)

    def _zoom_slider(self, v: int):
        t = self.cur()
        if t:
            t.viewer.zoom_mode = "fixed"; t.viewer.set_zoom(v / 100)

    def _remember_zoom(self, t: DocTab):
        if t is self.cur():   # 다음에 여는 문서가 따라갈 확대(환경 설정 「열 때 확대」가 마지막 확대일 때)
            self.settings["last_zoom"] = {"mode": t.viewer.zoom_mode, "zoom": round(t.viewer.zoom, 4)}

    def _sync_status(self, t: DocTab):
        if t is not self.cur():
            return
        v = t.viewer
        w, h = t.doc.page_size(v.current)
        self.size_label.setText(f"{w / 72 * 2.54:.2f} x {h / 72 * 2.54:.2f} cm")
        self.page_edit.setText(str(v.current + 1)); self.page_total.setText(f"/ {t.doc.page_count}")
        self.zoom_slider.blockSignals(True); self.zoom_slider.setValue(int(v.zoom * 100)); self.zoom_slider.blockSignals(False)
        self.zoom_box.setEditText(f"{int(round(v.zoom * 100))}%")
        self.sb_hand.setChecked(v.tool == "hand"); self.sb_sel.setChecked(v.tool == "select")
        i = self.tabs.index(t)
        self.doctabs.setTabText(i, ("● " if t.doc.modified else "") + t.doc.name)

    def _goto_entered(self):
        t = self.cur()
        if not t:
            return
        txt = self.page_edit.text().strip()
        labels = P.labels(t.doc)
        if txt in labels:
            t.viewer.goto(labels.index(txt))
        elif txt.isdigit():
            t.viewer.goto(int(txt) - 1)

    def _zoom_box(self, *_):
        t = self.cur()
        if not t:
            return
        txt = self.zoom_box.currentText().strip()
        if txt == "창 너비":
            t.viewer.set_zoom_mode("fit_width")
        elif txt == "창 크기":
            t.viewer.set_zoom_mode("fit_page")
        elif txt == "실제 크기":
            t.viewer.set_zoom_mode("actual")
        else:
            try:
                t.viewer.zoom_mode = "fixed"; t.viewer.set_zoom(float(txt.rstrip("%")) / 100)
            except ValueError:
                pass

    def toggle_side(self, key: str):
        t = self.cur()
        if not t:
            return
        t.show_side(key)
        for k, b in self.navbtns.items():
            b.setChecked(t.side.isVisible() and t.side_key == k)

    def right_action(self, key: str):
        t = self.cur()
        if not t:
            return
        if key == "props":
            t.props.setVisible(not t.props.isVisible()); self.rightbtns["props"].setChecked(t.props.isVisible())
        elif key == "fields":
            self.toggle_side("fields"); self.rightbtns["fields"].setChecked(t.side.isVisible() and t.side_key == "fields")
        elif key == "sign":
            self.on_action("sign_apply")
        elif key == "translate":
            s = t.viewer.selected_text()
            if s:
                webbrowser.open("https://papago.naver.com/?sk=auto&tk=ko&st=" + urllib.parse.quote(s[:500]))
            else:
                self.say("먼저 글자를 고르세요(선택 도구)")

    def _on_selection(self, t: DocTab, s):
        if s is not None and s.kind == "annot" and not t.props.isVisible() and t.viewer.tool == "select":
            pass  # 속성 패널은 사용자가 켠 것만 보인다
        self.ribbon.set_enabled("delete_sel", s is not None)

    def _tool_done(self, tool: str):
        t = self.cur()
        if not t:
            return
        # 한 번 쓰는 도구는 선택으로 돌아간다. 연필·강조·스탬프 등은 계속.
        if tool in ("textbox", "callout", "typewriter", "addtext", "addimage", "link", "crop", "snapshot", "ocrarea", "signature",
                    "escape") or tool.startswith("f"):
            self.on_tool("select" if not t.viewer.form_mode else "form")

    # ── 리본 처리 ─────────────────────────────────────────────
    def on_ribbon_tab(self, key: str):
        t = self.cur()
        if not t:
            return
        if key == "page":
            # 구성 화면은 문서가 바뀔 때마다 맞춰 두므로(changed) 다 그린 썸네일은 그대로 쓴다
            t.view_stack.setCurrentWidget(t.organize); t.organize.reload(list(range(t.doc.page_count)))
        else:
            t.view_stack.setCurrentWidget(t.viewer)
        if key == "form" and not t.viewer.form_mode:
            pass
        if key != "form" and t.viewer.form_mode:
            t.viewer.set_form_mode(False); self.ribbon.set_toggle("form_edit", False)
        if key != "edit" and key != "home" and t.viewer.edit_mode:
            t.viewer.set_edit_mode(False); self.ribbon.set_toggle("edit_all", False)

    def on_tool(self, key: str):
        t = self.cur()
        if not t:
            return
        if key == "stamp_pick":
            text, ok = QInputDialog.getItem(self, "스탬프", "글자", A.STAMP_TEXTS, 0, True)
            if not ok:
                return
            t.viewer.props["stamp_text"] = text; t.viewer.props["stamp_image"] = ""; key = "stamp"
        elif key == "stamp_image":
            p, _ = QFileDialog.getOpenFileName(self, "스탬프 이미지", "", "이미지 (*.png *.jpg *.jpeg *.bmp *.gif)")
            if not p:
                return
            t.viewer.props["stamp_image"] = p; key = "stamp"
        elif key == "addimage":
            p, _ = QFileDialog.getOpenFileName(self, "넣을 이미지", "", "이미지 (*.png *.jpg *.jpeg *.bmp *.gif *.tif *.webp)")
            if not p:
                return
            t.viewer.props["image_path"] = p
        elif key == "signature":
            p = self.prefs.get("signature_image")
            if not p or not os.path.exists(p):
                p, _ = QFileDialog.getOpenFileName(self, "서명 이미지(투명 PNG 권장)", "", "이미지 (*.png *.jpg *.jpeg *.bmp)")
                if not p:
                    return
                self.prefs["signature_image"] = p; save_settings(self.settings)
            t.viewer.props["stamp_image"] = p
        if key.startswith("f") and key in TOOL_KEYS and not t.viewer.form_mode:
            t.viewer.set_form_mode(True); self.ribbon.set_toggle("form_edit", True)
        if key in ("hand", "select") and t.viewer.form_mode:
            t.viewer.set_form_mode(False); self.ribbon.set_toggle("form_edit", False)
        if key in ("hand", "select") and t.viewer.edit_mode:
            t.viewer.set_edit_mode(False); self.ribbon.set_toggle("edit_all", False)
        t.viewer.set_tool(key)
        self.ribbon.reflect_tool(key)
        self._sync_status(t)
        if key in LABELS and key not in ("hand", "select", "form", "edit"):
            self.say(f"{LABELS[key]}: 문서 위에서 끌어 그리세요. Esc 로 그만둡니다.")

    def on_toggle(self, key: str, on: bool):
        t = self.cur()
        if key == "dark":
            self.prefs["dark"] = on; save_settings(self.settings); self.apply_theme(); return
        if not t:
            return
        if key == "edit_all":
            t.viewer.set_edit_mode(on); self.ribbon.reflect_tool(t.viewer.tool)
            self.say("글자 덩어리를 두 번 눌러 고치고, 그림은 눌러 고른 뒤 Delete 로 지웁니다." if on else "")
        elif key == "form_edit":
            t.viewer.set_form_mode(on); self.ribbon.reflect_tool(t.viewer.tool)
        elif key == "show_annots":
            # 보기 설정이다. 파일의 주석 숨김 표시를 바꾸면 저장할 때 주석이 숨은 채로 남는다.
            t.viewer.show_annots = on
            t.viewer.refresh_pages(None)
        elif key == "read_mode":
            self.ribbon.stack.setVisible(not on); self.navstrip.setVisible(not on); self.rightstrip.setVisible(not on)
            if on:
                t.side.hide(); t.props.hide()

    def on_action(self, key: str):  # noqa: C901
        t = self.cur()
        v = t.viewer if t else None
        d = t.doc if t else None
        try:
            # 파일
            if key == "open": self.open_dialog()
            elif key == "save" and t: self.save(t)
            elif key == "saveas" and t: self.save(t, ask=True)
            elif key == "close" and t: self.close_tab(self.tabs.index(t))
            elif key == "quit": self.close()
            elif key == "print" and t: self.print_doc(t)
            elif key == "print_batch": self.batch_dialog()
            elif key == "properties" and t: self.properties(t)
            elif key == "preferences": self.preferences()
            elif key == "help": self.show_help()
            elif key == "mail" and t: self.mail(t)
            elif key == "next_tab" and self.tabs: self.doctabs.setCurrentIndex((self.doctabs.currentIndex() + 1) % len(self.tabs))
            elif key == "prev_tab" and self.tabs: self.doctabs.setCurrentIndex((self.doctabs.currentIndex() - 1) % len(self.tabs))
            elif key == "undo" and t:
                if d.undo(): t.viewer.clear_selection(); t.changed(None, structural=True)
            elif key == "redo" and t:
                if d.redo(): t.viewer.clear_selection(); t.changed(None, structural=True)
            # 생성
            elif key == "create_blank": self.create_blank()
            elif key == "create_files": self.create_from_files()
            elif key == "create_clipboard": self.create_from_clipboard()
            elif key == "create_scan": QMessageBox.information(self, "스캐너", "스캐너 연동은 넣지 않았습니다. 스캔 프로그램으로 이미지를 받은 뒤 「파일에서 생성」 으로 PDF 를 만드세요.")
            # 보기
            elif not t: return
            elif key == "zoomin": v.zoom_in()
            elif key == "zoomout": v.zoom_out()
            elif key == "zoom_actual": v.set_zoom_mode("actual")
            elif key == "zoom_width": v.set_zoom_mode("fit_width")
            elif key == "zoom_height": v.set_zoom_mode("fit_height")
            elif key == "zoom_page": v.set_zoom_mode("fit_page")
            elif key.startswith("layout_"): v.set_layout_mode(key[len("layout_"):])
            elif key.startswith("canvas_"):
                col = {"gray": theme.CANVAS, "white": "#FFFFFF", "dark": "#2B2B2E", "sepia": "#F1E7D0"}[key[7:]]
                v.setBackgroundBrush(QColor(col))
            elif key.startswith("rotate_view_"):
                # 보기 회전 = 모든 쪽 회전(저장되는 회전). 알PDF 도 저장되는 쪽으로 동작한다.
                P.rotate(d, range(d.page_count), 90 if key.endswith("cw") else -90); t.changed(None, structural=True)
            elif key == "fullscreen":
                if self.isFullScreen(): self.showNormal(); self.ribbon.show(); self.status.show()
                else: self.showFullScreen(); self.ribbon.hide()
            elif key == "auto_bookmark": self.auto_bookmark(t)
            elif key == "toggle_grid": self.say("격자는 아직 그리지 않습니다")
            elif key == "toggle_ruler": self.say("눈금자는 아직 그리지 않습니다")
            elif key in ("tile_h", "tile_v"): self.tile(key == "tile_h")
            elif key == "first": v.goto(0)
            elif key == "last": v.goto(d.page_count - 1)
            elif key == "next": v.next_page()
            elif key == "prev": v.prev_page()
            elif key == "goto":
                n, ok = QInputDialog.getInt(self, "쪽으로 이동", "쪽 번호", v.current + 1, 1, d.page_count)
                if ok: v.goto(n - 1)
            elif key in ("thumbs", "bookmark", "comments", "attach", "search", "layers"): self.toggle_side(key)
            elif key == "bookmark_add": t.show_side("bookmark") if not t.side.isVisible() else None; t.bookmarks.add()
            elif key == "edit_all_toggle": on = not v.edit_mode; v.set_edit_mode(on); self.ribbon.set_toggle("edit_all", on)
            elif key == "tool_hand": self.on_tool("hand")
            elif key == "tool_select": self.on_tool("select")
            elif key == "snapshot_tool": self.on_tool("snapshot")
            elif key == "delete_sel": v.delete_selection()
            elif key == "replace_image":
                if v.sel and v.sel.kind == "image": self.on_request(t, "replace_image", v.sel)
                else: self.say("편집 모드에서 이미지를 먼저 고르세요")
            # 편집 꾸미기
            elif key == "watermark": self.run_tool(t, "watermark", D.WatermarkDialog(self, d.page_count, v.current, v.props))
            elif key == "background": self.run_tool(t, "background", D.BackgroundDialog(self, d.page_count, v.current))
            elif key == "headerfooter": self.run_tool(t, "header_footer", D.HeaderFooterDialog(self, d.page_count, v.current))
            elif key == "bates": self.run_tool(t, "bates_numbers", D.BatesDialog(self, d.page_count, v.current))
            elif key == "pagelabel":
                dlg = D.PageLabelDialog(self, d.page_count, v.current)
                if dlg.exec(): P.set_labels(d, **dlg.values()); t.changed(None, structural=True)
            # 구성
            elif key in ("thumb_in", "thumb_out"): t.organize.set_width(t.organize.thumb_w + (30 if key == "thumb_in" else -30))
            elif key in ("rotate_left", "rotate_right", "delete_pages", "extract", "insert_blank", "insert_pdf", "insert_images", "crop_pages",
                         "resize", "replace", "duplicate", "pagebox", "delete_blank", "reverse", "rotate_pages", "split"):
                pages = t.organize.selected_pages() if t.view_stack.currentWidget() is t.organize else []
                if not pages and t.side.isVisible() and t.side_key == "thumbs":
                    pages = t.thumbs.selected_pages()
                self.pages_action({"rotate_left": "rotate_left", "rotate_right": "rotate_right", "delete_pages": "delete", "extract": "extract",
                                   "crop_pages": "crop", "rotate_pages": "rotate", "delete_blank": "delete_blank"}.get(key, key), pages or [v.current])
            # 변환
            elif key.startswith("to_"): self.convert(t, key[3:])
            elif key == "convert_settings": self.convert(t, "docx")
            elif key == "batch": self.batch_dialog()
            # 도구
            elif key == "ocr": self.ocr(t)
            elif key == "merge": self.merge()
            elif key == "compare": self.compare(t)
            elif key == "compress": self.compress(t)
            elif key == "flatten": self.flatten(t)
            elif key == "deskew": self.deskew(t)
            elif key == "table": self.table_recognize(t)
            elif key == "extract_data": self.extract_data()
            elif key == "search": t.show_side("search") if not (t.side.isVisible() and t.side_key == "search") else t.search.focus(); self.navbtns["search"].setChecked(True)
            # 양식
            elif key == "recognize":
                n = A.recognize_fields(d); t.changed(None); self.say(f"양식 칸 {n}개를 만들었습니다")
            elif key == "form_export": self.form_data(t, "export")
            elif key == "form_import": self.form_data(t, "import")
            elif key == "form_reset": A.reset_fields(d); t.changed(None)
            elif key == "form_preview": v.set_form_mode(False); self.ribbon.set_toggle("form_edit", False); self.say("미리보기: 칸을 두 번 눌러 값을 넣습니다")
            elif key.startswith("align_"):
                items = [(v.sel.page, v.sel.xref)] if v.sel and v.sel.kind == "widget" else []
                if len(items) < 1: self.say("맞출 양식 칸을 먼저 고르세요"); return
                A.align_fields(d, items, key[6:]); t.changed([v.sel.page])
            # 보안
            elif key == "encrypt": self.encrypt(t)
            elif key == "decrypt": self.decrypt(t)
            elif key == "redact_search":
                dlg = D.RedactSearchDialog(self, d.page_count, v.current)
                if dlg.exec():
                    vals = dlg.values(); T = _engine("tools")
                    d.snapshot(); n = T.redact_search(d.fitz, vals["pattern"], regex=vals["regex"], pages=vals["pages"], fill=vals["fill"])
                    t.changed(None); self.say(f"{n}곳에 교정 표시를 했습니다. 「교정 적용」 을 누르면 지워집니다.")
            elif key == "redact_apply":
                if QMessageBox.question(self, "교정 적용", "교정 표시한 자리의 글자·그림을 영구히 지웁니다. 계속할까요?") == QMessageBox.StandardButton.Yes:
                    A.apply_redactions(d); t.changed(None)
            elif key == "hidden":
                dlg = D.HiddenDataDialog(self)
                if dlg.exec():
                    T = _engine("tools"); d.snapshot(); T.remove_hidden_data(d.fitz, **dlg.values()); t.changed(None, structural=True)
            elif key == "digital_sign":
                QMessageBox.information(self, "디지털 서명", "인증서 기반 디지털 서명은 넣지 않았습니다(공인 인증서·시점 확인 서버가 필요합니다).\n"
                                        "손글씨·도장 그림을 넣는 「서명 관리 → 서명 넣기」 를 쓰세요.")
            elif key == "sign_new":
                p, _ = QFileDialog.getOpenFileName(self, "서명 이미지(투명 PNG 권장)", "", "이미지 (*.png *.jpg *.jpeg *.bmp)")
                if p: self.prefs["signature_image"] = p; save_settings(self.settings); self.say("서명 이미지를 등록했습니다")
            elif key == "sign_apply": self.on_tool("signature")
            elif key == "sign_clear": self.prefs.pop("signature_image", None); save_settings(self.settings)
            elif key == "crop": self.on_tool("crop")
            elif key in TOOL_KEYS: self.on_tool(key)
            else: self.say(f"아직 없는 기능: {key}")
        except RuntimeError as e:
            QMessageBox.warning(self, "아직 못 씀", str(e))
        except ValueError as e:  # 엔진이 내는 입력 확인 문구(「찾을 글자를 넣어 주세요」 등)
            QMessageBox.warning(self, "확인", str(e))
        except Exception as e:  # noqa: BLE001
            if type(e).__name__ in ("EngineError", "DocumentError"):  # 한국어로 쓴 안내(코드 흔적 없이)
                QMessageBox.warning(self, "확인", str(e)); return
            import traceback
            QMessageBox.critical(self, "오류", f"{e}\n\n{traceback.format_exc()[-1200:]}")

    # ── 페이지 동작(썸네일·구성·리본 공용) ───────────────────────
    def pages_action(self, key: str, pages: list[int]):  # noqa: C901
        t = self.cur()
        if not t:
            return
        d, v = t.doc, t.viewer
        pages = sorted(set(pages)) or [v.current]
        cur = pages[0]
        try:
            if key == "rotate_left": P.rotate(d, pages, -90); t.changed(pages, structural=True)
            elif key == "rotate_right": P.rotate(d, pages, 90); t.changed(pages, structural=True)
            elif key == "rotate":
                deg, ok = QInputDialog.getItem(self, "회전", f"{len(pages)}쪽 회전", ["오른쪽 90°", "왼쪽 90°", "180°"], 0, False)
                if ok: P.rotate(d, pages, {"오른쪽 90°": 90, "왼쪽 90°": -90, "180°": 180}[deg]); t.changed(pages, structural=True)
            elif key == "delete":
                if d.page_count <= len(pages): self.say("쪽을 전부 지울 수는 없습니다"); return
                if len(pages) > 1 and QMessageBox.question(self, "페이지 삭제", f"{len(pages)}쪽을 지울까요?") != QMessageBox.StandardButton.Yes: return
                gone = set(pages); keep = [i for i in range(d.page_count) if i not in gone]
                P.delete(d, pages); v.clear_selection(); t.changed(None, structural=True, mapping=keep)
            elif key == "extract":
                out, _ = QFileDialog.getSaveFileName(self, "쪽 내보내기", self._suggest(t, "_추출.pdf"), PDF_FILTER)
                if out: P.extract(d, pages, out); self.say(f"{len(pages)}쪽을 저장했습니다: {out}"); self._offer_open(out)
            elif key == "insert_blank":
                dlg = D.InsertBlankDialog(self, d.page_count, cur)
                if dlg.exec(): P.insert_blank(d, **dlg.values()); t.changed(None, structural=True)
            elif key == "insert_pdf":
                p, _ = QFileDialog.getOpenFileName(self, "끼울 PDF", "", OPEN_FILTER)
                if not p: return
                src = fitz.open(p); n = src.page_count; src.close()
                dlg = D.InsertPdfDialog(self, d.page_count, cur, p, n)
                if dlg.exec(): P.insert_from_pdf(d, path=p, **dlg.values()); t.changed(None, structural=True)
            elif key == "insert_images":
                paths, _ = QFileDialog.getOpenFileNames(self, "끼울 이미지", "", "이미지 (*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.gif *.webp)")
                if paths: P.insert_images(d, cur + 1, paths); t.changed(None, structural=True)
            elif key == "duplicate":
                dup = set(pages); m = [k for i in range(d.page_count) for k in ((i, i) if i in dup else (i,))]
                P.duplicate(d, pages); t.changed(None, structural=True, mapping=m)
            elif key == "replace":
                p, _ = QFileDialog.getOpenFileName(self, "바꿀 내용이 든 PDF", "", OPEN_FILTER)
                if not p: return
                src = fitz.open(p); n = src.page_count; src.close()
                rng, ok = QInputDialog.getText(self, "바꾸기", f"{len(pages)}쪽을 「{os.path.basename(p)}」 의 어느 쪽으로 바꿀까요? (비우면 전부, 1-{n})")
                if ok:
                    P.replace(d, pages, p, P.parse_range(rng, n) if rng.strip() else None); t.changed(None, structural=True)
            elif key == "crop":
                # 자르기 상자는 지금 보이는 영역 기준의 회전 전 좌표다(page.rect 는 회전 뒤 크기라 MediaBox 를 넘었다)
                cb = d.fitz[cur].cropbox
                dlg = D.CropDialog(self, d.page_count, cur, None, fitz.Rect(0, 0, cb.width, cb.height))
                if dlg.exec(): vals = dlg.values(); P.crop(d, vals["pages"], vals["rect"]); t.changed(None, structural=True)
            elif key == "resize":
                dlg = D.ResizeDialog(self, d.page_count, cur)
                if dlg.exec(): vals = dlg.values(); P.resize(d, vals["pages"], vals["width"], vals["height"], vals["fit"]); t.changed(None, structural=True)
            elif key == "pagebox":
                dlg = D.PageBoxDialog(self, d.page_count, cur, d.fitz[cur].cropbox)
                if dlg.exec(): vals = dlg.values(); P.set_box(d, vals["pages"], vals["box"], vals["rect"]); t.changed(None, structural=True)
            elif key == "pagelabel": self.on_action("pagelabel")
            elif key == "delete_blank":
                T = _engine("tools"); d.snapshot(); gone = T.delete_blank_pages(d.fitz)
                t.changed(None, structural=True); self.say(f"빈 쪽 {len(gone)}장을 지웠습니다")
            elif key == "reverse":
                rev = list(range(d.page_count - 1, -1, -1)); P.reorder(d, rev); t.changed(None, structural=True, mapping=rev)
            elif key == "split":
                dlg = D.SplitDialog(self, d.page_count, os.path.dirname(d.path) if d.path else os.path.expanduser("~"))
                if dlg.exec():
                    vals = dlg.values(); outs = P.split(d, vals["out_dir"], vals["mode"], vals["value"], vals["ranges"])
                    self.say(f"{len(outs)}개 파일로 나눴습니다"); self._offer_open(vals["out_dir"])
        except ValueError as e:
            QMessageBox.warning(self, "페이지", str(e))
        except RuntimeError as e:
            QMessageBox.warning(self, "아직 못 씀", str(e))

    # ── 뷰어가 부탁한 대화상자 ─────────────────────────────────
    def on_request(self, t: DocTab, kind: str, data):  # noqa: C901
        d, v = t.doc, t.viewer
        if kind == "textbox":
            dlg = D.TextBoxDialog(self, data["kind"], v.props)
            if dlg.exec():
                vals = dlg.values()
                if vals["text"].strip():
                    callout = None
                    if data["kind"] == "callout":
                        r = data["rect"]; callout = fitz.Point(r.x0 - 40, r.y1 + 40)
                    A.free_text(d, data["page"], data["rect"], vals["text"], vals["fontsize"], vals["text_color"], vals["fill"],
                                vals["border"], vals["width"], v.props["opacity"], callout=callout, author=v.props["author"],
                                typewriter=(data["kind"] == "typewriter"))
                    t.changed([data["page"]])
            self.on_tool("select")
        elif kind == "addtext":
            dlg = D.TextBoxDialog(self, "addtext", v.props)
            if dlg.exec():
                vals = dlg.values()
                if vals["text"].strip():
                    A.add_text(d, data["page"], data["rect"], vals["text"], vals["fontsize"], vals["text_color"] or (0, 0, 0), vals["bold"], vals["align"])
                    t.changed([data["page"]])
            self.on_tool("select" if not v.edit_mode else "edit")
        elif kind == "edit_block":
            s = data
            dlg = D.TextBoxDialog(self, "edit_block", {"fontsize": s.extra.get("fontsize", 11), "text_color": (0, 0, 0)}, s.text)
            dlg.bold.setChecked(s.extra.get("bold", False))
            if dlg.exec():
                vals = dlg.values()
                A.edit_block(d, s.page, s.rect, vals["text"], vals["fontsize"], vals["text_color"] or (0, 0, 0), vals["bold"], vals["align"])
                v.clear_selection(); t.changed([s.page])
        elif kind == "addimage":
            p = v.props.get("image_path") or QFileDialog.getOpenFileName(self, "넣을 이미지", "", "이미지 (*.png *.jpg *.jpeg *.bmp *.gif *.tif *.webp)")[0]
            if p:
                A.add_image(d, data["page"], data["rect"], p); t.changed([data["page"]])
            self.on_tool("select" if not v.edit_mode else "edit")
        elif kind == "replace_image":
            s = data
            p, _ = QFileDialog.getOpenFileName(self, "바꿀 이미지", "", "이미지 (*.png *.jpg *.jpeg *.bmp *.gif *.tif *.webp)")
            if p:
                if s.extra.get("xref"):
                    A.replace_image(d, s.page, s.extra["xref"], p)
                else:
                    A.delete_image(d, s.page, s.rect); A.add_image(d, s.page, s.rect, p)
                v.clear_selection(); t.changed([s.page])
        elif kind == "save_image":
            s = data
            out, _ = QFileDialog.getSaveFileName(self, "이미지 저장", "이미지.png", "PNG (*.png)")
            if out:
                if s.extra.get("xref"):
                    pix = fitz.Pixmap(d.fitz, s.extra["xref"])
                    if pix.n - pix.alpha >= 4: pix = fitz.Pixmap(fitz.csRGB, pix)
                    pix.save(out)
                else:
                    d.render(s.page, 3.0, clip=s.rect).save(out)
                self.say(f"저장했습니다: {out}")
        elif kind == "signature":
            p = v.props.get("stamp_image") or self.prefs.get("signature_image")
            if p and os.path.exists(p):
                A.signature_image(d, data["page"], data["rect"], p); t.changed([data["page"]])
            self.on_tool("select")
        elif kind == "note":
            text, ok = QInputDialog.getMultiLineText(self, "노트", "내용")
            if ok and text.strip():
                A.note(d, data["page"], data["point"], text, color=v.props["hl_color"], author=v.props["author"]); t.changed([data["page"]])
        elif kind == "caret":
            text, ok = QInputDialog.getText(self, "삽입 기호", "끼워 넣을 글자(메모)")
            if ok:
                A.caret(d, data["page"], data["point"], text, v.props["author"]); t.changed([data["page"]])
        elif kind == "attach":
            p, _ = QFileDialog.getOpenFileName(self, "첨부할 파일")
            if p:
                A.attachment(d, data["page"], data["point"], p, v.props["author"]); t.changed([data["page"]])
        elif kind == "link":
            dlg = D.LinkDialog(self, d.page_count)
            if dlg.exec():
                A.add_link(d, data["page"], data["rect"], **dlg.values()); t.changed([data["page"]])
            self.on_tool("select")
        elif kind == "crop":
            cb = d.fitz[data["page"]].cropbox
            dlg = D.CropDialog(self, d.page_count, data["page"], data["rect"], fitz.Rect(0, 0, cb.width, cb.height))
            if dlg.exec():
                vals = dlg.values(); P.crop(d, vals["pages"], vals["rect"]); t.changed(None, structural=True)
            self.on_tool("select")
        elif kind == "ocrarea":
            try:
                O = _engine("ocr")
                text = O.ocr_area(d.fitz, data["page"], data["rect"], lang=self.prefs.get("ocr_lang", "ko"))
            except Exception as e:  # noqa: BLE001
                QMessageBox.warning(self, "영역 OCR", str(e)); return
            dlg = D.FormDialog("영역 OCR 결과", self, "복사")
            from PySide6.QtWidgets import QPlainTextEdit
            te = QPlainTextEdit(text); te.setMinimumSize(420, 200); dlg.form.addRow(te)
            if dlg.exec():
                QApplication.clipboard().setText(te.toPlainText()); self.say("복사했습니다")
            self.on_tool("select")
        elif kind == "annot_props":
            s = data
            dlg = D.AnnotPropsDialog(self, s)
            if dlg.exec():
                vals = dlg.values()
                A.update(d, s.page, s.xref, stroke=vals["stroke"], fill=vals["fill"], width=vals["width"], opacity=vals["opacity"],
                         content=vals["content"], author=vals["author"])
                t.changed([s.page])
        elif kind == "field_props":
            s = data
            pg = d.fitz[s.page]  # 쪽 객체를 붙잡아 둔다(놓치면 칸이 풀려 w.update 가 멈춘다)
            w = next((w for w in pg.widgets() if w.xref == s.xref), None)
            if not w: return
            radio = w.field_type == fitz.PDF_WIDGET_TYPE_RADIOBUTTON
            dlg = D.FieldPropsDialog(self, w.field_name, w.field_type_string, w.field_value, w.choice_values or [], radio=radio)
            if dlg.exec():
                vals = dlg.values(); d.snapshot()
                if w.field_type in (fitz.PDF_WIDGET_TYPE_COMBOBOX, fitz.PDF_WIDGET_TYPE_LISTBOX): w.choice_values = vals["options"]
                if w.field_type in (fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.PDF_WIDGET_TYPE_RADIOBUTTON):
                    v = vals["value"].strip()
                    w.field_value = v.lower() in ("yes", "on", "true", "1", "예") or v == (w.on_state() or "")
                elif w.field_type != fitz.PDF_WIDGET_TYPE_BUTTON: w.field_value = vals["value"]
                A.save_widget(w)
                A.rename_field(d, s.page, s.xref, vals["name"])  # 라디오는 같은 이름 묶음으로 옮겨 간다
                t.changed([s.page])
        elif kind == "fill_field":
            w = data["widget"]
            val = D.fill_field(self, w)
            if val is not None:
                A.set_field_value(d, data["page"], data["xref"], val); t.changed([data["page"]])
        elif kind == "open_uri":
            if QMessageBox.question(self, "링크", f"이 주소를 열까요?\n{data}") == QMessageBox.StandardButton.Yes:
                webbrowser.open(data)
        elif kind == "open_file":
            if data and os.path.exists(data) and QMessageBox.question(self, "링크", f"이 파일을 열까요?\n{data}") == QMessageBox.StandardButton.Yes:
                os.startfile(data)
        elif kind == "rotate": P.rotate(d, data["pages"], data["deg"]); t.changed(list(data["pages"]), structural=True)
        elif kind == "delete_pages": self.pages_action("delete", data["pages"])

    # ── 파일 메뉴 ─────────────────────────────────────────────
    def show_file_menu(self):
        m = QMenu(self)
        m.setMinimumWidth(240)
        t = self.cur()
        m.addAction(icons.icon("open"), "열기\tCtrl+O", self.open_dialog)
        rm = m.addMenu(icons.icon("recent"), "최근 파일 열기")
        for r in self.settings["recent"][:15]:
            rm.addAction(os.path.basename(r["path"]), lambda p=r["path"]: self.open_path(p))
        cm = m.addMenu(icons.icon("create"), "생성")
        cm.addAction("파일에서 생성", lambda: self.on_action("create_files")); cm.addAction("빈 PDF\tCtrl+N", lambda: self.on_action("create_blank"))
        cm.addAction("클립보드에서 가져오기", lambda: self.on_action("create_clipboard")); cm.addAction("스캐너에서 생성", lambda: self.on_action("create_scan"))
        m.addSeparator()
        a = m.addAction(icons.icon("save"), "저장\tCtrl+S", lambda: self.on_action("save")); a.setEnabled(bool(t))
        a = m.addAction(icons.icon("saveas"), "다른 이름으로 저장\tCtrl+Shift+S", lambda: self.on_action("saveas")); a.setEnabled(bool(t))
        vm = m.addMenu("다른 버전으로 저장"); vm.setEnabled(bool(t))
        vm.addAction("PDF/A", lambda: self.on_action("to_pdfa")); vm.addAction("압축한 PDF", lambda: self.on_action("compress")); vm.addAction("평면화한 PDF", lambda: self.on_action("flatten"))
        sm = m.addMenu(icons.icon("folder"), "최근 폴더에 저장"); sm.setEnabled(bool(t))
        seen = []
        for r in self.settings["recent"]:
            dd = os.path.dirname(r["path"])
            if dd not in seen and os.path.isdir(dd):
                seen.append(dd); sm.addAction(dd, lambda p=dd: self.save_to_dir(p))
            if len(seen) >= 8: break
        if self.prefs.get("share_dir"):
            m.addAction(icons.icon("share"), "공유폴더에 저장", lambda: self.save_to_dir(self.prefs["share_dir"])).setEnabled(bool(t))
        m.addSeparator()
        em = m.addMenu(icons.icon("export"), "내보내기"); em.setEnabled(bool(t))
        for label, key in (("Word", "to_docx"), ("Excel", "to_xlsx"), ("PowerPoint", "to_pptx"), ("한글(HWPX)", "to_hwpx"), ("이미지", "to_png"),
                           ("텍스트", "to_txt"), ("HTML", "to_html"), ("RTF", "to_rtf"), ("EPUB", "to_epub"), ("PDF/A", "to_pdfa")):
            em.addAction(label, lambda k=key: self.on_action(k))
        a = m.addAction(icons.icon("print"), "인쇄하기\tCtrl+P", lambda: self.on_action("print")); a.setEnabled(bool(t))
        a = m.addAction(icons.icon("mail"), "메일로 보내기", lambda: self.on_action("mail")); a.setEnabled(bool(t))
        m.addSeparator()
        a = m.addAction(icons.icon("info"), "파일 속성\tCtrl+D", lambda: self.on_action("properties")); a.setEnabled(bool(t))
        m.addAction(icons.icon("settings"), "환경 설정\tCtrl+K", lambda: self.on_action("preferences"))
        m.addAction(icons.icon("update"), self.updates.menu_label(), self.updates.check_now)
        m.addSeparator()
        a = m.addAction("닫기\tCtrl+W", lambda: self.on_action("close")); a.setEnabled(bool(t))
        m.addAction("끝내기", self.close)
        m.exec(self.ribbon.mapToGlobal(self.ribbon.rect().bottomLeft() - self.ribbon.rect().bottomLeft() + QPointF(6, 30).toPoint()))

    # ── 저장 ──────────────────────────────────────────────────
    def save(self, t: DocTab, ask: bool = False) -> bool:
        path = t.doc.path
        if ask or not path or not path.lower().endswith(".pdf"):
            path, _ = QFileDialog.getSaveFileName(self, "저장", self._suggest(t, ".pdf") if not path else path, PDF_FILTER)
            if not path:
                return False
        try:
            t.doc.save(path)
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "저장", f"저장하지 못했습니다: {e}")
            return False
        self.add_recent(path)
        i = self.tabs.index(t)
        self.doctabs.setTabText(i, t.doc.name); self.doctabs.setTabToolTip(i, path)
        self.setWindowTitle(f"{t.doc.name} - {APP_NAME}")
        t.viewer.refresh_pages(None)
        self.say(f"저장했습니다: {path}")
        return True

    def save_to_dir(self, folder: str):
        t = self.cur()
        if not t:
            return
        path = os.path.join(folder, t.doc.name if t.doc.name.lower().endswith(".pdf") else "문서.pdf")
        if os.path.exists(path) and QMessageBox.question(self, "저장", f"같은 이름이 있습니다. 덮어쓸까요?\n{path}") != QMessageBox.StandardButton.Yes:
            return
        t.doc.save(path); self.add_recent(path); self.say(f"저장했습니다: {path}")

    def _suggest(self, t: DocTab, suffix: str) -> str:
        base = os.path.splitext(t.doc.path)[0] if t.doc.path else os.path.join(self.prefs.get("default_dir") or os.path.expanduser("~"), "문서")
        return base + suffix

    def _offer_open(self, path: str):
        if QMessageBox.question(self, "완료", f"저장했습니다.\n{path}\n\n폴더를 열까요?") == QMessageBox.StandardButton.Yes:
            folder = path if os.path.isdir(path) else os.path.dirname(path)
            if os.path.isfile(path):
                subprocess.Popen(["explorer", "/select,", os.path.normpath(path)])
            else:
                os.startfile(folder)

    # ── 생성 ──────────────────────────────────────────────────
    def create_blank(self):
        dlg = D.InsertBlankDialog(self, 1, 0); dlg.setWindowTitle("빈 PDF"); dlg.where.hide(); dlg.form.labelForField(dlg.where).hide()
        dlg.size.setCurrentText("A4")
        if dlg.exec():
            vals = dlg.values()
            doc = Document.blank(vals["width"] or 595, vals["height"] or 842, vals["count"])
            self.add_doc(doc)

    def create_from_files(self):
        dlg = D.FileListDialog(self, "파일에서 생성", "만들기",
                               "문서·이미지 (*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff *.gif *.webp *.docx *.doc *.xlsx *.xls *.pptx *.ppt *.hwp *.hwpx *.txt *.html *.htm *.rtf *.epub *.xps);;모든 파일 (*)")
        if not dlg.exec() or not dlg.paths():
            return
        paths = dlg.paths()
        out, _ = QFileDialog.getSaveFileName(self, "만들 PDF", os.path.splitext(paths[0])[0] + ".pdf", PDF_FILTER)
        if not out:
            return
        C = _engine("convert")
        res = D.ProgressDialog(self, "PDF 생성", C.from_files, paths, out).run()
        if res:
            self.open_path(out)

    def create_from_clipboard(self):
        cb = QApplication.clipboard()
        md = cb.mimeData()
        if md is not None and md.hasImage():
            img = cb.image()
            pfd, png = tempfile.mkstemp(suffix=".png"); os.close(pfd); img.save(png)
            pdfbytes = fitz.open(png).convert_to_pdf(); os.remove(png)
            # 이름(경로)을 주면 Ctrl+S 가 묻지 않고 지금 폴더에 「클립보드.pdf」 로 저장해 버린다
            doc = Document.from_bytes(pdfbytes)
        elif md is not None and md.hasText():
            doc = Document.blank()
            A.add_text(doc, 0, fitz.Rect(56, 56, 539, 786), md.text(), 11)
        else:
            self.say("클립보드에 그림이나 글자가 없습니다"); return
        self.add_doc(doc)

    # ── 변환·도구 ─────────────────────────────────────────────
    def convert(self, t: DocTab, fmt: str):
        d, v = t.doc, t.viewer
        default = self._suggest(t, {"pdfa": "_PDFA.pdf", "png": "", "jpg": ""}.get(fmt, "." + fmt))
        if fmt in ("png", "jpg"):
            default = os.path.dirname(self._suggest(t, ".x")) or os.path.expanduser("~")
        dlg = D.ConvertDialog(self, d.page_count, v.current, fmt, default)
        if not dlg.exec():
            return
        vals = dlg.values(); fmt = vals["fmt"]; out = vals["out"]; pages = vals["pages"]
        if pages == list(range(d.page_count)):
            pages = None
        C = _engine("convert")
        src = d.copy()  # 지금 상태 그대로 변환(열린 문서를 그냥 save 하면 암호 문서가 망가진다)

        def job(progress=None):
            try:
                if pages and fmt in ("rtf", "pdfa", "hwpx"):
                    src.select(pages)  # 이 셋은 쪽 범위를 받지 않아 고른 범위가 무시됐다
                if fmt in ("png", "jpg"):
                    return C.to_images(src, out, fmt=fmt, dpi=vals["dpi"], pages=pages, progress=progress, stem=os.path.splitext(d.name)[0])
                if fmt == "docx": return C.to_word(src, out, pages=pages, progress=progress)
                if fmt == "xlsx": return C.to_excel(src, out, pages=pages, progress=progress)
                if fmt == "pptx": return C.to_pptx(src, out, pages=pages, dpi=vals["dpi"], progress=progress)
                if fmt == "txt": return C.to_text(src, out, pages=pages)
                if fmt == "html": return C.to_html(src, out, pages=pages)
                if fmt == "rtf": return C.to_rtf(src, out)
                if fmt == "epub": return C.to_epub(src, out, pages=pages)
                if fmt == "pdfa": return C.to_pdfa(src, out)
                if fmt == "hwpx": return C.to_hwpx(src, out, progress=progress)
                raise RuntimeError(f"모르는 형식: {fmt}")
            finally:
                src.close()
        res = D.ProgressDialog(self, "변환", job).run()
        if res:
            path = res[0] if isinstance(res, list) else res
            self.say(f"변환했습니다: {path}")
            if vals["open_after"]:
                self._offer_open(path if isinstance(res, str) else out)

    def run_tool(self, t: DocTab, fn_name: str, dlg):
        """워터마크·배경·머리글·베이츠처럼 문서를 바로 고치는 도구."""
        if not dlg.exec():
            return
        vals = dlg.values()
        T = _engine("tools")
        t.doc.snapshot()
        getattr(T, fn_name)(t.doc.fitz, **vals)
        t.changed(None)

    def ocr(self, t: DocTab):
        O = _engine("ocr")
        try:
            langs = O.available_languages()
        except Exception:  # noqa: BLE001
            langs = ["ko", "en"]
        dlg = D.OcrDialog(self, t.doc.page_count, t.viewer.current, langs)
        if not dlg.exec():
            return
        vals = dlg.values()
        # 일꾼에는 사본을 준다. 보는 문서를 그대로 주면 그리는 일꾼과 같은 문서를 동시에 만져 튕길 수 있다
        work = t.doc.copy()
        n = D.ProgressDialog(self, "OCR", O.ocr_pages, work, pages=vals["pages"], lang=vals["lang"], dpi=vals["dpi"], force=vals["force"]).run()
        if n is not None:
            t.doc.adopt(work)   # 되돌리기 사본을 뜬 뒤 바꿔 끼운다
            t.changed(None); self.say(f"글자 {n}개를 인식해 넣었습니다. 이제 찾기·복사가 됩니다.")
        else:
            work.close()

    def merge(self):
        t = self.cur()
        dlg = D.FileListDialog(self, "PDF 병합", "합치기", OPEN_FILTER, paths=[t.doc.path] if t and t.doc.path else None)
        if not dlg.exec() or len(dlg.paths()) < 1:
            return
        paths = dlg.paths()
        out, _ = QFileDialog.getSaveFileName(self, "합친 PDF", os.path.splitext(paths[0])[0] + "_합침.pdf", PDF_FILTER)
        if not out:
            return
        res = D.ProgressDialog(self, "PDF 병합", P.merge, paths, out).run()
        if res:
            self.open_path(out)

    def compare(self, t: DocTab):
        p, _ = QFileDialog.getOpenFileName(self, "비교할 다른 PDF", "", PDF_FILTER)
        if not p:
            return
        other = self._open_other(p)
        if other is None:
            return
        T = _engine("tools")
        mine = t.doc.copy()  # 보는 문서는 그리는 일꾼이 같이 쓴다. 비교 일꾼에는 사본을 준다
        res = D.ProgressDialog(self, "PDF 비교", T.compare, mine, other.fitz).run()
        mine.close(); other.close()
        if res is not None:
            D.CompareDialog(self, t.doc.name, os.path.basename(p), res).exec()

    def _open_other(self, path: str) -> Document | None:
        """비교 등에 쓸 다른 PDF 를 연다. 암호가 걸려 있으면 물어서 열고, 틀리면 다시 묻는다."""
        pw = None
        while True:
            try:
                return Document(path, pw)
            except NeedsPassword:
                if pw is not None:
                    QMessageBox.warning(self, "암호", "암호가 맞지 않습니다")
                pw = D.ask_password(self, os.path.basename(path))
                if pw is None:
                    return None
            except DocumentError as e:
                QMessageBox.critical(self, "열기", str(e))
                return None

    def compress(self, t: DocTab):
        size = os.path.getsize(t.doc.path) if t.doc.path and os.path.exists(t.doc.path) else len(t.doc.to_bytes())
        dlg = D.CompressDialog(self, size)
        if not dlg.exec():
            return
        out, _ = QFileDialog.getSaveFileName(self, "압축한 PDF", self._suggest(t, "_압축.pdf"), PDF_FILTER)
        if not out:
            return
        T = _engine("tools")
        src = t.doc.copy()
        res = D.ProgressDialog(self, "PDF 압축", T.compress, src, out, level=dlg.values()["level"], password=t.doc.password).run()
        src.close()
        if res:
            if res["after"] >= size:
                QMessageBox.information(self, "PDF 압축", f"더 줄일 수 없습니다 · 원본 {size / 1024:.0f} KB, 결과 {res['after'] / 1024:.0f} KB\n{out}")
            else:
                QMessageBox.information(self, "PDF 압축", f"{size / 1024:.0f} KB → {res['after'] / 1024:.0f} KB\n{out}")
            self.open_path(out)

    def flatten(self, t: DocTab):
        out, _ = QFileDialog.getSaveFileName(self, "평면화한 PDF", self._suggest(t, "_평면화.pdf"), PDF_FILTER)
        if not out:
            return
        T = _engine("tools")
        src = t.doc.copy()
        res = D.ProgressDialog(self, "평면화", T.flatten, src, out, password=t.doc.password).run()
        src.close()
        if res:
            self.open_path(out)

    def deskew(self, t: DocTab):
        T = _engine("tools")
        work = t.doc.copy()   # OCR 과 같은 까닭으로 사본에서
        angles = D.ProgressDialog(self, "기울기 보정", T.deskew, work).run()
        if angles is None:
            work.close()
        else:
            t.doc.adopt(work)
            t.changed(None, structural=True); self.say(f"{sum(1 for a in angles if abs(a) > 0.1)}쪽을 바로잡았습니다")

    def table_recognize(self, t: DocTab):
        O = _engine("ocr")
        rows = O.recognize_table(t.doc.fitz, t.viewer.current)
        if not rows:
            self.say("표를 찾지 못했습니다"); return
        out, _ = QFileDialog.getSaveFileName(self, "표 저장", self._suggest(t, f"_{t.viewer.current + 1}쪽_표.xlsx"), "Excel (*.xlsx)")
        if out:
            import openpyxl
            wb = openpyxl.Workbook(); ws = wb.active
            for r in rows: ws.append(r)
            wb.save(out); self._offer_open(out)

    def extract_data(self):
        dlg = D.FileListDialog(self, "데이터 추출(양식 값 모으기)", "추출", PDF_FILTER)
        if not dlg.exec() or not dlg.paths():
            return
        out, _ = QFileDialog.getSaveFileName(self, "엑셀로 저장", "양식 데이터.xlsx", "Excel (*.xlsx)")
        if out:
            T = _engine("tools"); T.extract_data(dlg.paths(), out); self._offer_open(out)

    def form_data(self, t: DocTab, how: str):
        T = _engine("tools")
        if how == "export":
            out, _ = QFileDialog.getSaveFileName(self, "양식 데이터 내보내기", self._suggest(t, "_양식.json"), "JSON (*.json);;CSV (*.csv)")
            if out: T.export_form_data(t.doc.fitz, out); self._offer_open(out)
        else:
            p, _ = QFileDialog.getOpenFileName(self, "양식 데이터 가져오기", "", "JSON (*.json)")
            if p:
                t.doc.snapshot(); n = T.import_form_data(t.doc.fitz, p); t.changed(None); self.say(f"{n}칸을 채웠습니다")

    def batch_dialog(self):
        dlg = D.BatchDialog(self)
        if not dlg.exec():
            return
        job = dlg.values()
        if job["op"] == "print":  # 윈도우 「인쇄」 동작에 맡기면 PDF 기본 앱에 따라 늘 실패해서 앱이 직접 보낸다
            if not job["inputs"]:
                self.say("인쇄할 파일을 고르세요"); return
            self._batch_print(job["inputs"], job["options"].get("printer") or QPrinterInfo.defaultPrinterName()); return
        if not job["inputs"] or not job["out_dir"]:
            self.say("파일과 저장 폴더를 고르세요"); return
        T = _engine("tools")
        res = D.ProgressDialog(self, "일괄 작업", T.batch, [job]).run()
        if res is not None:
            ok = sum(1 for r in res if r.get("ok")); bad = [r for r in res if not r.get("ok")]
            msg = f"{ok}개 끝남" + (f", {len(bad)}개 실패:\n" + "\n".join(f"{os.path.basename(r['input'])}: {r.get('error')}" for r in bad[:8]) if bad else "")
            QMessageBox.information(self, "일괄 작업", msg); self._offer_open(job["out_dir"])

    def encrypt(self, t: DocTab):
        dlg = D.EncryptDialog(self)
        if not dlg.exec():
            return
        vals = dlg.values()
        if vals is None:
            return
        if not vals["user_pw"] and not vals["owner_pw"]:
            QMessageBox.warning(self, "암호화", "암호를 넣어 주세요"); return
        out, _ = QFileDialog.getSaveFileName(self, "암호 건 PDF 저장", self._suggest(t, "_암호.pdf"), PDF_FILTER)
        if not out:
            return
        T = _engine("tools")
        T.encrypt(t.doc.fitz, out, **vals)
        self.say(f"암호를 걸어 저장했습니다: {out}")

    def decrypt(self, t: DocTab):
        # 권한 암호만 걸린 문서는 is_encrypted 가 False 다. 암호화 정보로 본다.
        if not (t.doc.fitz.metadata or {}).get("encryption") and not t.doc.password:
            self.say("이 문서에는 암호가 없습니다"); return
        out, _ = QFileDialog.getSaveFileName(self, "암호 푼 PDF 저장", self._suggest(t, "_암호해제.pdf"), PDF_FILTER)
        if out:
            c = t.doc.copy()  # 열린 문서를 encryption 없이 저장하면 그 뒤 저장이 깨진다. 사본에서 푼다
            try:
                c.save(out, encryption=fitz.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
            finally:
                c.close()
            self.say(f"암호를 풀어 저장했습니다: {out}")

    def auto_bookmark(self, t: DocTab):
        """글자 크기가 큰 줄을 제목으로 보고 북마크를 만든다."""
        d = t.doc
        sizes: dict[float, int] = {}
        heads = []
        for i in range(d.page_count):
            prot = d.fitz[i].rotation
            for b in d.text_blocks(i):
                if b.get("type") != 0: continue
                for l in b["lines"]:
                    # 화면에서 바로 선 줄만(비스듬한 워터마크·옆으로 누운 글은 제목이 아니다). → 0, ↑ 90, ← 180, ↓ 270
                    dx, dy = l["dir"]
                    if abs(dx) > 0.01 and abs(dy) > 0.01: continue
                    if int(round(math.degrees(math.atan2(-dy, dx)) / 90)) * 90 % 360 != prot: continue
                    # 수식 기호 글꼴의 사용자 영역 글자는 걷고, 글자·숫자가 남는 줄만 제목 후보로
                    txt = " ".join(re.sub(r"[-]", " ", "".join(s["text"] for s in l["spans"])).split())
                    if not txt or len(txt) > 80 or not re.search(r"[0-9A-Za-z가-힣]", txt): continue
                    sz = round(max(s["size"] for s in l["spans"]), 1)
                    sizes[sz] = sizes.get(sz, 0) + len(txt)
                    heads.append((i, sz, txt, l["bbox"][1]))
        if not sizes:
            self.say("글자가 없는 문서입니다"); return
        body = max(sizes, key=sizes.get)
        big = sorted({s for s in sizes if s >= body * 1.25}, reverse=True)[:3]
        if not big:
            self.say("제목으로 볼 만한 큰 글자가 없습니다"); return
        toc = [[big.index(sz) + 1, txt, i + 1] for i, sz, txt, _ in heads if sz in big]
        # 단계가 건너뛰지 않게
        fixed, prev = [], 0
        for lvl, txt, pg in toc:
            lvl = min(lvl, prev + 1); fixed.append([lvl, txt, pg]); prev = lvl
        d.set_flat_toc(fixed); t.bookmarks.reload(); t.show_side("bookmark") if not t.side.isVisible() else None
        self.say(f"북마크 {len(fixed)}개를 만들었습니다")

    def tile(self, horizontal: bool):
        """두 문서를 나란히: 지금 문서를 새 창에 하나 더 연다."""
        t = self.cur()
        if not t or not t.doc.path:
            self.say("저장된 문서만 나란히 볼 수 있습니다"); return
        w = MainWindow(); w.open_path(t.doc.path); w.show()
        g = self.frameGeometry(); scr = self.screen().availableGeometry()
        if horizontal:
            self.setGeometry(scr.x(), scr.y(), scr.width() // 2, scr.height()); w.setGeometry(scr.x() + scr.width() // 2, scr.y(), scr.width() // 2, scr.height())
        else:
            self.setGeometry(scr.x(), scr.y(), scr.width(), scr.height() // 2); w.setGeometry(scr.x(), scr.y() + scr.height() // 2, scr.width(), scr.height() // 2)
        self._tiled = w

    def annots_import_export(self, key: str):
        t = self.cur()
        if not t:
            return
        if key == "export":
            out, _ = QFileDialog.getSaveFileName(self, "주석 내보내기", self._suggest(t, "_주석.json"), "JSON (*.json)")
            if not out: return
            data = A.export_annots(t.doc)
            with open(out, "w", encoding="utf-8") as fp:
                json.dump(data, fp, ensure_ascii=False, indent=1)
            self.say(f"주석 {len(data)}개를 내보냈습니다")
        else:
            p, _ = QFileDialog.getOpenFileName(self, "주석 가져오기", "", "JSON (*.json)")
            if not p: return
            try:
                with open(p, encoding="utf-8") as fp:
                    data = json.load(fp)
                if not isinstance(data, list): raise ValueError("주석 목록이 아님")
            except (OSError, ValueError) as e:
                QMessageBox.warning(self, "주석 가져오기", f"읽을 수 없는 파일입니다 ({e})"); return
            n, skipped = A.import_annots(t.doc, data)
            t.changed(None); self.say(f"주석 {n}개를 가져왔습니다" + (f" · 못 가져온 것 {skipped}개" if skipped else ""))

    # ── 인쇄·메일·속성·설정·도움말 ──────────────────────────────
    def print_doc(self, t: DocTab):
        """알PDF 와 같은 인쇄 창(미리보기·크기·포스터·여러 페이지·소책자·양면·흑백). 고른 설정은 다음에도 그대로."""
        from .printing import PrintDialog, send
        if not QPrinterInfo.availablePrinterNames():
            QMessageBox.warning(self, "인쇄", "이 PC 에 설치된 프린터가 없습니다"); return
        dlg = PrintDialog(self, t.doc.fitz, t.viewer.current, self.prefs.get("print"), lock=t.doc.lock,
                          view=self._print_view(t), name=t.doc.name)
        if not dlg.exec():
            return
        self.prefs["print"] = dlg.remember(); save_settings(self.settings)
        opts = dlg.options()
        err = send(self, t.doc.fitz, opts, t.doc.name, lock=t.doc.lock)
        if err == "그만둠":
            self.say("인쇄를 그만뒀습니다")
        elif err:
            QMessageBox.warning(self, "인쇄", err)
        else:
            self.say(f"인쇄를 보냈습니다 · {opts.printer} · {len(opts.pages) if opts.view is None else 1}쪽")

    def _print_view(self, t: DocTab):
        """인쇄 창의 「최근 보기」: 뷰어 지금 쪽에서 화면에 보이는 곳(쪽 좌표). 안 보이면 None."""
        v = t.viewer
        if not v.items:
            return None
        i = max(0, min(v.current, len(v.items) - 1))
        it = v.items[i]
        pr = it.mapRectToScene(it.page_rect())
        inter = v.mapToScene(v.viewport().rect()).boundingRect().intersected(pr)
        if inter.isEmpty() or v.zoom <= 0:
            return None
        z = v.zoom
        r = fitz.Rect((inter.left() - pr.left()) / z, (inter.top() - pr.top()) / z,
                      (inter.right() - pr.left()) / z, (inter.bottom() - pr.top()) / z) & t.doc.fitz[i].rect
        return None if r.is_empty else (i, r)

    def _batch_print(self, inputs: list[str], printer: str):
        """일괄 인쇄: 파일마다 열어 인쇄 창과 같은 길(send)로 보낸다. 용지에 맞추기·쪽 모양 따라·컬러."""
        from .printing import PrintOptions, send
        ok, bad = 0, []
        for p in inputs:
            try:
                d = fitz.open(p)
                if d.needs_pass:
                    bad.append(f"{os.path.basename(p)}: 암호가 걸린 파일"); continue
                if not d.is_pdf:
                    d = fitz.open("pdf", d.convert_to_pdf())
                err = send(self, d, PrintOptions(printer=printer, pages=list(range(d.page_count))), os.path.basename(p))
                d.close()
            except Exception as e:  # noqa: BLE001
                err = f"열 수 없는 파일 ({e})"
            if err == "그만둠":
                break
            if err:
                bad.append(f"{os.path.basename(p)}: {err}")
            else:
                ok += 1
        QMessageBox.information(self, "일괄 인쇄", f"{ok}개 보냄" + (f", {len(bad)}개 못 보냄:\n" + "\n".join(bad[:8]) if bad else ""))

    def mail(self, t: DocTab):
        if not t.doc.path or t.doc.modified:
            if QMessageBox.question(self, "메일", "먼저 저장해야 보낼 수 있습니다. 저장할까요?") != QMessageBox.StandardButton.Yes or not self.save(t):
                return
        try:
            import win32com.client  # noqa: F401
            ol = win32com.client.Dispatch("Outlook.Application"); mail = ol.CreateItem(0)
            mail.Subject = t.doc.name; mail.Attachments.Add(t.doc.path); mail.Display()
        except Exception:  # noqa: BLE001
            webbrowser.open("mailto:?subject=" + urllib.parse.quote(t.doc.name))
            self.say("메일 프로그램을 열었습니다. 파일은 직접 붙이세요: " + t.doc.path)

    def properties(self, t: DocTab):
        dlg = D.PropertiesDialog(self, t.doc, t.doc.path)
        if dlg.exec():
            t.doc.set_metadata(dlg.values()); self.mark_modified()

    def preferences(self):
        dlg = D.PreferencesDialog(self, self.prefs)
        if dlg.exec():
            self.prefs.update(dlg.values()); save_settings(self.settings); self.apply_theme()
            for t in self.tabs:
                t.viewer.props["author"] = self.prefs.get("author", ""); t.viewer.props["measure_unit"] = self.prefs.get("unit", "cm")

    def show_help(self):
        QMessageBox.information(self, "도움말",
                                f"{APP_NAME} (알PDF 대체)\n\n"
                                "홈: 손·선택·확대, 모두 편집(원문 글자·그림 고치기), 텍스트 추가, OCR, 자르기, 검색\n"
                                "편집: 텍스트·이미지·링크 추가, 워터마크·배경·머리글/바닥글·베이츠 번호\n"
                                "주석: 연필·지우개·강조·밑줄·취소선·글상자·말풍선·노트·도형·스탬프·서명·첨부·측정\n"
                                "변환: HWP·Word·Excel·PowerPoint·이미지·텍스트·HTML·RTF·EPUB·PDF/A, 파일에서 PDF 생성, 일괄\n"
                                "보기: 확대 방식·보기 방식·바탕색·읽기 모드·자동 북마크·전체 화면·어두운 화면\n"
                                "구성: 썸네일 격자에서 끌어 순서 바꾸기, 회전·삭제·추출·분할·삽입·자르기·크기·바꾸기·레이블\n"
                                "도구: OCR·영역 OCR·병합·비교·압축·평면화·일괄·기울기 보정·표 인식·데이터 추출\n"
                                "양식: 칸 만들기(텍스트·체크·라디오·콤보·목록·버튼·이미지·서명·날짜)·양식 인식·데이터 내보내기/가져오기\n"
                                "보안: 암호화·암호 풀기·교정 표시/검색/적용·숨겨진 데이터 삭제·서명 이미지\n\n"
                                "단축키: Ctrl+O 열기 · Ctrl+S 저장 · Ctrl+P 인쇄 · Ctrl+F 검색 · Ctrl+Z/Y 되돌리기 · Ctrl+Tab 문서 넘기기\n"
                                "Ctrl+0 창 너비 · Ctrl+1 실제 크기 · Ctrl+2 창 크기 · Home/End 첫·마지막 쪽 · Esc 도구 끝")

    def quick_search(self, text: str):
        t = self.cur()
        if not t:
            return
        if not (t.side.isVisible() and t.side_key == "search"):
            t.show_side("search"); self.navbtns["search"].setChecked(True)
        t.search.edit.setText(text); t.search.run()

    def start_tool(self, key: str):
        """시작 화면 빠른 도구."""
        if key in ("edit", "ocr", "compress", "protect", "convert"):
            self.open_dialog()
            t = self.cur()
            if not t:
                return
            if key == "edit": self.ribbon.set_tab("edit"); self.on_toggle("edit_all", True); self.ribbon.set_toggle("edit_all", True)
            elif key == "ocr": self.ocr(t)
            elif key == "compress": self.compress(t)
            elif key == "protect": self.encrypt(t)
            elif key == "convert": self.ribbon.set_tab("convert"); self.convert(t, "docx")
        elif key == "merge": self.merge()
        elif key == "batch": self.batch_dialog()
        elif key == "create": self.create_from_files()
