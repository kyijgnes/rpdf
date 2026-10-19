# -*- coding: utf-8 -*-
"""변환 엔진: PDF → 그림·글자·HTML·Word·RTF·엑셀·PPT·EPUB·PDF/A·한글, 다른 파일 → PDF.

화면(Qt) 코드는 들여오지 않는다. 사용자에게 보일 실패는 EngineError(한국어) 로 올린다.
넘겨받은 Document 는 건드리지 않는다.
"""
from __future__ import annotations

import datetime as _dt
import html as _html
import io
import os
import re
import shutil
import subprocess
import tempfile
import threading
import uuid
import zipfile
from typing import Callable, Iterable, Sequence

import fitz

from .tools import EngineError, _pages, _say, korean_font, page_size_points

Progress = Callable[[int, int, str], None] | None

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".pnm", ".pgm", ".ppm", ".jxr", ".jpx", ".jp2"}
FITZ_DOC_EXTS = {".txt", ".xps", ".oxps", ".epub", ".mobi", ".fb2", ".cbz", ".svg"}
WORD_EXTS = {".docx", ".doc", ".rtf", ".odt", ".docm", ".dotx", ".dot"}
EXCEL_EXTS = {".xlsx", ".xls", ".xlsm", ".csv", ".ods"}
PPT_EXTS = {".pptx", ".ppt", ".pptm", ".odp"}
HWP_EXTS = {".hwp", ".hwpx"}
HTML_EXTS = {".html", ".htm", ".mht", ".mhtml"}
COM_TIMEOUT = 120


# ── 안쪽 도움 ───────────────────────────────────────────────────────────
def _ensure_dir(path: str) -> None:
    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _clean(s: str) -> str:
    """XML(Word·Excel·EPUB·HWPX)에 못 들어가는 제어 글자를 지운다. 공식 요강처럼 글자층에 \\x07·\\x08 이 섞인 PDF 가 있다."""
    return _CTRL.sub("", s)


def _page_text_dict(page: fitz.Page) -> dict:
    d = page.get_text("dict", flags=fitz.TEXT_PRESERVE_IMAGES | fitz.TEXT_PRESERVE_WHITESPACE)
    for b in d.get("blocks", []):
        for l in b.get("lines", []):
            for s in l.get("spans", []):
                s["text"] = _clean(s.get("text", ""))
    return d


def _visible_blocks(page: fitz.Page, pd: dict) -> list[dict]:
    """글·그림 덩어리를 보이는 화면 차례(위→아래, 왼→오른)로. get_text 좌표는 회전 전이라 돈 쪽에서 차례가 꼬였다."""
    rm = page.rotation_matrix

    def key(b):
        r = fitz.Rect(b["bbox"]) * rm
        return (round(r.y0 / 4), r.x0)
    return sorted(pd.get("blocks", []), key=key)


def _upright_image(page: fitz.Page, block: dict) -> tuple[bytes, str, float] | None:
    """그림 덩어리를 보이는 방향 그대로 (바이트, 확장자, 보이는 폭 pt). 쪽이 돌았거나 그림이 돌아 들어간 것
    (회전된 스캔 쪽)은 원본 그림이 누워 있으니 그 자리를 그려서 쓴다."""
    data = block.get("image")
    if not data:
        return None
    vis = fitz.Rect(block["bbox"]) * page.rotation_matrix
    vis.normalize()
    t = block.get("transform")
    m = (fitz.Matrix(t) if t else fitz.Matrix(1, 1)) * page.rotation_matrix
    if abs(m.b) < 1e-3 and abs(m.c) < 1e-3 and m.a > 0 and m.d > 0:
        return data, block.get("ext") or "png", vis.width
    clip = vis & page.rect
    if clip.is_empty:
        return None
    px_w = float(block.get("width") or 0)
    src_w = max(abs(m.a), abs(m.b), abs(m.c), abs(m.d)) or clip.width  # 그림 한 변이 놓인 길이(pt)
    dpi = int(max(96, min(300, 72 * max(px_w, float(block.get("height") or 0)) / max(1.0, src_w)))) if px_w else 150
    pix = page.get_pixmap(dpi=dpi, clip=clip, alpha=False)
    return pix.tobytes("png"), "png", clip.width


def _com_init():
    try:
        import pythoncom
        pythoncom.CoInitialize()
    except Exception:
        pass


def _com_uninit():
    try:
        import pythoncom
        pythoncom.CoUninitialize()
    except Exception:
        pass


def _procs(name: str) -> set[int]:
    """이름이 같은 프로세스 PID 집합(psutil 이 없으면 tasklist)."""
    try:
        import psutil
        return {p.pid for p in psutil.process_iter(["name"]) if (p.info["name"] or "").lower() == name.lower()}
    except Exception:
        pass
    try:
        out = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {name}", "/FO", "CSV", "/NH"],
                             capture_output=True, text=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)).stdout
        pids = set()
        for line in out.splitlines():
            parts = [c.strip('"') for c in line.split('","')]
            if len(parts) >= 2 and parts[1].isdigit():
                pids.add(int(parts[1]))
        return pids
    except Exception:
        return set()


def _kill(pids: Iterable[int]) -> None:
    for pid in pids:
        try:
            subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except Exception:
            pass


def _run_with_watchdog(fn: Callable[[], object], exe_name: str, timeout: int = COM_TIMEOUT, what: str = "변환"):
    """COM 호출을 다른 스레드에서 돌리고, 시한을 넘기면 새로 뜬 프로세스를 죽인다."""
    before = _procs(exe_name)
    result: dict = {}

    def runner():
        _com_init()
        try:
            result["value"] = fn()
        except BaseException as e:  # noqa: BLE001
            result["error"] = e
        finally:
            _com_uninit()

    t = threading.Thread(target=runner, daemon=True)
    t.start()
    t.join(timeout)
    if t.is_alive():
        _kill(_procs(exe_name) - before)
        raise EngineError(f"{what}이(가) {timeout}초 안에 끝나지 않아 멈췄습니다")
    if "error" in result:
        err = result["error"]
        if isinstance(err, EngineError):
            raise err
        raise EngineError(f"{what} 실패 ({_com_message(err)})") from err
    return result.get("value")


def _com_message(e: BaseException) -> str:
    try:
        import pywintypes
        if isinstance(e, pywintypes.com_error):
            parts = e.args
            if len(parts) >= 3 and parts[2]:
                return str(parts[2][2] or parts[1])
            return str(parts[1] if len(parts) > 1 else e)
    except Exception:
        pass
    return str(e)


# ── PDF → 그림 ─────────────────────────────────────────────────────────
def to_images(doc: fitz.Document, out_dir: str, fmt: str = "png", dpi: int = 150, pages=None,
              progress: Progress = None, stem: str = "page") -> list[str]:
    """쪽마다 그림 파일 하나(<stem>_p0001.png). 만든 파일 경로 목록."""
    fmt = fmt.lower().lstrip(".")
    if fmt == "jpeg":
        fmt = "jpg"
    if fmt not in ("png", "jpg", "bmp", "tiff", "tif"):
        raise EngineError(f"지원하지 않는 그림 형식입니다: {fmt}")
    os.makedirs(out_dir, exist_ok=True)
    targets = _pages(doc, pages)
    out: list[str] = []
    for k, pno in enumerate(targets):
        _say(progress, k, len(targets), f"그림으로 뜨는 중 {pno + 1}쪽")
        page = doc[pno]
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        path = os.path.join(out_dir, f"{stem}_p{pno + 1:04d}.{fmt}")
        if fmt in ("png", "jpg"):
            pix.save(path, output="png" if fmt == "png" else "jpeg", jpg_quality=90)
        else:
            from PIL import Image
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            im.save(path, "TIFF" if fmt in ("tif", "tiff") else "BMP")
        out.append(path)
    _say(progress, len(targets), len(targets), "그림 변환 끝")
    return out


# ── PDF → 글자 · HTML ───────────────────────────────────────────────────
def to_text(doc: fitz.Document, out_path: str, pages=None, layout: bool = False) -> str:
    """글자만 뽑아 .txt 로. layout=True 면 읽는 차례로 정렬."""
    _ensure_dir(out_path)
    chunks = []
    for pno in _pages(doc, pages):
        page = doc[pno]
        chunks.append(page.get_text("text", sort=layout))
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\f".join(chunks) if len(chunks) > 1 else (chunks[0] if chunks else ""))
    return out_path


def _xhtml_body(page: fitz.Page) -> str:
    x = page.get_text("xhtml")
    m = re.search(r"<div[^>]*>(.*)</div>\s*$", x, re.S)
    body = m.group(1) if m else x
    # 한글이 &#xc9f8; 꼴 숫자 참조로 나오므로 글자로 되돌린다(마크업 글자 &lt; 등은 그대로). 제어 글자는 버린다.
    return re.sub(r"&#x([0-9a-fA-F]+);", lambda mm: _clean(chr(int(mm.group(1), 16))), body)


def to_html(doc: fitz.Document, out_path: str, pages=None) -> str:
    """쪽마다 xhtml 을 이어 붙인 HTML 한 파일(그림은 base64 로 안에)."""
    _ensure_dir(out_path)
    parts = ['<!DOCTYPE html><html lang="ko"><head><meta charset="utf-8">',
             f"<title>{_html.escape(os.path.splitext(os.path.basename(out_path))[0])}</title>",
             "<style>body{font-family:'Malgun Gothic',sans-serif;max-width:900px;margin:24px auto;padding:0 16px}"
             ".page{border-bottom:1px solid #ccc;padding:16px 0}img{max-width:100%}</style></head><body>"]
    for pno in _pages(doc, pages):
        parts.append(f'<div class="page" id="p{pno + 1}">')
        parts.append(_xhtml_body(doc[pno]))
        parts.append("</div>")
    parts.append("</body></html>")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    return out_path


# ── PDF → Word · RTF ────────────────────────────────────────────────────
def to_word(doc: fitz.Document, out_path: str, pages=None, progress: Progress = None) -> str:
    """문단·글자 크기·굵게·기울임·그림을 살려 .docx 로(판형은 첫 쪽 크기)."""
    from docx import Document as DocxDocument
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, Emu

    _ensure_dir(out_path)
    targets = _pages(doc, pages)
    dx = DocxDocument()
    section = dx.sections[0]
    if targets:
        r = doc[targets[0]].rect
        _apply_page_size(section, r, WD_ORIENT, Emu)
    style = dx.styles["Normal"]
    style.font.name = "맑은 고딕"
    try:
        from docx.oxml.ns import qn
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    except Exception:
        pass
    for k, pno in enumerate(targets):
        _say(progress, k, len(targets), f"Word 로 옮기는 중 {pno + 1}쪽")
        page = doc[pno]
        if k > 0:
            dx.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if abs(page.rect.width - doc[targets[k - 1]].rect.width) > 1 or abs(page.rect.height - doc[targets[k - 1]].rect.height) > 1:
                from docx.enum.section import WD_SECTION
                section = dx.add_section(WD_SECTION.NEW_PAGE)
                _apply_page_size(section, page.rect, WD_ORIENT, Emu)
        pd = _page_text_dict(page)
        usable_w = page.rect.width - 2 * 36
        for block in _visible_blocks(page, pd):
            if block.get("type") == 1:
                _docx_image(dx, page, block, usable_w, Pt)
                continue
            for line in block.get("lines", []):
                para = dx.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    if not text:
                        continue
                    run = para.add_run(text)
                    size = float(span.get("size", 10))
                    run.font.size = Pt(max(4, min(size, 72)))
                    flags = int(span.get("flags", 0))
                    run.font.bold = bool(flags & 16) or "bold" in span.get("font", "").lower()
                    run.font.italic = bool(flags & 2)
                    col = span.get("color")
                    if isinstance(col, int) and col != 0:
                        from docx.shared import RGBColor
                        run.font.color.rgb = RGBColor((col >> 16) & 255, (col >> 8) & 255, col & 255)
    dx.save(out_path)
    _say(progress, len(targets), len(targets), "Word 변환 끝")
    return out_path


def _apply_page_size(section, rect: fitz.Rect, WD_ORIENT, Emu) -> None:
    pt = 12700  # EMU per point
    w, h = rect.width, rect.height
    section.orientation = WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT
    section.page_width = Emu(int(w * pt))
    section.page_height = Emu(int(h * pt))
    m = Emu(int(36 * pt))
    section.left_margin = section.right_margin = m
    section.top_margin = section.bottom_margin = m


def _docx_image(dx, page: fitz.Page, block: dict, usable_w: float, Pt) -> None:
    got = _upright_image(page, block)
    if not got:
        return
    data, _ext, bw = got
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        if im.mode not in ("RGB", "L", "RGBA"):
            im = im.convert("RGB")
        if im.mode == "RGBA":
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "PNG")
        buf.seek(0)
        width_pt = max(10.0, min(bw, usable_w))
        dx.add_picture(buf, width=Pt(width_pt))
    except Exception:
        return


def to_rtf(doc: fitz.Document, out_path: str, pages=None) -> str:
    """간단한 RTF(문단 + 글자 크기, 한글은 \\uN 으로)."""
    _ensure_dir(out_path)
    out = ["{\\rtf1\\ansi\\ansicpg949\\deff0{\\fonttbl{\\f0\\fnil\\fcharset129 Malgun Gothic;}}\\f0"]
    targets = _pages(doc, pages)
    for k, pno in enumerate(targets):
        if k > 0:
            out.append("\\page")
        pd = doc[pno].get_text("dict")
        for block in pd.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                out.append("\\pard\\sa0 ")
                for span in line.get("spans", []):
                    size = int(round(float(span.get("size", 10)) * 2))
                    flags = int(span.get("flags", 0))
                    s = f"\\fs{size}"
                    if flags & 16:
                        s += "\\b"
                    if flags & 2:
                        s += "\\i"
                    out.append("{" + s + " " + _rtf_escape(span.get("text", "")) + "}")
                out.append("\\par")
    out.append("}")
    with open(out_path, "w", encoding="ascii", errors="replace") as f:
        f.write("\n".join(out))
    return out_path


def _rtf_escape(s: str) -> str:
    buf = []
    for ch in s:
        o = ord(ch)
        if ch in "\\{}":
            buf.append("\\" + ch)
        elif o < 128:
            buf.append(ch)
        else:
            if o > 32767:
                o -= 65536
            buf.append(f"\\u{o}?")
    return "".join(buf)


# ── PDF → 엑셀 ─────────────────────────────────────────────────────────
def _sheet_name(s: str) -> str:
    s = re.sub(r"[\[\]:*?/\\]", "_", s)
    return s[:31] or "Sheet"


def _append_text_row(ws, values: list[str]):
    """글자로만 넣는다. 「= 3x + 1」 같은 줄이 수식 칸이 되면 엑셀이 파일을 못 연다. 제어 글자도 걷는다."""
    ws.append([_clean(v) for v in values])
    for c in ws[ws.max_row]:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"


def to_excel(doc: fitz.Document, out_path: str, pages=None, progress: Progress = None) -> str:
    """쪽마다 표를 찾아 시트 하나씩(p1_t1), 표가 없으면 글자 줄을 시트(p1) 에."""
    from openpyxl import Workbook

    _ensure_dir(out_path)
    wb = Workbook()
    wb.remove(wb.active)
    targets = _pages(doc, pages)
    for k, pno in enumerate(targets):
        _say(progress, k, len(targets), f"표 찾는 중 {pno + 1}쪽")
        page = doc[pno]
        tables = []
        try:
            tables = list(page.find_tables().tables)
        except Exception:
            tables = []
        boxes = []
        for ti, tb in enumerate(tables):
            try:
                rows = tb.extract()
            except Exception:
                continue
            if not rows:
                continue
            ws = wb.create_sheet(_sheet_name(f"p{pno + 1}_t{ti + 1}"))
            for row in rows:
                _append_text_row(ws, [("" if c is None else str(c).replace("\n", " ")) for c in row])
            boxes.append(fitz.Rect(tb.bbox))
        # 표 밖 본문도 쓴다(전에는 표가 하나라도 있으면 쪽의 나머지 글을 통째로 버렸다)
        lines = []
        for b in page.get_text("dict", sort=True).get("blocks", []):
            for l in b.get("lines", []):
                if any(bx.intersects(fitz.Rect(l["bbox"])) for bx in boxes):
                    continue
                t = "".join(s["text"] for s in l["spans"])
                if t.strip():
                    lines.append(t)
        if lines or not boxes:
            ws = wb.create_sheet(_sheet_name(f"p{pno + 1}" if not boxes else f"p{pno + 1}_본문"))
            for line in lines:
                _append_text_row(ws, [line])
    if not wb.sheetnames:
        wb.create_sheet("빈 문서")
    wb.save(out_path)
    _say(progress, len(targets), len(targets), "엑셀 변환 끝")
    return out_path


# ── PDF → PPT ──────────────────────────────────────────────────────────
def to_pptx(doc: fitz.Document, out_path: str, pages=None, dpi: int = 150, progress: Progress = None) -> str:
    """쪽마다 슬라이드 하나(쪽 전체 그림), 글자는 슬라이드 노트에."""
    from pptx import Presentation
    from pptx.util import Emu

    _ensure_dir(out_path)
    targets = _pages(doc, pages)
    prs = Presentation()
    pt = 12700
    if targets:
        r = doc[targets[0]].rect
        prs.slide_width = Emu(int(r.width * pt))
        prs.slide_height = Emu(int(r.height * pt))
    blank = prs.slide_layouts[6]
    for k, pno in enumerate(targets):
        _say(progress, k, len(targets), f"슬라이드 만드는 중 {pno + 1}쪽")
        page = doc[pno]
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        buf = io.BytesIO(pix.tobytes("png"))
        slide = prs.slides.add_slide(blank)
        # 슬라이드 크기는 하나뿐이라 판형이 다른 쪽은 비율을 지켜 가운데에 둔다(늘이면 찌그러진다)
        sw, sh = int(prs.slide_width), int(prs.slide_height)
        k_ = min(sw / (page.rect.width * pt), sh / (page.rect.height * pt))
        w_, h_ = int(page.rect.width * pt * k_), int(page.rect.height * pt * k_)
        slide.shapes.add_picture(buf, (sw - w_) // 2, (sh - h_) // 2, width=w_, height=h_)
        text = page.get_text("text").strip()
        if text:
            slide.notes_slide.notes_text_frame.text = text[:30000]
    prs.save(out_path)
    _say(progress, len(targets), len(targets), "PPT 변환 끝")
    return out_path


# ── PDF → EPUB ─────────────────────────────────────────────────────────
def to_epub(doc: fitz.Document, out_path: str, pages=None) -> str:
    """최소 EPUB 3(쪽마다 xhtml 하나, 그림 base64 로 안에)."""
    _ensure_dir(out_path)
    targets = _pages(doc, pages)
    title = (doc.metadata or {}).get("title") or os.path.splitext(os.path.basename(out_path))[0] or "문서"
    uid = f"urn:uuid:{uuid.uuid4()}"
    now = _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    container = ('<?xml version="1.0" encoding="UTF-8"?>\n'
                 '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                 '</rootfiles></container>')
    manifest = ['<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>']
    spine = []
    nav_items = []
    pages_xhtml: list[tuple[str, str]] = []
    for pno in targets:
        name = f"page{pno + 1:04d}.xhtml"
        body = _xhtml_body(doc[pno])
        xhtml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
                 '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
                 f'<head><title>{_html.escape(title)} {pno + 1}</title></head><body>{body}</body></html>')
        pages_xhtml.append((name, xhtml))
        pid = f"p{pno + 1}"
        manifest.append(f'<item id="{pid}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="{pid}"/>')
        nav_items.append(f'<li><a href="{name}">{pno + 1}쪽</a></li>')
    nav = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
           f'<head><title>차례</title></head><body><nav epub:type="toc" id="toc"><h1>차례</h1><ol>'
           f'{"".join(nav_items)}</ol></nav></body></html>')
    opf = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid">'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           f'<dc:identifier id="uid">{uid}</dc:identifier><dc:title>{_html.escape(title)}</dc:title>'
           f'<dc:language>ko</dc:language><meta property="dcterms:modified">{now}</meta></metadata>'
           f'<manifest>{"".join(manifest)}</manifest><spine>{"".join(spine)}</spine></package>')
    with zipfile.ZipFile(out_path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        for name, xhtml in pages_xhtml:
            z.writestr(f"OEBPS/{name}", xhtml, compress_type=zipfile.ZIP_DEFLATED)
    return out_path


# ── PDF → PDF/A (근사) ──────────────────────────────────────────────────
_SRGB_ICC = r"C:\Windows\System32\spool\drivers\color\sRGB Color Space Profile.icm"


def _pdf_date(t: _dt.datetime) -> str:
    """D:20261019193000+09'00' 꼴."""
    z = t.strftime("%z") or "+0000"
    return t.strftime("D:%Y%m%d%H%M%S") + f"{z[:3]}'{z[3:]}'"


def _iso_date(s: str | None) -> str | None:
    """PDF 날짜(D:YYYYMMDDHHmmSS+HH'mm') → XMP 날짜(YYYY-MM-DDTHH:MM:SS+HH:MM). 못 읽으면 None."""
    m = re.match(r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?([Zz+-])?(\d{2})?'?(\d{2})?", s or "")
    if not m:
        return None
    y, mo, d, h, mi, se = (m.group(i) for i in range(1, 7))
    out = f"{y}-{mo or '01'}-{d or '01'}T{h or '00'}:{mi or '00'}:{se or '00'}"
    sign = m.group(7)
    if sign in ("+", "-"):
        out += f"{sign}{m.group(8) or '00'}:{m.group(9) or '00'}"
    elif sign:
        out += "Z"
    return out


def to_pdfa(doc: fitz.Document, out_path: str) -> str:
    """근사 PDF/A-2b. 글꼴 내장·투명도 제거 같은 완전한 검증은 하지 않는다.

    하는 것: 쓰레기 걷어 저장, XMP 에 pdfaid part=2 conformance=B 선언, sRGB OutputIntent(ICC 가 있으면).
    """
    _ensure_dir(out_path)
    work = fitz.open("pdf", doc.tobytes())
    try:
        # PDF/A \ub294 \ubb38\uc11c \uc815\ubcf4(Info)\uc640 XMP \uc758 \ub0a0\uc9dc\u00b7\uc81c\ubaa9\u00b7\ub9cc\ub4e0 \uc774\u00b7\uc81c\uc791 \ud504\ub85c\uadf8\ub7a8\uc774 \uc11c\ub85c \uac19\uc544\uc57c \ud55c\ub2e4.
        # \uc804\uc5d0\ub294 XMP \ub0a0\uc9dc\ub9cc \uc9c0\uae08\uc73c\ub85c \uc368\uc11c Info \uc758 \uc61b \ub0a0\uc9dc\uc640 \uc5b4\uae0b\ub0ac\ub2e4. \ub9cc\ub4e0 \ub0a0\uc740 \uc6d0\ub798 \uac83\uc744 \ub450\uace0 \uace0\uce5c \ub0a0\uc740 \uc9c0\uae08\uc73c\ub85c \ub9de\ucd98\ub2e4.
        meta = {k: v for k, v in (work.metadata or {}).items() if k not in ("format", "encryption")}
        now = _dt.datetime.now().astimezone()
        pdf_now = _pdf_date(now)
        create = meta.get("creationDate") if _iso_date(meta.get("creationDate")) else pdf_now
        meta["creationDate"], meta["modDate"] = create, pdf_now
        work.set_metadata(meta)
        e = lambda k: _html.escape(meta.get(k) or "")  # noqa: E731
        dc = f'<dc:title><rdf:Alt><rdf:li xml:lang="x-default">{e("title")}</rdf:li></rdf:Alt></dc:title>' if e("title") else ""
        dc += f'<dc:creator><rdf:Seq><rdf:li>{e("author")}</rdf:li></rdf:Seq></dc:creator>' if e("author") else ""
        dc += f'<dc:description><rdf:Alt><rdf:li xml:lang="x-default">{e("subject")}</rdf:li></rdf:Alt></dc:description>' if e("subject") else ""
        pdf_ns = (f'<pdf:Producer>{e("producer")}</pdf:Producer>' if e("producer") else "") + \
                 (f'<pdf:Keywords>{e("keywords")}</pdf:Keywords>' if e("keywords") else "")
        tool = f'<xmp:CreatorTool>{e("creator")}</xmp:CreatorTool>' if e("creator") else ""
        xmp = (
            '<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>'
            '<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
            '<rdf:Description rdf:about="" xmlns:pdfaid="http://www.aiim.org/pdfa/ns/id/">'
            '<pdfaid:part>2</pdfaid:part><pdfaid:conformance>B</pdfaid:conformance></rdf:Description>'
            f'<rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/">{dc}</rdf:Description>'
            f'<rdf:Description rdf:about="" xmlns:pdf="http://ns.adobe.com/pdf/1.3/">{pdf_ns}</rdf:Description>'
            '<rdf:Description rdf:about="" xmlns:xmp="http://ns.adobe.com/xap/1.0/">'
            f'<xmp:CreateDate>{_iso_date(create)}</xmp:CreateDate><xmp:ModifyDate>{_iso_date(pdf_now)}</xmp:ModifyDate>'
            f'<xmp:MetadataDate>{_iso_date(pdf_now)}</xmp:MetadataDate>{tool}</rdf:Description>'
            '</rdf:RDF></x:xmpmeta><?xpacket end="w"?>'
        )
        work.set_xml_metadata(xmp)
        tmp = out_path + ".tmp.pdf"
        work.save(tmp, garbage=4, deflate=True, clean=True)
    finally:
        work.close()
    try:
        _add_output_intent(tmp, out_path)
    except Exception:
        shutil.move(tmp, out_path)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
    return out_path


def _add_output_intent(src: str, dst: str) -> None:
    import pikepdf

    with pikepdf.open(src) as pdf:
        if os.path.exists(_SRGB_ICC):
            with open(_SRGB_ICC, "rb") as f:
                icc = f.read()
            stream = pikepdf.Stream(pdf, icc)
            stream.N = 3
            stream.Alternate = pikepdf.Name("/DeviceRGB")
            intent = pikepdf.Dictionary(
                Type=pikepdf.Name("/OutputIntent"), S=pikepdf.Name("/GTS_PDFA1"),
                OutputConditionIdentifier=pikepdf.String("sRGB IEC61966-2.1"),
                Info=pikepdf.String("sRGB IEC61966-2.1"), RegistryName=pikepdf.String("http://www.color.org"),
                DestOutputProfile=pdf.make_indirect(stream))
            pdf.Root.OutputIntents = pikepdf.Array([pdf.make_indirect(intent)])
        if "/MarkInfo" not in pdf.Root:
            pdf.Root.MarkInfo = pikepdf.Dictionary(Marked=True)
        pdf.save(dst, min_version="1.7", object_stream_mode=pikepdf.ObjectStreamMode.disable)


# ── PDF → 한글(.hwpx) ──────────────────────────────────────────────────
def _hwp_dispatch():
    try:
        import pywintypes
        import win32com.client
    except ImportError as e:
        raise EngineError("pywin32 가 없어 한글을 부르지 못합니다") from e
    try:
        try:
            hwp = win32com.client.gencache.EnsureDispatch("HWPFrame.HwpObject")
        except Exception:
            hwp = win32com.client.Dispatch("HWPFrame.HwpObject")
    except pywintypes.com_error as e:
        raise EngineError("이 PC 에 한컴 한글이 설치돼 있지 않습니다") from e
    try:
        hwp.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
    except Exception:
        pass
    # 이 PC 의 한글(NEO)은 보안모듈 등록이 안 먹혀 「파일 접근 허용」 창이 뜬다.
    # 안내창은 자동 응답 모드로, 보안창은 공용 워치독(있으면)으로 누른다.
    try:
        hwp.SetMessageBoxMode(0x22011)
    except Exception:
        pass
    _start_hwp_dialog_watchdog()
    try:
        hwp.XHwpWindows.Item(0).Visible = False
    except Exception:
        pass
    return hwp


# 이 파일 → engine → app → PDF 앱 → 14. 도구 → 코딩(루트)
_CODING_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".."))
_HWP_WATCHDOG = os.path.join(_CODING_ROOT, "11. 한글 변환기", "스킬", "공용", "한글_대화상자_자동응답.py")


def _start_hwp_dialog_watchdog() -> bool:
    """한글 변환기 공용 「대화상자 자동응답」 워치독을 띄운다(없으면 False)."""
    if not os.path.exists(_HWP_WATCHDOG):
        return False
    try:
        import importlib.util
        spec = importlib.util.spec_from_file_location("hwp_dialog_watchdog", _HWP_WATCHDOG)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        mod.start_watchdog()
        return True
    except Exception:
        return False


def _html_for_hwp(doc: fitz.Document, out_html: str, pages=None) -> str:
    """한글 HTML 필터용 중간 HTML(문단·글자 크기·굵게·기울임, 그림은 옆 파일로, 쪽 사이 쪽 나눔)."""
    folder = os.path.dirname(out_html)
    # 한글 HTML 필터는 <meta charset> 을 무시하고 CP949 로 읽는다. http-equiv 꼴이라야 UTF-8 로 읽는다.
    parts = ['<html><head><meta http-equiv="Content-Type" content="text/html; charset=utf-8">'
             '<style>body{font-family:"맑은 고딕"}p{margin:0}</style></head><body>']
    img_no = 0
    for k, pno in enumerate(_pages(doc, pages)):
        if k > 0:
            parts.append('<br style="page-break-before:always">')
        page = doc[pno]
        pd = _page_text_dict(page)
        for block in _visible_blocks(page, pd):
            if block.get("type") == 1:
                got = _upright_image(page, block)
                if not got:
                    continue
                data, ext, w = got
                img_no += 1
                name = f"img{img_no:04d}.{ext}"
                with open(os.path.join(folder, name), "wb") as f:
                    f.write(data)
                parts.append(f'<p><img src="{name}" width="{int(w)}"></p>')
                continue
            for line in block.get("lines", []):
                spans = []
                for span in line.get("spans", []):
                    text = _html.escape(span.get("text", ""))
                    if not text:
                        continue
                    size = max(4.0, min(float(span.get("size", 10)), 72.0))
                    flags = int(span.get("flags", 0))
                    if flags & 16:
                        text = f"<b>{text}</b>"
                    if flags & 2:
                        text = f"<i>{text}</i>"
                    spans.append(f'<span style="font-size:{size:.1f}pt">{text}</span>')
                if spans:
                    parts.append("<p>" + "".join(spans) + "</p>")
    parts.append("</body></html>")
    with open(out_html, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    return out_html


def to_hwpx(doc: fitz.Document, out_path: str, progress: Progress = None) -> str:
    """한컴 한글 COM 으로 .hwpx 저장. Word(.docx) 중간 파일을 먼저 시도하고, 이 PC 의 한글처럼
    MSWord 필터가 없으면 HTML 중간 파일로 연다(글자 크기·굵게·기울임·그림은 살고 자리 배치는 근사)."""
    _ensure_dir(out_path)
    tmp = tempfile.mkdtemp(prefix="pdfapp_hwpx_")        # 영문 임시 경로(한글 경로는 한글 COM 이 조용히 실패하기도 한다)
    try:
        docx = os.path.join(tmp, "mid.docx")
        html = os.path.join(tmp, "mid.html")
        to_word(doc, docx, progress=progress)
        _html_for_hwp(doc, html)
        tmp_out = os.path.join(tmp, "out.hwpx")
        _say(progress, 0, 1, "한글로 저장하는 중")

        def job():
            hwp = _hwp_dispatch()
            try:
                ok = False
                try:
                    ok = bool(hwp.Open(docx, "MSWord", "forceopen:true"))
                except Exception:
                    ok = False
                if not ok:
                    ok = bool(hwp.Open(html, "HTML", "forceopen:true"))
                if not ok:
                    raise EngineError("한글이 중간 파일(Word·HTML)을 열지 못했습니다")
                if not hwp.SaveAs(tmp_out, "HWPX", ""):
                    raise EngineError("한글이 .hwpx 로 저장하지 못했습니다")
            finally:
                try:
                    hwp.Clear(1)
                except Exception:
                    pass
                try:
                    hwp.Quit()
                except Exception:
                    pass
            return tmp_out

        _run_with_watchdog(job, "Hwp.exe", what="한글 변환")
        if not os.path.exists(tmp_out):
            raise EngineError("한글 파일이 만들어지지 않았습니다")
        shutil.move(tmp_out, out_path)
        _say(progress, 1, 1, "한글 변환 끝")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return out_path


# ── 다른 파일 → PDF ─────────────────────────────────────────────────────
def _office_to_pdf(path: str, out_pdf: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    src = os.path.abspath(path)
    out = os.path.abspath(out_pdf)

    def word():
        import win32com.client
        app = win32com.client.DispatchEx("Word.Application")
        app.Visible = False
        app.DisplayAlerts = 0
        try:
            d = app.Documents.Open(src, ReadOnly=True, AddToRecentFiles=False)
            d.ExportAsFixedFormat(out, 17)
            d.Close(False)
        finally:
            app.Quit()
        return out

    def excel():
        import win32com.client
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        try:
            wb = app.Workbooks.Open(src, ReadOnly=True)
            wb.ExportAsFixedFormat(0, out)
            wb.Close(False)
        finally:
            app.Quit()
        return out

    def ppt():
        import win32com.client
        app = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            pres = app.Presentations.Open(src, True, False, False)
            pres.SaveAs(out, 32)
            pres.Close()
        finally:
            app.Quit()
        return out

    def hwp():
        h = _hwp_dispatch()
        tmpd = tempfile.mkdtemp(prefix="pdfapp_hwp_")      # 영문 임시 경로로 복사해 연다
        try:
            fmt = "HWPX" if ext == ".hwpx" else "HWP"
            tsrc = os.path.join(tmpd, "in" + ext)
            shutil.copy2(src, tsrc)
            tout = os.path.join(tmpd, "out.pdf")
            if not h.Open(tsrc, fmt, "forceopen:true"):
                raise EngineError(f"한글이 파일을 열지 못했습니다: {os.path.basename(path)}")
            if not h.SaveAs(tout, "PDF", ""):
                raise EngineError("한글이 PDF 로 저장하지 못했습니다")
            shutil.move(tout, out)
        finally:
            try:
                h.Clear(1)
            except Exception:
                pass
            try:
                h.Quit()
            except Exception:
                pass
            shutil.rmtree(tmpd, ignore_errors=True)
        return out

    if ext in WORD_EXTS or ext in HTML_EXTS:
        return _run_with_watchdog(word, "WINWORD.EXE", what="Word 변환")
    if ext in EXCEL_EXTS:
        return _run_with_watchdog(excel, "EXCEL.EXE", what="Excel 변환")
    if ext in PPT_EXTS:
        return _run_with_watchdog(ppt, "POWERPNT.EXE", what="PowerPoint 변환")
    if ext in HWP_EXTS:
        return _run_with_watchdog(hwp, "Hwp.exe", what="한글 변환")
    raise EngineError(f"PDF 로 바꿀 수 없는 파일입니다: {os.path.basename(path)}")


def _image_to_pdf_bytes(path: str) -> bytes:
    try:
        img = fitz.open(path)
    except Exception:
        img = None
    if img is not None and img.page_count:
        try:
            return img.convert_to_pdf()
        except Exception:
            pass
        finally:
            img.close()
    # fitz 가 못 여는 그림(webp 등)은 Pillow 로 PNG 를 거쳐
    from PIL import Image
    try:
        im = Image.open(path)
    except Exception as e:
        raise EngineError(f"그림을 열지 못했습니다: {os.path.basename(path)}") from e
    out = fitz.open()
    try:
        frames = getattr(im, "n_frames", 1)
        for i in range(frames):
            im.seek(i)
            fr = im.convert("RGB")
            buf = io.BytesIO()
            fr.save(buf, "PNG")
            w, h = fr.size
            dpi = (im.info.get("dpi") or (72, 72))[0] or 72
            pw, ph = w * 72 / dpi, h * 72 / dpi
            page = out.new_page(width=pw, height=ph)
            page.insert_image(page.rect, stream=buf.getvalue())
        return out.tobytes()
    finally:
        out.close()


def decode_text(data: bytes) -> str:
    """글자 파일 읽기: BOM(UTF-8·UTF-16) → UTF-8 → CP949(메모장 ANSI 로 저장한 한글) 차례로.
    전에는 늘 UTF-8 로 읽어 CP949 txt 가 깨졌다."""
    if data[:3] == b"\xef\xbb\xbf":
        return data[3:].decode("utf-8", "replace")
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", "replace")
    for enc in ("utf-8", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _html_text(data: bytes) -> str:
    m = re.search(rb"""charset\s*=\s*["']?([A-Za-z0-9_-]+)""", data[:4096])
    if m:
        try:
            return data.decode(m.group(1).decode("ascii").lower().replace("ks_c_5601-1987", "cp949"))
        except (LookupError, UnicodeDecodeError):
            pass
    return decode_text(data)


def _html_to_pdf_bytes(path: str) -> bytes | None:
    try:
        with open(path, "rb") as f:
            data = f.read()
        # archive: 그 파일이 있는 폴더. 없으면 <img src="그림.png"> 같은 상대 경로 그림이 빠진다
        story = fitz.Story(html=_html_text(data), archive=os.path.dirname(os.path.abspath(path)))
        buf = io.BytesIO()
        writer = fitz.DocumentWriter(buf)
        w, h = page_size_points("A4")
        rect = fitz.Rect(36, 36, w - 36, h - 36)
        more = True
        while more:
            dev = writer.begin_page(fitz.Rect(0, 0, w, h))
            more, _ = story.place(rect)
            story.draw(dev)
            writer.end_page()
        writer.close()
        return buf.getvalue()
    except Exception:
        return None


def from_files(paths: Sequence[str], out_path: str, progress: Progress = None) -> str:
    """여러 파일(PDF·그림·문서)을 차례대로 PDF 하나로."""
    if not paths:
        raise EngineError("합칠 파일이 없습니다")
    _ensure_dir(out_path)
    tmp = tempfile.mkdtemp(prefix="pdfapp_from_")
    out = fitz.open()
    try:
        n = len(paths)
        for i, p in enumerate(paths):
            _say(progress, i, n, f"넣는 중 {os.path.basename(p)}")
            if not os.path.exists(p):
                raise EngineError(f"파일이 없습니다: {p}")
            ext = os.path.splitext(p)[1].lower()
            if ext == ".pdf":
                src = fitz.open(p)
                try:
                    if src.needs_pass:
                        raise EngineError(f"암호가 걸린 PDF 입니다: {os.path.basename(p)}")
                    out.insert_pdf(src)
                finally:
                    src.close()
            elif ext in IMAGE_EXTS:
                _insert_bytes(out, _image_to_pdf_bytes(p))
            elif ext == ".txt":  # MuPDF 는 txt 를 UTF-8 로만 읽고 한글 글꼴도 없다
                _insert_bytes(out, text_to_pdf_bytes(decode_text(_read(p))))
            elif ext in FITZ_DOC_EXTS:
                src = fitz.open(p)
                try:
                    _insert_bytes(out, src.convert_to_pdf())
                finally:
                    src.close()
            elif ext in HTML_EXTS:
                data = _html_to_pdf_bytes(p)
                if data:
                    _insert_bytes(out, data)
                else:
                    _insert_bytes(out, _read(_office_to_pdf(p, os.path.join(tmp, f"{i:03d}.pdf"))))
            elif ext in WORD_EXTS | EXCEL_EXTS | PPT_EXTS | HWP_EXTS:
                _insert_bytes(out, _read(_office_to_pdf(p, os.path.join(tmp, f"{i:03d}.pdf"))))
            else:
                raise EngineError(f"지원하지 않는 파일입니다: {os.path.basename(p)}")
        if not out.page_count:
            raise EngineError("넣을 쪽이 하나도 없습니다")
        try:
            out.subset_fonts()
        except Exception:
            pass
        out.save(out_path, garbage=3, deflate=True)
        _say(progress, n, n, "PDF 만들기 끝")
    finally:
        out.close()
        shutil.rmtree(tmp, ignore_errors=True)
    return out_path


def _read(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()


def _insert_bytes(out: fitz.Document, data: bytes) -> None:
    src = fitz.open("pdf", data)
    try:
        out.insert_pdf(src)
    finally:
        src.close()


def text_to_pdf_bytes(text: str, size="A4", font_size: float = 11) -> bytes:
    """글자를 한글 글꼴로 흘려 PDF 바이트로(클립보드 글자 등)."""
    w, h = page_size_points(size)
    out = fitz.open()
    font = korean_font()
    margin = 56
    line_h = font_size * 1.5
    lines: list[str] = []
    maxw = w - 2 * margin
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if not raw:
            lines.append("")
            continue
        cur = ""
        for ch in raw:
            if font.text_length(cur + ch, fontsize=font_size) > maxw and cur:
                lines.append(cur)
                cur = ch
            else:
                cur += ch
        lines.append(cur)
    per_page = max(1, int((h - 2 * margin) // line_h))
    for start in range(0, max(1, len(lines)), per_page):
        page = out.new_page(width=w, height=h)
        tw = fitz.TextWriter(page.rect)
        y = margin + font_size
        for line in lines[start:start + per_page]:
            if line:
                tw.append((margin, y), line, font=font, fontsize=font_size)
            y += line_h
        tw.write_text(page)
    try:
        out.subset_fonts()  # 맑은 고딕 통째(쪽마다 수 MB)가 실리지 않게
        return out.tobytes(garbage=3, deflate=True)
    finally:
        out.close()


def from_clipboard(out_path: str) -> str:
    """클립보드의 그림 또는 글자를 PDF 로."""
    from PIL import ImageGrab

    _ensure_dir(out_path)
    data = None
    try:
        grabbed = ImageGrab.grabclipboard()
    except Exception:
        grabbed = None
    if grabbed is not None and hasattr(grabbed, "size"):
        im = grabbed.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "PNG")
        pdf = fitz.open()
        page = pdf.new_page(width=im.width * 72 / 96, height=im.height * 72 / 96)
        page.insert_image(page.rect, stream=buf.getvalue())
        data = pdf.tobytes()
        pdf.close()
    elif isinstance(grabbed, list):          # 파일 목록이 복사돼 있으면 그대로 합친다
        files = [p for p in grabbed if os.path.isfile(p)]
        if files:
            return from_files(files, out_path)
    if data is None:
        text = _clipboard_text()
        if not text:
            raise EngineError("클립보드에 그림이나 글자가 없습니다")
        data = text_to_pdf_bytes(text)
    with open(out_path, "wb") as f:
        f.write(data)
    return out_path


def _clipboard_text() -> str:
    try:
        import win32clipboard
        import win32con
    except ImportError:
        return ""
    try:
        win32clipboard.OpenClipboard()
        try:
            if win32clipboard.IsClipboardFormatAvailable(win32con.CF_UNICODETEXT):
                return str(win32clipboard.GetClipboardData(win32con.CF_UNICODETEXT) or "")
            return ""
        finally:
            win32clipboard.CloseClipboard()
    except Exception:
        return ""


def blank_pdf(out_path: str, size="A4", orientation: str = "portrait", count: int = 1) -> str:
    """빈 쪽 count 장짜리 PDF."""
    _ensure_dir(out_path)
    w, h = page_size_points(size, orientation)
    d = fitz.open()
    try:
        for _ in range(max(1, int(count))):
            d.new_page(width=w, height=h)
        d.save(out_path)
    finally:
        d.close()
    return out_path
