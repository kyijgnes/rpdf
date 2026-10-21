# -*- coding: utf-8 -*-
"""PDF 도구 엔진: 압축·워터마크·머리글·보안·가리기·비교·양식·일괄.

화면(Qt) 코드는 들여오지 않는다. 사용자에게 보일 실패는 EngineError(한국어) 로 올린다.
「in place」 로 적힌 함수는 넘겨받은 Document 를 고치고 None 을 돌려준다. 그 밖의 함수는
넘겨받은 Document 를 건드리지 않는다.
"""
from __future__ import annotations

import csv
import datetime as _dt
import difflib
import json
import math
import os
import re
from typing import Callable, Iterable, Sequence

import fitz

from ..appinfo import APP_NAME
from .annots import reset_widget, save_widget

Progress = Callable[[int, int, str], None] | None


class EngineError(Exception):
    """사용자에게 보여 줄 실패(메시지는 한국어)."""


def _subset(doc: fitz.Document):
    """저장 직전에 글꼴을 쓴 글자만 남긴다. 글자 하나만 넣어도 맑은 고딕 통째(7MB 넘게)가 실린다."""
    try:
        doc.subset_fonts()
    except Exception:  # noqa: BLE001
        pass


# ── 공용 ────────────────────────────────────────────────────────────────
PAGE_SIZES: dict[str, tuple[float, float]] = {       # mm (세로 방향 기준)
    "A3": (297.0, 420.0),
    "A4": (210.0, 297.0),
    "A5": (148.0, 210.0),
    "B4": (257.0, 364.0),      # JIS B4 (한국 인쇄 관행)
    "B5": (182.0, 257.0),
    "Letter": (215.9, 279.4),
    "Legal": (215.9, 355.6),
}
MM = 72.0 / 25.4

_KOREAN_FONT_FILES = ("malgun.ttf", "NanumGothic.ttf", "gulim.ttc", "batang.ttc")
_FONT_CACHE: dict[str, str] = {}


def korean_font_path() -> str:
    """한글이 되는 트루타입 글꼴 파일 경로(맑은 고딕 우선)."""
    if "path" in _FONT_CACHE:
        return _FONT_CACHE["path"]
    fonts_dir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for name in _KOREAN_FONT_FILES:
        p = os.path.join(fonts_dir, name)
        if os.path.exists(p):
            _FONT_CACHE["path"] = p
            return p
    raise EngineError("한글 글꼴(맑은 고딕)을 찾지 못했습니다")


def korean_font() -> fitz.Font:
    """글자를 그릴 때 늘 쓰는 한글 글꼴(fitz.Font). 내장 글꼴은 한글이 없다."""
    return fitz.Font(fontfile=korean_font_path())


def _say(progress: Progress, done: int, total: int, msg: str) -> None:
    if progress:
        try:
            progress(done, total, msg)
        except Exception:
            pass


def _pages(doc: fitz.Document, pages: Iterable[int] | None) -> list[int]:
    if pages is None:
        return list(range(doc.page_count))
    out = []
    for p in pages:
        if 0 <= int(p) < doc.page_count:
            out.append(int(p))
    return out


def _color(c) -> tuple[float, float, float] | None:
    """'#rrggbb' · (r,g,b) 0~1 · (r,g,b) 0~255 → (r,g,b) 0~1."""
    if c is None:
        return None
    if isinstance(c, str):
        s = c.strip().lstrip("#")
        if len(s) != 6:
            raise EngineError(f"색 값이 잘못됐습니다: {c}")
        return tuple(int(s[i:i + 2], 16) / 255 for i in (0, 2, 4))  # type: ignore[return-value]
    t = tuple(float(v) for v in c)[:3]
    if any(v > 1 for v in t):
        t = tuple(v / 255 for v in t)
    return t  # type: ignore[return-value]


def page_size_points(size="A4", orientation: str | None = "portrait") -> tuple[float, float]:
    """용지 이름 또는 (가로mm, 세로mm) → (가로pt, 세로pt)."""
    if isinstance(size, str):
        key = size.strip()
        hit = next((k for k in PAGE_SIZES if k.lower() == key.lower()), None)
        if hit is None:
            raise EngineError(f"모르는 용지 크기입니다: {size}")
        w_mm, h_mm = PAGE_SIZES[hit]
    else:
        w_mm, h_mm = float(size[0]), float(size[1])
    w, h = w_mm * MM, h_mm * MM
    if orientation == "landscape" and w < h:
        w, h = h, w
    elif orientation == "portrait" and w > h:
        w, h = h, w
    return w, h


def _file_stem(doc: fitz.Document) -> str:
    name = getattr(doc, "name", "") or ""
    return os.path.splitext(os.path.basename(name))[0] if name else ""


def _write_text(page: fitz.Page, point, text: str, fontsize: float, color, opacity=1.0,
                rotate: float = 0, render_mode: int = 0, pivot=None) -> None:
    """한글 글꼴로 글을 쓴다. point·pivot 은 보이는 화면 좌표(page.rect 기준).

    전에는 보이는 좌표로 잰 자리를 회전 전 좌표로 그대로 써서, 회전된 쪽에서는 글이 옆으로 눕거나 쪽 밖으로 나갔다.
    rotate 는 보이는 화면에서 pivot(없으면 point)을 축으로 더 돌릴 각도(워터마크).
    """
    from .document import korean_font_name

    dm = page.derotation_matrix
    p = fitz.Point(point) * dm
    morph = (fitz.Point(pivot or point) * dm, fitz.Matrix(rotate)) if rotate else None
    page.insert_text(p, text, fontsize=fontsize, fontname=korean_font_name(page), fontfile=korean_font_path(),
                     color=color, fill_opacity=opacity, rotate=page.rotation, morph=morph, render_mode=render_mode)


def _text_width(text: str, fontsize: float) -> float:
    # 맑은 고딕을 부를 때마다 파일에서 읽으면 쪽 번호 300쪽에 몇 초가 걸린다. 한 번 읽어 둔다
    f = _FONT_CACHE.get("font")
    if f is None:
        f = _FONT_CACHE["font"] = korean_font()
    return f.text_length(text, fontsize=fontsize)


# ── 압축 ────────────────────────────────────────────────────────────────
# 알PDF 단계값(높음 72dpi·45 / 중간 120dpi·70 / 낮음 150dpi·85). 목표 해상도보다 1/4 넘게 고운 그림만 줄인다
# (조금 넘는 것까지 다시 담으면 화질만 깎이고 크기는 거의 그대로다)
_LEVELS = {
    "high": dict(dpi_threshold=90, dpi_target=72, quality=45),
    "medium": dict(dpi_threshold=150, dpi_target=120, quality=70),
    "low": dict(dpi_threshold=188, dpi_target=150, quality=85),
}
COMPRESS_LEVELS = ("high", "medium", "low")   # 화면 차례(기본은 첫째 「높음」)


class NotSmaller(EngineError):
    """압축해도 원본보다 작아지지 않았다(결과 파일은 만들지 않음)."""


def _keep_copy(doc: fitz.Document, password: str | None = None) -> fitz.Document:
    """암호·권한을 그대로 둔 작업 사본. 여기서 고치고 encryption=KEEP 으로 저장하면 결과에도 같은 암호가 남는다.
    (그냥 tobytes 로 뜬 사본은 암호가 풀려, 압축·평면화·일괄 결과가 암호 없이 나갔다)"""
    from .document import DocumentError, keep_copy
    try:
        return keep_copy(doc, password)
    except DocumentError as e:
        raise EngineError(str(e)) from e


def compress(doc: fitz.Document, out_path: str, level: str = "high",
             progress: Progress = None, password: str | None = None, remove_metadata: bool = True,
             before: int | None = None) -> dict:
    """그림을 다시 눌러 담고 쓰레기 객체를 걷어 저장. {"before": 바이트, "after": 바이트}. 암호는 그대로.

    흑백(1비트) 그림도 같은 해상도로 줄인다. 원본보다 작아지지 않으면 out_path 를 만들지 않고 NotSmaller.
    before: 원본 파일 크기(사본을 넘길 때. 없으면 doc.name 파일이나 지금 바이트 길이로 잰다)."""
    if level not in _LEVELS:
        raise EngineError(f"압축 단계가 잘못됐습니다: {level}")
    opt = _LEVELS[level]
    work = _keep_copy(doc, password)
    # 다시 쓴 바이트 길이가 아니라 실제 파일 크기와 대조한다(원본보다 커진 것도 줄었다고 보였다)
    # 사본을 encryption 없이 tobytes 하면 사본의 암호가 떨어져 KEEP 저장이 깨진다(길이만 잴 때도 KEEP)
    if before is None:
        before = os.path.getsize(doc.name) if doc.name and os.path.exists(doc.name) else len(doc.tobytes(encryption=fitz.PDF_ENCRYPT_KEEP))
    tmp = out_path + ".rpdf-tmp"
    try:
        _say(progress, 0, 2, "그림 다시 담는 중")
        if hasattr(work, "rewrite_images"):
            try:
                work.rewrite_images(dpi_threshold=opt["dpi_threshold"], dpi_target=opt["dpi_target"],
                                    quality=opt["quality"], lossy=True, lossless=True, bitonal=True)
            except Exception:
                _manual_recompress(work, opt, progress)
        else:
            _manual_recompress(work, opt, progress)
        if remove_metadata:
            work.set_metadata({})
            try:
                work.del_xml_metadata()
            except Exception:  # noqa: BLE001
                pass
        _say(progress, 1, 2, "저장하는 중")
        _subset(work)
        work.save(tmp, garbage=4, deflate=True, clean=True, deflate_images=True, deflate_fonts=True,
                  encryption=fitz.PDF_ENCRYPT_KEEP)
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise
    finally:
        work.close()
    after = os.path.getsize(tmp)
    if after >= before:  # 커진 파일을 「압축」 이라고 남기지 않는다
        os.remove(tmp)
        raise NotSmaller(f"더 줄일 수 없는 PDF 입니다(원본 {before / 1024:.0f} KB, 압축해 보니 {after / 1024:.0f} KB)")
    os.replace(tmp, out_path)
    _say(progress, 2, 2, "압축 끝")
    return {"before": before, "after": after}


def _manual_recompress(work: fitz.Document, opt: dict, progress: Progress) -> None:
    """rewrite_images 가 없을 때: 그림마다 Pillow 로 다시 눌러 담는다."""
    import io

    from PIL import Image

    done_xrefs: set[int] = set()
    n = work.page_count
    for i, page in enumerate(work):
        _say(progress, i, n, f"그림 다시 담는 중 {i + 1}/{n}쪽")
        for info in page.get_images(full=True):
            xref = info[0]
            if xref in done_xrefs:
                continue
            done_xrefs.add(xref)
            try:
                raw = work.extract_image(xref)
                im = Image.open(io.BytesIO(raw["image"]))
                if im.mode in ("RGBA", "P", "LA"):
                    im = im.convert("RGB")
                elif im.mode not in ("RGB", "L"):
                    im = im.convert("RGB")
                rects = page.get_image_rects(xref)
                if rects:
                    r = rects[0]
                    want = max(1, int(r.width / 72 * opt["dpi_target"]))
                    if im.width > want * 1.1:
                        im = im.resize((want, max(1, int(im.height * want / im.width))), Image.LANCZOS)
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=opt["quality"], optimize=True)
                if buf.tell() < len(raw["image"]):
                    page.replace_image(xref, stream=buf.getvalue())
            except Exception:
                continue


# ── 평탄화 ──────────────────────────────────────────────────────────────
def flatten(doc: fitz.Document, out_path: str, password: str | None = None) -> str:
    """주석·입력칸을 본문 그림으로 굳혀 저장. 암호는 그대로."""
    work = _keep_copy(doc, password)
    try:
        if not hasattr(work, "bake"):
            raise EngineError("이 PyMuPDF 에는 평탄화(bake)가 없습니다")
        work.bake(annots=True, widgets=True)
        _subset(work)
        work.save(out_path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_KEEP)
    finally:
        work.close()
    return out_path


# ── 워터마크 · 배경 · 머리글/바닥글 · 베이츠 ──────────────────────────────
def _anchor(rect: fitz.Rect, w: float, h: float, position: str, margin: float = 36) -> fitz.Point:
    """글자/그림 상자의 왼쪽 위 점."""
    pos = (position or "center").lower()
    cx = rect.x0 + (rect.width - w) / 2
    cy = rect.y0 + (rect.height - h) / 2
    x = {"left": rect.x0 + margin, "right": rect.x1 - margin - w}.get(pos.split("_")[-1], cx)
    y = {"top": rect.y0 + margin, "bottom": rect.y1 - margin - h}.get(pos.split("_")[0], cy)
    if pos == "center":
        x, y = cx, cy
    return fitz.Point(x, y)


def watermark(doc: fitz.Document, text: str | None = None, image_path: str | None = None,
              opacity: float = 0.3, rotate: float = 45, font_size: float = 48,
              color=(0.5, 0.5, 0.5), position: str = "center", pages=None, scale: float = 1.0) -> None:
    """워터마크(글자 또는 그림)를 쪽 위에 얹는다. in place.

    position: center · top_left · top_right · bottom_left · bottom_right · top · bottom · left · right · tile(바둑판)
    """
    if not text and not image_path:
        raise EngineError("워터마크 글자나 그림 가운데 하나는 있어야 합니다")
    col = _color(color) or (0.5, 0.5, 0.5)
    img_bytes = None
    img_wh = None
    if image_path:
        img_bytes, img_wh = _blend_image(image_path, opacity)
    tile = (position or "").lower() == "tile"
    for pno in _pages(doc, pages):
        page = doc[pno]
        rect = page.rect
        if text:
            fs = font_size  # scale 은 그림 배율이다(대화상자가 늘 0.5 를 넘겨 48pt 글자가 22pt 로 나왔다)
            w = _text_width(text, fs)
            h = fs
            if tile:
                centers = _tile_centers(rect, w, h, rotate, gap=fs * 1.2)
            else:
                tl = _anchor(rect, w, h, position)
                centers = [fitz.Point(tl.x + w / 2, tl.y + h / 2)]
            _write_texts(page, [(fitz.Point(c.x - w / 2, c.y - h / 2 + fs * 0.8), c) for c in centers],
                         text, fs, col, opacity=opacity, rotate=rotate)
        if img_bytes:
            iw, ih = img_wh
            tw_ = min(rect.width * 0.5 * scale, iw * scale)
            th_ = tw_ * ih / iw
            if tile:
                tls = [fitz.Point(c.x - tw_ / 2, c.y - th_ / 2)
                       for c in _tile_centers(rect, tw_, th_, round(rotate / 90) * 90, gap=max(tw_, th_) * 0.4)]
            else:
                tls = [_anchor(rect, tw_, th_, position)]
            xref = 0
            for tl in tls:
                r = fitz.Rect(tl.x, tl.y, tl.x + tw_, tl.y + th_) * page.derotation_matrix  # 보이는 좌표 → 회전 전
                r.normalize()
                # 그림은 90° 단위로만 돌릴 수 있다. 가까운 쪽으로(전에는 내림이라 -45° 가 -90° 가 됐다)
                # 바둑판은 같은 그림 객체를 다시 쓴다(xref)
                xref = page.insert_image(r, stream=None if xref else img_bytes, xref=xref, overlay=True,
                                         rotate=(page.rotation + int(round(rotate / 90)) * 90) % 360)


def _tile_centers(rect: fitz.Rect, w: float, h: float, angle: float, gap: float) -> list[fitz.Point]:
    """바둑판 워터마크 자리(가운데 점). 돌린 상자 크기만큼 띄우고 줄마다 반 칸씩 엇갈린다. 가장자리는 잘려도 넣는다."""
    t = math.radians(angle)
    bw = abs(w * math.cos(t)) + abs(h * math.sin(t))
    bh = abs(w * math.sin(t)) + abs(h * math.cos(t))
    sx, sy = bw + gap, bh + gap
    out = []
    row = 0
    y = rect.y0 + sy / 2
    while y - bh / 2 < rect.y1:
        x = rect.x0 + (sx / 2 if row % 2 else 0)
        while x - bw / 2 < rect.x1:
            out.append(fitz.Point(x, y))
            x += sx
        y += sy
        row += 1
    return out


def _write_texts(page: fitz.Page, spots: list[tuple[fitz.Point, fitz.Point]], text: str, fontsize: float, color,
                 opacity=1.0, rotate: float = 0) -> None:
    """같은 글을 여러 자리에(바둑판 워터마크). spots = [(글 시작점, 돌릴 축)], 보이는 화면 좌표. 한 번에 쪽에 붙인다."""
    from .document import korean_font_name

    dm = page.derotation_matrix
    name = korean_font_name(page)
    page.insert_font(fontname=name, fontfile=korean_font_path())
    shape = page.new_shape()
    for point, pivot in spots:
        morph = (fitz.Point(pivot) * dm, fitz.Matrix(rotate)) if rotate else None
        shape.insert_text(fitz.Point(point) * dm, text, fontsize=fontsize, fontname=name, color=color,
                          fill_opacity=opacity, rotate=page.rotation, morph=morph)
    shape.commit()


def _blend_image(image_path: str, opacity: float) -> tuple[bytes, tuple[int, int]]:
    """그림을 주어진 불투명도로 미리 섞어 PNG 바이트로."""
    import io

    from PIL import Image

    try:
        im = Image.open(image_path).convert("RGBA")
    except Exception as e:
        raise EngineError(f"그림을 열지 못했습니다: {image_path}") from e
    if opacity < 1:
        a = im.getchannel("A").point(lambda v: int(v * max(0.0, min(1.0, opacity))))
        im.putalpha(a)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue(), im.size


def background(doc: fitz.Document, color=None, image_path: str | None = None,
               opacity: float = 1.0, pages=None) -> None:
    """바탕색 또는 바탕 그림을 본문 아래에 깐다. in place."""
    if color is None and not image_path:
        raise EngineError("바탕색이나 바탕 그림 가운데 하나는 있어야 합니다")
    col = _color(color)
    img_bytes = _blend_image(image_path, opacity)[0] if image_path else None
    for pno in _pages(doc, pages):
        page = doc[pno]
        full = page.rect * page.derotation_matrix  # 회전 전 좌표로 쪽 전체(page.rect 로는 90·270° 쪽의 70% 만 덮었다)
        full.normalize()
        if col is not None:
            shape = page.new_shape()
            shape.draw_rect(full)
            shape.finish(color=None, fill=col, fill_opacity=opacity)
            shape.commit(overlay=False)
        if img_bytes:
            page.insert_image(full, stream=img_bytes, overlay=False, keep_proportion=False, rotate=page.rotation)


def _fill_placeholders(s: str, page_no: int, total: int, date_s: str, fname: str) -> str:
    return (s.replace("{page}", str(page_no)).replace("{total}", str(total))
             .replace("{date}", date_s).replace("{file}", fname))


def header_footer(doc: fitz.Document, left: str = "", center: str = "", right: str = "",
                  footer_left: str = "", footer_center: str = "", footer_right: str = "",
                  font_size: float = 10, margin_mm: float = 12, pages=None, start_number: int = 1,
                  date_format: str = "%Y-%m-%d", color=(0, 0, 0), file_name: str | None = None,
                  total: int | None = None) -> None:
    """머리글·바닥글 여섯 자리. 자리표: {page} {total} {date} {file}. in place.

    {total} 은 (적용할 쪽 수 + 첫 번호 - 1), 곧 마지막 쪽에 찍히는 번호다(알PDF 의 n). total 로 따로 줄 수 있다."""
    col = _color(color) or (0, 0, 0)
    date_s = _dt.datetime.now().strftime(date_format)
    fname = file_name if file_name is not None else _file_stem(doc)
    targets = _pages(doc, pages)
    if total is None:
        total = len(targets) + start_number - 1
    m = margin_mm * MM
    for k, pno in enumerate(targets):
        page = doc[pno]
        rect = page.rect
        no = start_number + k
        slots = (
            (left, "left", "top"), (center, "center", "top"), (right, "right", "top"),
            (footer_left, "left", "bottom"), (footer_center, "center", "bottom"), (footer_right, "right", "bottom"),
        )
        for raw, hx, vy in slots:
            if not raw:
                continue
            s = _fill_placeholders(raw, no, total, date_s, fname)
            w = _text_width(s, font_size)
            x = {"left": rect.x0 + m, "center": rect.x0 + (rect.width - w) / 2, "right": rect.x1 - m - w}[hx]
            y = rect.y0 + m + font_size * 0.8 if vy == "top" else rect.y1 - m
            _write_text(page, (x, y), s, font_size, col)


def bates_numbers(doc: fitz.Document, prefix: str = "", start: int = 1, digits: int = 6,
                  suffix: str = "", position: str = "footer_right", font_size: float = 10,
                  pages=None, color=(0, 0, 0), reverse: bool = False, margin_mm: float = 12) -> None:
    """베이츠 번호(접두어 + 0 채운 번호 + 접미어). position: header_/footer_ + left/center/right. in place.

    reverse: 마지막 쪽이 start 이고 앞으로 갈수록 커진다(알PDF 「베이츠 번호 거꾸로」)."""
    col = _color(color) or (0, 0, 0)
    pos = position.lower()
    vy, _, hx = pos.partition("_")
    hx = hx or "right"
    m = margin_mm * MM
    targets = _pages(doc, pages)
    for k, pno in enumerate(targets):
        page = doc[pno]
        rect = page.rect
        no = start + (len(targets) - 1 - k if reverse else k)
        s = f"{prefix}{no:0{digits}d}{suffix}"
        w = _text_width(s, font_size)
        x = {"left": rect.x0 + m, "center": rect.x0 + (rect.width - w) / 2, "right": rect.x1 - m - w}.get(hx, rect.x1 - m - w)
        y = rect.y0 + m + font_size * 0.8 if vy == "header" else rect.y1 - m
        _write_text(page, (x, y), s, font_size, col)


# 쪽 번호 모양(화면 이름, 머리글 자리표). 쪽 번호는 머리글/바닥글 한 칸에 번호만 든 것이다(알PDF 와 같은 뜻)
PAGE_NUMBER_STYLES = [("1", "{page}"), ("1 / n", "{page} / {total}"), ("- 1 -", "- {page} -"),
                      ("페이지 1", "페이지 {page}"), ("n 페이지 중 1", "{total} 페이지 중 {page}")]
HF_SLOTS = ("left", "center", "right", "footer_left", "footer_center", "footer_right")  # header_footer 의 여섯 칸 이름


def page_numbers(doc: fitz.Document, style: str = "{page}", position: str = "footer_center", start: int = 1,
                 font_size: float = 11, margin_mm: float = 12.7, color=(0, 0, 0), pages=None,
                 total: int | None = None) -> None:
    """쪽 번호. style 은 PAGE_NUMBER_STYLES 의 자리표 글, position 은 HF_SLOTS 가운데 하나. in place."""
    if position not in HF_SLOTS:
        raise EngineError(f"쪽 번호 자리가 잘못됐습니다: {position}")
    header_footer(doc, **{position: style}, font_size=font_size, margin_mm=margin_mm, pages=pages,
                  start_number=start, color=color, total=total)


# ── 꾸미기 표시: 넣은 것을 나중에 찾아 고치고 지운다 ─────────────────────
# 꾸미기 하나를 쪽마다 폼 객체(XObject) 하나로 넣고, 그 객체에 /RPDFDeco(종류)·/RPDFSet(설정 JSON)을 단다.
# 그냥 글·그림으로 그리면 본문과 섞여 다시 찾을 수 없었다(알PDF 의 추가·업데이트·삭제를 따라 함).
DECO_KINDS = {"watermark": "워터마크", "background": "배경", "header_footer": "머리글 & 바닥글",
              "page_number": "쪽 번호", "bates": "베이츠 번호"}


def _deco_draw(kind: str, s: dict):
    """종류·설정 → (그리는 함수(임시 쪽, 차례 k, 쪽 수 n), 모든 쪽이 같은 그림인지)."""
    s = {k: v for k, v in s.items() if k != "pages"}
    if kind == "watermark":
        return (lambda d, i, k, n: watermark(d, pages=[i], **s)), True
    if kind == "background":
        return (lambda d, i, k, n: background(d, pages=[i], **s)), True
    if kind == "header_footer":
        st = int(s.pop("start_number", 1))
        return (lambda d, i, k, n: header_footer(d, pages=[i], start_number=st + k, total=n + st - 1, **s)), False
    if kind == "page_number":
        st = int(s.pop("start", 1))
        return (lambda d, i, k, n: page_numbers(d, pages=[i], start=st + k, total=n + st - 1, **s)), False
    if kind == "bates":
        st, rev = int(s.pop("start", 1)), bool(s.pop("reverse", False))
        return (lambda d, i, k, n: bates_numbers(d, pages=[i], start=st + (n - 1 - k if rev else k), **s)), False
    raise EngineError(f"모르는 꾸미기입니다: {kind}")


def add_decoration(doc: fitz.Document, kind: str, pages=None, file_name: str | None = None, **settings) -> int:
    """꾸미기를 넣는다(in place). 넣은 쪽 수. settings 는 종류마다 watermark·background·header_footer·page_numbers·
    bates_numbers 의 인자 그대로이고, 문서 안에 같이 적어 두어 「업데이트」 때 대화상자에 다시 채운다."""
    if kind == "header_footer":   # 파일 이름은 넣을 때의 이름으로 굳는다(임시 문서에는 이름이 없다)
        settings["file_name"] = file_name if file_name is not None else _file_stem(doc)
    draw, same = _deco_draw(kind, settings)
    targets = _pages(doc, pages)
    n = len(targets)
    tmp = fitz.open()
    try:
        # 1) 임시 문서에 쪽마다 그린다. 쪽과 같은 크기·회전으로 그린 뒤 회전만 0 으로 되돌리면
        #    그 쪽에 바로 그렸을 때와 같은 회전 전 좌표의 그림이 된다
        src_of: list[int] = []
        seen: dict[tuple, int] = {}
        fonts: dict[str, int] = {}
        for k, pno in enumerate(targets):
            page = doc[pno]
            cb = page.cropbox
            key = (round(cb.width, 2), round(cb.height, 2), page.rotation)
            if same and key in seen:
                src_of.append(seen[key]); continue
            t = tmp.new_page(width=cb.width, height=cb.height)
            # 앞 쪽에 넣은 글꼴을 미리 걸어 둔다. 안 그러면 쪽마다 맑은 고딕을 새로 실어 300쪽에 10초 가까이 걸렸다
            rt, rv = tmp.xref_get_key(t.xref, "Resources")
            res = int(rv.split()[0]) if rt == "xref" else t.xref
            for name, fx in fonts.items():
                tmp.xref_set_key(res, f"Font/{name}" if rt == "xref" else f"Resources/Font/{name}", f"{fx} 0 R")
            t.set_rotation(page.rotation)
            draw(tmp, t.number, k, n)
            t.set_rotation(0)
            for f in t.get_fonts():
                fonts.setdefault(f[4], f[0])
            seen[key] = t.number; src_of.append(t.number)
        # 2) 다 그린 뒤에 붙인다(붙이는 중에 임시 문서가 바뀌면 PyMuPDF 의 옮김 표가 어긋난다)
        meta = fitz.get_pdf_str(json.dumps({k: v for k, v in settings.items()}, ensure_ascii=False))
        for pno, src in zip(targets, src_of):
            page = doc[pno]
            rot = page.rotation
            before = {x[1] for x in page.get_xobjects()}
            page.set_rotation(0)   # 돌린 쪽에 그대로 붙이면 그림이 옆으로 눕는다
            try:
                page.show_pdf_page(page.rect, tmp, src, overlay=kind != "background", keep_proportion=False)
            finally:
                page.set_rotation(rot)
            for x in page.get_xobjects():
                if x[2] == 0 and x[1] not in before:
                    doc.xref_set_key(x[0], "RPDFDeco", f"/{kind}")
                    doc.xref_set_key(x[0], "RPDFSet", meta)
    finally:
        tmp.close()
    return n


def _deco_items(doc: fitz.Document, pno: int):
    """쪽 pno 에 붙은 꾸미기: [(xref, 이름, 종류, 설정 JSON 글)]."""
    out = []
    for xref, name, ref, _bbox in doc[pno].get_xobjects():
        if ref != 0:
            continue
        t, v = doc.xref_get_key(xref, "RPDFDeco")
        if t == "name":
            out.append((xref, name, v.lstrip("/"), doc.xref_get_key(xref, "RPDFSet")[1]))
    return out


def decorations(doc: fitz.Document, kind: str | None = None) -> list[dict]:
    """문서 안의 꾸미기. 같은 설정끼리 묶어 [{"kind", "settings", "pages"}] (넣은 차례)."""
    groups: dict[tuple, dict] = {}
    for pno in range(doc.page_count):
        for _x, _n, k, raw in _deco_items(doc, pno):
            if kind and k != kind:
                continue
            g = groups.get((k, raw))
            if g is None:
                try:
                    s = json.loads(raw)
                except ValueError:
                    s = {}
                g = groups[(k, raw)] = {"kind": k, "settings": s, "pages": []}
            if pno not in g["pages"]:
                g["pages"].append(pno)
    return list(groups.values())


def remove_decorations(doc: fitz.Document, kind: str) -> int:
    """kind 꾸미기를 모두 지운다(in place). 지운 수. 그린 내용 줄(「q /이름 Do Q」)과 쪽 자원의 이름을 같이 뺀다."""
    gone = 0
    for pno in range(doc.page_count):
        items = [it for it in _deco_items(doc, pno) if it[2] == kind]
        if not items:
            continue
        page = doc[pno]
        names = [re.escape(name) for _x, name, _k, _r in items]
        pat = re.compile(rb"q\s*/(?:" + "|".join(names).encode() + rb")\s+Do\s*Q")
        for cx in page.get_contents():
            data = doc.xref_stream(cx)
            new = pat.sub(b"", data)
            if new != data:
                doc.update_stream(cx, new)
        for _x, name, _k, _r in items:
            doc.xref_set_key(page.xref, f"Resources/XObject/{name}", "null")
            gone += 1
    return gone


def replace_decorations(doc: fitz.Document, kind: str, groups: Sequence[dict]) -> None:
    """kind 꾸미기를 모두 지우고 groups([{"settings", "pages"}])로 다시 넣는다(업데이트). in place."""
    remove_decorations(doc, kind)
    for g in groups:
        add_decoration(doc, kind, pages=g["pages"], **g["settings"])


# ── 빈 쪽 삭제 · 기울기 보정 ────────────────────────────────────────────
def _white_fraction(page: fitz.Page, dpi: int = 36, threshold_gray: int = 245) -> float:
    import numpy as np

    pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
    arr = np.frombuffer(pix.samples, np.uint8)
    if arr.size == 0:
        return 1.0
    return float((arr >= threshold_gray).mean())


def _ink_pixels(page: fitz.Page, dpi: int = 50) -> int:
    """종이 바탕(중앙값)보다 눈에 띄게 어두운 점 수. 낮은 해상도로 떠서 잡티는 흐려지고, 가장자리 그림자는 빼고 센다."""
    import numpy as np

    pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
    a = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.stride)[:, :pix.width].astype(np.int16)
    if a.size == 0:
        return 0
    dark = a < int(np.median(a)) - 60
    h, w = dark.shape
    mh, mw = max(1, h * 4 // 100), max(1, w * 4 // 100)
    return int(dark[mh:h - mh, mw:w - mw].sum())


def is_blank_page(page: fitz.Page, threshold: float = 0.995) -> bool:
    """글자층·벡터 그림·주석이 없고, 그림으로 봐도 잉크가 거의 없는 쪽.

    전에는 36dpi 흰 비율 99.5% 만 봐서 벡터 그래프만 있는 쪽, 「정답 ⑤」 한 줄 스캔 쪽, 서명 한 획 쪽까지 지웠다.
    threshold 는 옛 호출과 맞추려고 남겨 둔다.
    """
    if page.get_text("text").strip():
        return False
    if page.get_drawings() or page.first_annot or page.first_widget:
        return False
    return _ink_pixels(page) < 8


def delete_blank_pages(doc: fitz.Document, threshold: float = 0.995) -> list[int]:
    """글자가 없고 거의 흰 쪽을 지운다. 지운 쪽 번호(원래 번호, 0부터) 목록. in place."""
    blanks = [i for i in range(doc.page_count) if is_blank_page(doc[i], threshold)]
    if blanks and len(blanks) < doc.page_count:
        doc.delete_pages(blanks)
    elif blanks:
        # 전부 빈 쪽이면 첫 쪽은 남긴다(쪽 없는 PDF 는 못 만든다)
        doc.delete_pages(blanks[1:])
        blanks = blanks[1:]
    return blanks


def _estimate_skew(gray, max_angle: float = 10.0, step: float = 0.25) -> float:
    """가로 투영의 분산이 가장 큰 각도(도). 양수 = 시계 반대 방향으로 돌려야 바로 선다."""
    import numpy as np
    from PIL import Image

    im = Image.fromarray(gray)
    if im.width > 800:
        im = im.resize((800, max(1, int(im.height * 800 / im.width))))
    g = np.asarray(im, dtype=np.int16)
    # 종이 바탕(중앙값)보다 어두운 것만 잉크로 본다. 전에는 회색 195 미만을 다 잉크로 봐서
    # 색 바탕 표지 전체가 잉크가 되고 반듯한 쪽을 찾는 범위 끝(10°)까지 돌렸다.
    ink = (g < int(np.median(g)) - 60).astype(np.uint8) * 255
    base = Image.fromarray(ink)
    best, best_score = 0.0, -1.0
    angles = np.arange(-max_angle, max_angle + step / 2, step)
    for a in angles:
        rot = base.rotate(float(a), resample=Image.NEAREST, expand=False, fillcolor=0)
        prof = np.asarray(rot, dtype=np.float32).sum(axis=1)
        score = float(prof.var())
        if score > best_score:
            best_score, best = score, float(a)
    return best


def deskew(doc: fitz.Document, pages=None, max_angle: float = 10.0,
           progress: Progress = None) -> list[float]:
    """글자층이 없는 쪽(스캔본)의 기울기를 바로잡는다. 쪽마다 돌린 각도 목록(안 돌린 쪽은 0). in place.

    쪽을 그림으로 다시 뜨지 않고 쪽 내용 앞뒤에 회전 행렬(q … cm … Q)을 끼운다. 그래서 화질·크기가 그대로이고
    쪽 객체가 남아 북마크·링크·주석도 그대로다.
    """
    import numpy as np

    targets = _pages(doc, pages)
    angles: list[float] = []
    step = 0.25
    for k, pno in enumerate(targets):
        _say(progress, k, len(targets), f"기울기 재는 중 {pno + 1}쪽")
        page = doc[pno]
        if page.get_text("text").strip():
            angles.append(0.0)
            continue
        pix = page.get_pixmap(dpi=150, colorspace=fitz.csGRAY)
        gray = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.stride)[:, :pix.width].copy()
        ang = _estimate_skew(gray, max_angle, step)
        # 0.5° 미만은 눈에 안 띄고, 찾는 범위 끝에 걸린 값은 잘못 잰 것이다(반듯한 표지가 10° 돌던 일)
        if abs(ang) < 0.5 or abs(ang) >= max_angle - step:
            angles.append(0.0)
            continue
        _rotate_content(doc, page, ang)
        angles.append(ang)
    _say(progress, len(targets), len(targets), "기울기 보정 끝")
    return angles


def _rotate_content(doc: fitz.Document, page: fitz.Page, deg: float):
    """쪽 내용을 보이는 영역 가운데를 축으로 deg 만큼(화면에서 시계 반대) 돌린다."""
    import math

    cb = page.cropbox
    c = fitz.Point(cb.width / 2, cb.height / 2) * ~page.transformation_matrix  # 쪽 좌표 → PDF 좌표
    t = math.radians(deg)
    cs, sn = math.cos(t), math.sin(t)
    e = c.x - c.x * cs + c.y * sn
    f = c.y - c.x * sn - c.y * cs
    page.wrap_contents()
    head = doc.get_new_xref()
    doc.update_object(head, "<<>>")
    doc.update_stream(head, f"q {cs:.6f} {sn:.6f} {-sn:.6f} {cs:.6f} {e:.4f} {f:.4f} cm\n".encode())
    tail = doc.get_new_xref()
    doc.update_object(tail, "<<>>")
    doc.update_stream(tail, b"\nQ\n")
    xrefs = [head] + list(page.get_contents()) + [tail]
    doc.xref_set_key(page.xref, "Contents", "[" + " ".join(f"{x} 0 R" for x in xrefs) + "]")


# ── 쪽 상자 · 크기 ─────────────────────────────────────────────────────
def crop_pages(doc: fitz.Document, rect, pages=None) -> None:
    """자르기(CropBox, 포인트). in place."""
    r = fitz.Rect(rect)
    for pno in _pages(doc, pages):
        page = doc[pno]
        page.set_cropbox(r & page.mediabox if (r & page.mediabox).is_valid else r)


def set_page_boxes(doc: fitz.Document, box: str, rect, pages=None) -> None:
    """media/crop/trim/bleed/art 상자를 정한다. in place."""
    r = fitz.Rect(rect)
    setters = {
        "media": "set_mediabox", "crop": "set_cropbox", "trim": "set_trimbox",
        "bleed": "set_bleedbox", "art": "set_artbox",
    }
    name = setters.get(box.lower().replace("box", ""))
    if not name:
        raise EngineError(f"모르는 쪽 상자입니다: {box}")
    for pno in _pages(doc, pages):
        getattr(doc[pno], name)(r)


def resize_pages(doc: fitz.Document, size="A4", orientation: str | None = None,
                 scale_content: bool = True, pages=None) -> None:
    """쪽 크기를 바꾼다. scale_content 면 내용을 비율 맞춰 채우고, 아니면 가운데 그대로 둔다. in place."""
    targets = _pages(doc, pages)
    from .document import plain_bytes
    src = fitz.open("pdf", plain_bytes(doc))
    try:
        for pno in targets:
            old = doc[pno]
            orient = orientation
            if orient is None:
                orient = "landscape" if old.rect.width > old.rect.height else "portrait"
            w, h = page_size_points(size, orient)
            new = doc.new_page(pno=pno + 1, width=w, height=h)
            if scale_content:
                new.show_pdf_page(new.rect, src, pno, keep_proportion=True)
            else:
                ow, oh = old.rect.width, old.rect.height
                x0 = (w - ow) / 2
                y0 = (h - oh) / 2
                new.show_pdf_page(fitz.Rect(x0, y0, x0 + ow, y0 + oh), src, pno, keep_proportion=True)
            doc.delete_page(pno)
    finally:
        src.close()


def set_page_labels(doc: fitz.Document, start_page: int = 0, style: str = "D",
                    prefix: str = "", first_number: int = 1) -> None:
    """쪽 라벨(D 숫자 · r/R 로마 · a/A 알파벳). in place. 한글 접두어 때문에 pages 의 읽기·쓰기를 쓴다."""
    from .pages import read_page_labels, write_page_labels
    labels = [l for l in read_page_labels(doc) if l.get("startpage") != start_page]
    labels.append({"startpage": int(start_page), "prefix": prefix, "style": style, "firstpagenum": int(first_number)})
    write_page_labels(doc, labels)


# ── 보안 ────────────────────────────────────────────────────────────────
def encrypt(doc: fitz.Document, out_path: str, user_pw: str = "", owner_pw: str = "",
            allow_print: bool = True, allow_copy: bool = True, allow_modify: bool = True,
            allow_annotate: bool = True, allow_forms: bool = True, allow_accessibility: bool = True,
            allow_assemble: bool = True, allow_print_hq: bool = True) -> str:
    """AES-256 으로 암호를 걸어 저장."""
    perm = 0
    for flag, allowed in (
        (fitz.PDF_PERM_PRINT, allow_print), (fitz.PDF_PERM_COPY, allow_copy),
        (fitz.PDF_PERM_MODIFY, allow_modify), (fitz.PDF_PERM_ANNOTATE, allow_annotate),
        (fitz.PDF_PERM_FORM, allow_forms), (fitz.PDF_PERM_ACCESSIBILITY, allow_accessibility),
        (fitz.PDF_PERM_ASSEMBLE, allow_assemble), (fitz.PDF_PERM_PRINT_HQ, allow_print_hq),
    ):
        if allowed:
            perm |= flag
    if not user_pw and not owner_pw:
        raise EngineError("사용자 암호나 소유자 암호 가운데 하나는 넣어야 합니다")
    # 권한 암호를 비우면 전에는 열기 암호가 권한 암호로도 쓰여, 열기 암호로 열면 인쇄·복사 막기가 무의미했다
    import secrets
    from .document import plain_bytes
    work = fitz.open("pdf", plain_bytes(doc))  # 열린 문서의 글꼴을 줄이면 이어서 넣는 글자가 빠진다. 사본에서.
    try:
        _subset(work)
        work.save(out_path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=owner_pw or secrets.token_urlsafe(18),
                  user_pw=user_pw, permissions=perm, garbage=3, deflate=True)
    finally:
        work.close()
    return out_path


def is_encrypted(path: str) -> bool:
    try:
        d = fitz.open(path)
    except Exception as e:
        raise EngineError(f"파일을 열지 못했습니다: {os.path.basename(path)}") from e
    try:
        return bool(d.needs_pass or d.is_encrypted)
    finally:
        d.close()


def authenticate(doc: fitz.Document, password: str) -> bool:
    """암호를 넣어 연다. 맞으면 True."""
    if not doc.needs_pass:
        return True
    return bool(doc.authenticate(password))


def decrypt(path: str, password: str, out_path: str) -> str:
    """암호를 풀어 저장."""
    d = fitz.open(path)
    try:
        if d.needs_pass and not d.authenticate(password):
            raise EngineError("암호가 맞지 않습니다")
        d.save(out_path, encryption=fitz.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
    finally:
        d.close()
    return out_path


# ── 가리기(리댁션) · 숨은 정보 ───────────────────────────────────────────
def redact_search(doc: fitz.Document, pattern: str, regex: bool = False, pages=None,
                  fill=(0, 0, 0)) -> int:
    """찾은 글자마다 가리기 주석을 붙인다(아직 적용은 안 함). 붙인 수."""
    if not pattern:
        raise EngineError("찾을 글자를 넣어 주세요")
    from .document import find_text, search_pattern
    col = _color(fill)
    count = 0
    try:
        # 정규식은 낱말 하나가 아니라 줄 글자 전체에 건다(띄어 쓴 전화번호 「010 1234 5678」 도 잡게)
        rx = search_pattern(pattern, regex=regex, case=regex)
    except re.error as e:
        raise EngineError(f"정규식이 맞지 않습니다 ({e})") from e
    for pno in _pages(doc, pages):
        page = doc[pno]
        for _text, rects in find_text(page, rx):
            for r in rects:
                page.add_redact_annot(r, fill=col)
            count += 1
    return count


def apply_redactions(doc: fitz.Document, pages=None, images: int = fitz.PDF_REDACT_IMAGE_PIXELS) -> None:
    """붙여 둔 가리기 주석을 실제로 적용. in place."""
    for pno in _pages(doc, pages):
        doc[pno].apply_redactions(images=images)


def remove_hidden_data(doc: fitz.Document, metadata: bool = True, attachments: bool = True,
                       javascript: bool = True, annotations: bool = False, bookmarks: bool = False,
                       form_fields: bool = False) -> None:
    """숨은 정보(메타데이터·첨부·스크립트 등)를 걷어낸다. in place."""
    doc.scrub(attached_files=attachments, embedded_files=attachments, javascript=javascript,
              metadata=metadata, xml_metadata=metadata, reset_fields=False, reset_responses=False,
              remove_links=False, redactions=False, hidden_text=False, clean_pages=False, thumbnails=True)
    if annotations:
        for page in doc:
            for a in list(page.annots()):
                page.delete_annot(a)
    if bookmarks:
        doc.set_toc([])
    if form_fields:
        for page in doc:
            for w in list(page.widgets()):
                page.delete_widget(w)


# ── 비교 ────────────────────────────────────────────────────────────────
def _words(page: fitz.Page) -> list[tuple]:
    return page.get_text("words", sort=True)


def _image_diff_rects(pa: fitz.Page | None, pb: fitz.Page | None, dpi: int = 50, cell: int = 16):
    """두 쪽을 그려 다른 칸(16px 격자)을 뭉쳐 포인트 네모로."""
    import numpy as np

    def render(p):
        pix = p.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        return np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width).copy()

    a = render(pa) if pa is not None else None
    b = render(pb) if pb is not None else None
    if a is None and b is None:
        return [], []
    h = max(x.shape[0] for x in (a, b) if x is not None)
    w = max(x.shape[1] for x in (a, b) if x is not None)

    def pad(x):
        out = np.full((h, w), 255, np.uint8)
        if x is not None:
            out[:x.shape[0], :x.shape[1]] = x
        return out

    a, b = pad(a), pad(b)
    diff = np.abs(a.astype(np.int16) - b.astype(np.int16)) > 40
    rows, cols = (h + cell - 1) // cell, (w + cell - 1) // cell
    grid = np.zeros((rows, cols), bool)
    for r in range(rows):
        for c in range(cols):
            if diff[r * cell:(r + 1) * cell, c * cell:(c + 1) * cell].any():
                grid[r, c] = True
    scale = 72.0 / dpi
    rects = []
    seen = np.zeros_like(grid)
    for r in range(rows):
        for c in range(cols):
            if not grid[r, c] or seen[r, c]:
                continue
            stack = [(r, c)]
            seen[r, c] = True
            r0 = r1 = r
            c0 = c1 = c
            while stack:
                y, x = stack.pop()
                r0, r1, c0, c1 = min(r0, y), max(r1, y), min(c0, x), max(c1, x)
                for ny, nx in ((y - 1, x), (y + 1, x), (y, x - 1), (y, x + 1)):
                    if 0 <= ny < rows and 0 <= nx < cols and grid[ny, nx] and not seen[ny, nx]:
                        seen[ny, nx] = True
                        stack.append((ny, nx))
            rects.append([c0 * cell * scale, r0 * cell * scale, (c1 + 1) * cell * scale, (r1 + 1) * cell * scale])
    return rects, [list(x) for x in rects]


def _page_key(page: fitz.Page):
    """쪽 짝 맞추기용 지문: 글자가 있으면 낱말 묶음, 없으면(스캔) 작은 회색 그림(평균을 뺀 값)."""
    import numpy as np

    words = [w[4] for w in page.get_text("words")]
    if words:
        return ("t", set(words))
    pix = page.get_pixmap(matrix=fitz.Matrix(96 / max(1, page.rect.width), 128 / max(1, page.rect.height)),
                          colorspace=fitz.csGRAY, alpha=False)
    a = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.stride)[:, :pix.width].astype(float)
    a = a[:128, :96]
    a = a - a.mean()
    return ("i", a)


def _page_sim(a, b) -> float:
    """0~1. 글자 쪽은 낱말 겹침, 그림 쪽은 상관계수(글자가 빽빽한 스캔 쪽끼리는 회색 평균이 다 비슷해 차이로는 못 가른다)."""
    import numpy as np

    if a[0] != b[0]:
        return 0.0
    if a[0] == "t":
        return len(a[1] & b[1]) / max(1, len(a[1] | b[1]))
    x, y = a[1], b[1]
    if x.shape != y.shape:
        h, w = min(x.shape[0], y.shape[0]), min(x.shape[1], y.shape[1])
        x, y = x[:h, :w], y[:h, :w]
    den = float(np.sqrt((x * x).sum() * (y * y).sum()))
    if den == 0:
        return 1.0 if not x.any() and not y.any() else 0.0  # 둘 다 빈 쪽이면 같다
    return max(0.0, float((x * y).sum()) / den)


_MATCH_MIN = {"t": 0.35, "i": 0.8}  # 이만큼 비슷해야 같은 쪽으로 본다(그림 쪽은 빽빽한 글자 쪽끼리도 꽤 닮는다)


def match_pages(doc_a: fitz.Document, doc_b: fitz.Document, progress: Progress = None) -> list[tuple[int | None, int | None]]:
    """두 문서의 쪽 짝(차례 유지). 끼우거나 뺀 쪽이 있어도 뒤쪽이 한 칸씩 밀려 전부 다르다고 나오지 않게.
    비슷한 정도(낱말 겹침·작은 그림)에서 문턱을 뺀 만큼을 점수로, 순서를 지키는 짝 맞추기(동적 계획법)를 한다.
    문턱을 빼야 「조금 닮은 짝 여럿」 이 「꼭 같은 짝 몇」 을 이기지 못한다. 짝 없는 쪽은 (i, None)·(None, j)."""
    na, nb = doc_a.page_count, doc_b.page_count
    ka, kb = [], []
    for i in range(na):
        _say(progress, i, na + nb, f"쪽 지문 뜨는 중 {i + 1}/{na}")
        ka.append(_page_key(doc_a[i]))
    for j in range(nb):
        _say(progress, na + j, na + nb, f"쪽 지문 뜨는 중 {j + 1}/{nb}")
        kb.append(_page_key(doc_b[j]))
    band = 40 + abs(na - nb)  # 아주 긴 문서는 대각선 둘레만
    score = [[0.0] * (nb + 1) for _ in range(na + 1)]
    move = [[0] * (nb + 1) for _ in range(na + 1)]  # 1 짝 · 2 A만 · 3 B만
    for i in range(na + 1):
        for j in range(nb + 1):
            if i == 0 and j == 0:
                continue
            best, mv = -1.0, 0
            if i and j and abs(i - j) <= band:
                gain = _page_sim(ka[i - 1], kb[j - 1]) - _MATCH_MIN[ka[i - 1][0]]
                if gain >= 0 and score[i - 1][j - 1] + gain + 1e-9 > best:
                    best, mv = score[i - 1][j - 1] + gain + 1e-9, 1
            if i and score[i - 1][j] > best:
                best, mv = score[i - 1][j], 2
            if j and score[i][j - 1] > best:
                best, mv = score[i][j - 1], 3
            score[i][j], move[i][j] = best, mv
    pairs = []
    i, j = na, nb
    while i or j:
        mv = move[i][j]
        if mv == 1:
            pairs.append((i - 1, j - 1)); i -= 1; j -= 1
        elif mv == 2:
            pairs.append((i - 1, None)); i -= 1
        else:
            pairs.append((None, j - 1)); j -= 1
    # 같은 틈(앞뒤 짝 사이)에 양쪽 다 짝 없는 쪽이 남으면 「뺀 쪽+끼운 쪽」 이 아니라 「많이 바뀐 쪽」 으로 차례대로 짝짓는다
    out, gap_a, gap_b = [], [], []

    def flush():
        n = min(len(gap_a), len(gap_b))
        out.extend(zip(gap_a[:n], gap_b[:n]))
        out.extend((x, None) for x in gap_a[n:])
        out.extend((None, y) for y in gap_b[n:])
        gap_a.clear(); gap_b.clear()
    for a, b in pairs[::-1]:
        if a is not None and b is not None:
            flush(); out.append((a, b))
        elif a is not None:
            gap_a.append(a)
        else:
            gap_b.append(b)
    flush()
    return out


def compare(doc_a: fitz.Document, doc_b: fitz.Document, progress: Progress = None) -> list[dict]:
    """쪽마다 글자 차이(difflib)와 달라진 자리(네모)를 돌려준다. 쪽 수가 달라도 되고, 끼우거나 뺀 쪽은 짝을 맞춰
    한쪽에만 있는 쪽(only_in)으로 낸다. page_a·page_b 는 각 문서의 쪽(없으면 None)."""
    pairs = match_pages(doc_a, doc_b, progress)
    n = len(pairs)
    out = []
    for k, (ia, ib) in enumerate(pairs):
        _say(progress, k, n, f"비교하는 중 {k + 1}/{n}")
        pa = doc_a[ia] if ia is not None else None
        pb = doc_b[ib] if ib is not None else None
        wa = _words(pa) if pa is not None else []
        wb = _words(pb) if pb is not None else []
        ta = [w[4] for w in wa]
        tb = [w[4] for w in wb]
        text_diff = []
        rects_a: list[list[float]] = []
        rects_b: list[list[float]] = []
        if ta or tb:
            sm = difflib.SequenceMatcher(a=ta, b=tb, autojunk=False)
            for tag, i1, i2, j1, j2 in sm.get_opcodes():
                if tag == "equal":
                    continue
                text_diff.append((tag, " ".join(ta[i1:i2]), " ".join(tb[j1:j2])))
                rects_a.extend(list(w[:4]) for w in wa[i1:i2])
                rects_b.extend(list(w[:4]) for w in wb[j1:j2])
        else:
            rects_a, rects_b = _image_diff_rects(pa, pb)
        out.append({"page": ia if ia is not None else ib, "page_a": ia, "page_b": ib, "text_diff": text_diff,
                    "changed_rects_a": rects_a, "changed_rects_b": rects_b,
                    "only_in": "a" if pb is None else "b" if pa is None else None})
    _say(progress, n, n, "비교 끝")
    return out


# ── 속성 ────────────────────────────────────────────────────────────────
def pdf_version(doc: fitz.Document) -> str:
    if hasattr(doc, "pdf_version"):
        try:
            return str(doc.pdf_version())
        except Exception:
            pass
    fmt = (doc.metadata or {}).get("format", "")
    m = re.search(r"(\d+\.\d+)", fmt or "")
    if m:
        return m.group(1)
    try:
        head = doc.tobytes(encryption=fitz.PDF_ENCRYPT_KEEP)[:12].decode("latin-1")
        m = re.search(r"%PDF-(\d+\.\d+)", head)
        return m.group(1) if m else ""
    except Exception:
        return ""


def properties(doc: fitz.Document) -> dict:
    """메타데이터 + 쪽 수 + 쪽 크기(mm) + 암호 여부 + PDF 버전 + 파일 크기."""
    meta = dict(doc.metadata or {})
    size_mm = None
    if doc.page_count:
        r = doc[0].rect
        size_mm = (round(r.width / MM, 1), round(r.height / MM, 1))
    name = getattr(doc, "name", "") or ""
    fsize = os.path.getsize(name) if name and os.path.exists(name) else None
    return {
        **meta,
        "page_count": doc.page_count,
        "file_size": fsize,
        "file_path": name,
        "page_size_mm": size_mm,
        "encrypted": bool((doc.metadata or {}).get("encryption")),  # is_encrypted 는 푼 뒤·권한 암호만이면 False
        "pdf_version": pdf_version(doc),
        "has_form": bool(doc.is_form_pdf),
        "has_toc": bool(doc.get_toc(simple=True)),
    }


_META_KEYS = ("title", "author", "subject", "keywords", "creator", "producer", "creationDate", "modDate", "trapped")


def set_properties(doc: fitz.Document, **metadata) -> None:
    """메타데이터 바꾸기(title author subject keywords creator producer creationDate modDate). in place."""
    cur = dict(doc.metadata or {})
    cur.pop("format", None)
    cur.pop("encryption", None)
    for k, v in metadata.items():
        if k not in _META_KEYS:
            raise EngineError(f"모르는 속성입니다: {k}")
        cur[k] = "" if v is None else str(v)
    doc.set_metadata(cur)


# ── 양식 ────────────────────────────────────────────────────────────────
def form_fields(doc: fitz.Document) -> list[dict]:
    out = []
    for pno, page in enumerate(doc):
        for w in page.widgets():
            out.append({
                "name": w.field_name, "type": w.field_type_string, "type_id": w.field_type,
                "value": w.field_value, "page": pno, "rect": list(w.rect),
                "choices": list(w.choice_values or []) if w.field_type in (
                    fitz.PDF_WIDGET_TYPE_COMBOBOX, fitz.PDF_WIDGET_TYPE_LISTBOX) else None,
                "xref": w.xref,
            })
    return out


def _form_values(fields: list[dict]) -> dict[str, object]:
    """칸 이름 → 값. 라디오 묶음은 동그라미마다 줄이 있으니 켜진 것의 켜짐 이름(없으면 Off) 하나로."""
    data: dict[str, object] = {}
    for fd in fields:
        if fd["type_id"] == fitz.PDF_WIDGET_TYPE_RADIOBUTTON:
            if fd["value"] not in (None, "", "Off", False) or fd["name"] not in data:
                data[fd["name"]] = fd["value"] if fd["value"] not in (None, "", False) else "Off"
        else:
            data[fd["name"]] = fd["value"]
    return data


def export_form_data(doc: fitz.Document, out_path: str) -> str:
    """양식 값을 JSON(기본) 또는 CSV(.csv) 로."""
    fields = form_fields(doc)
    data = _form_values(fields)
    if out_path.lower().endswith(".csv"):
        first = {}
        for fd in fields:
            first.setdefault(fd["name"], fd)
        with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
            wr = csv.writer(f)
            wr.writerow(["name", "type", "value", "page"])
            for name, v in data.items():
                wr.writerow([name, first[name]["type"], v, first[name]["page"] + 1])
    else:
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump({"fields": data, "detail": fields}, f, ensure_ascii=False, indent=2)
    return out_path


_ON = ("1", "true", "yes", "on", "예")


def import_form_data(doc: fitz.Document, path: str) -> int:
    """JSON({name: value} 또는 {"fields": {...}}) · CSV(name,value 열) 를 읽어 채운다. 채운 칸 이름 수. in place."""
    values: dict[str, object] = {}
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            rd = csv.DictReader(f)
            for row in rd:
                if row.get("name") is not None:
                    values[row["name"]] = row.get("value", "")
    else:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict) and isinstance(data.get("fields"), dict):
            data = data["fields"]
        if not isinstance(data, dict):
            raise EngineError("양식 자료 파일 모양이 맞지 않습니다")
        values = data
    done: set[str] = set()
    for page in doc:
        radios = []
        for w in page.widgets():
            if w.field_name not in values:
                continue
            v = values[w.field_name]
            if w.field_type == fitz.PDF_WIDGET_TYPE_RADIOBUTTON:
                radios.append(w)  # 끄기를 먼저 하고 켜기를 나중에(켤 때 같은 묶음 나머지를 끈다)
                continue
            if w.field_type == fitz.PDF_WIDGET_TYPE_CHECKBOX:
                on = (w.on_state() or "").lower()
                w.field_value = bool(v) if not isinstance(v, str) else v.strip().lower() in _ON + (on,)
            else:
                w.field_value = "" if v is None else str(v)
            save_widget(w)
            done.add(w.field_name)
        radios.sort(key=lambda w: str(values[w.field_name]) == (w.on_state() or ""))
        for w in radios:
            w.field_value = str(values[w.field_name]) == (w.on_state() or "")
            save_widget(w)
            done.add(w.field_name)
    return len(done)


def reset_form(doc: fitz.Document) -> None:
    """양식 값을 전부 비운다. in place."""
    for page in doc:
        for w in page.widgets():
            reset_widget(w)


def extract_data(docs_paths: Sequence[str], out_xlsx: str, progress: Progress = None) -> str:
    """여러 PDF 의 양식 값을 엑셀 한 장으로(PDF 하나 = 한 줄, 칸 이름 = 열)."""
    from openpyxl import Workbook

    rows = []
    names: list[str] = []
    n = len(docs_paths)
    for i, p in enumerate(docs_paths):
        _say(progress, i, n, f"양식 읽는 중 {os.path.basename(p)}")
        try:
            d = fitz.open(p)
        except Exception:
            rows.append((p, {}, "열지 못함"))
            continue
        try:
            vals = {fd["name"]: fd["value"] for fd in form_fields(d)}
        finally:
            d.close()
        for k in vals:
            if k not in names:
                names.append(k)
        rows.append((p, vals, ""))
    wb = Workbook()
    ws = wb.active
    ws.title = "양식 자료"
    ws.append(["파일", *names, "비고"])
    for p, vals, note in rows:
        ws.append([os.path.basename(p), *[_cell(vals.get(k, "")) for k in names], note])
    wb.save(out_xlsx)
    _say(progress, n, n, "엑셀 저장 끝")
    return out_xlsx


def _cell(v):
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    return v


# ── 인쇄 ────────────────────────────────────────────────────────────────
def list_printers() -> list[str]:
    try:
        import win32print
    except ImportError as e:
        raise EngineError("pywin32 가 없어 프린터 목록을 못 읽습니다") from e
    try:
        return [p[2] for p in win32print.EnumPrinters(win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS)]
    except Exception as e:
        raise EngineError(f"프린터 목록을 못 읽었습니다 ({e})") from e


def default_printer() -> str | None:
    try:
        import win32print
        return win32print.GetDefaultPrinter()
    except Exception:
        return None


def print_pdf(path: str, printer: str | None = None, pages: Iterable[int] | None = None,
              copies: int = 1) -> str:
    """윈도우 기본 PDF 처리기로 인쇄한다. pages 가 있으면 그 쪽만 임시 PDF 로 만들어 보낸다."""
    if not os.path.exists(path):
        raise EngineError(f"파일이 없습니다: {path}")
    send = path
    if pages is not None:
        import tempfile

        d = fitz.open(path)
        try:
            sel = _pages(d, pages)
            if not sel:
                raise EngineError("인쇄할 쪽이 없습니다")
            d.select(sel)
            fd, send = tempfile.mkstemp(prefix="pdfapp_print_", suffix=".pdf")
            os.close(fd)
            d.save(send, garbage=3)
        finally:
            d.close()
    try:
        import win32api
    except ImportError as e:
        raise EngineError("pywin32 가 없어 인쇄를 못 합니다") from e
    for _ in range(max(1, int(copies))):
        try:
            if printer:
                win32api.ShellExecute(0, "printto", send, f'"{printer}"', ".", 0)
            else:
                win32api.ShellExecute(0, "print", send, None, ".", 0)
        except Exception as e:
            if getattr(e, "winerror", None) == 31 or (e.args and e.args[0] == 31):
                # SE_ERR_NOASSOC: PDF 기본 앱에 「인쇄」 동작이 없다(알PDF·이 앱 둘 다). 윈도우 문구는 엉뚱하다
                raise EngineError(f"PDF 기본 앱에 인쇄 기능이 없어 보내지 못했습니다({APP_NAME}의 인쇄·일괄 인쇄를 쓰세요)") from e
            raise EngineError(f"인쇄를 보내지 못했습니다 ({e})") from e
    return send


# ── 일괄 처리 ───────────────────────────────────────────────────────────
def _out_name(inp: str, out_dir: str, suffix: str = "", ext: str | None = None) -> str:
    stem, old_ext = os.path.splitext(os.path.basename(inp))
    return os.path.join(out_dir, f"{stem}{suffix}{ext if ext is not None else old_ext}")


def batch(jobs: Sequence[dict], progress: Progress = None) -> list[dict]:
    """여러 일을 차례로. jobs = [{"op", "inputs", "out_dir", "options"}].

    op: convert(options.target = images/text/html/word/rtf/excel/pptx/epub/pdfa/hwpx, 그 밖의 kw 는 변환 함수로)
        compress(level) · watermark(text/image_path/...) · encrypt(user_pw/owner_pw/allow_*)
        ocr(lang/dpi/force) · create(여러 입력을 PDF 하나로; options.name) · print(printer/copies)
    결과: [{"input", "output", "ok", "error"}]
    """
    from . import convert as _cv
    from . import ocr as _ocr

    results: list[dict] = []
    total = sum(1 if j.get("op") == "create" else len(j.get("inputs", [])) for j in jobs) or 1
    done = 0
    for job in jobs:
        op = job.get("op")
        inputs = list(job.get("inputs", []))
        out_dir = job.get("out_dir") or (os.path.dirname(inputs[0]) if inputs else os.getcwd())
        opts = dict(job.get("options") or {})
        os.makedirs(out_dir, exist_ok=True)
        if op == "create":
            name = opts.pop("name", "합친 문서.pdf")
            out = os.path.join(out_dir, name)
            try:
                _cv.from_files(inputs, out)
                results.append({"input": inputs, "output": out, "ok": True, "error": ""})
            except Exception as e:
                results.append({"input": inputs, "output": out, "ok": False, "error": str(e)})
            done += 1
            _say(progress, done, total, name)
            continue
        for inp in inputs:
            _say(progress, done, total, os.path.basename(inp))
            rec = {"input": inp, "output": "", "ok": False, "error": ""}
            try:
                rec["output"] = _batch_one(op, inp, out_dir, dict(opts), _cv, _ocr)
                rec["ok"] = True
            except Exception as e:
                rec["error"] = str(e)
            results.append(rec)
            done += 1
    _say(progress, total, total, "일괄 처리 끝")
    return results


def _batch_one(op: str, inp: str, out_dir: str, opts: dict, _cv, _ocr) -> str:
    if op == "print":
        return print_pdf(inp, printer=opts.get("printer"), pages=opts.get("pages"), copies=opts.get("copies", 1))
    doc = fitz.open(inp)
    pw = opts.pop("password", "")
    doc._pdfapp_pw = pw
    try:
        if doc.needs_pass:
            if not doc.authenticate(pw):
                raise EngineError("암호가 걸린 파일입니다")
        if op == "convert":
            target = opts.pop("target", "text")
            table = {
                "text": (_cv.to_text, ".txt"), "html": (_cv.to_html, ".html"), "word": (_cv.to_word, ".docx"),
                "rtf": (_cv.to_rtf, ".rtf"), "excel": (_cv.to_excel, ".xlsx"), "pptx": (_cv.to_pptx, ".pptx"),
                "epub": (_cv.to_epub, ".epub"), "pdfa": (_cv.to_pdfa, ".pdf"), "hwpx": (_cv.to_hwpx, ".hwpx"),
            }
            if target == "images":
                fmt = opts.pop("fmt", "png")
                sub = os.path.join(out_dir, os.path.splitext(os.path.basename(inp))[0])
                os.makedirs(sub, exist_ok=True)
                files = _cv.to_images(doc, sub, fmt=fmt, stem=os.path.splitext(os.path.basename(inp))[0], **opts)
                return sub if files else ""
            if target not in table:
                raise EngineError(f"모르는 변환 종류입니다: {target}")
            fn, ext = table[target]
            out = _out_name(inp, out_dir, "_pdfa" if target == "pdfa" else "", ext)
            return fn(doc, out, **opts)
        # 결과에도 원본의 암호·권한을 그대로 둔다(encryption=KEEP)
        if op == "compress":
            out = _out_name(inp, out_dir, "_압축")
            compress(doc, out, level=opts.get("level", "medium"), password=pw)
            return out
        if op == "watermark":
            out = _out_name(inp, out_dir, "_워터마크")
            watermark(doc, **opts)
            _subset(doc)
            doc.save(out, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_KEEP)
            return out
        if op == "encrypt":
            out = _out_name(inp, out_dir, "_암호")
            return encrypt(doc, out, **opts)
        if op == "ocr":
            out = _out_name(inp, out_dir, "_ocr")
            _ocr.ocr_pages(doc, **opts)
            _subset(doc)
            doc.save(out, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_KEEP)
            return out
        raise EngineError(f"모르는 일괄 작업입니다: {op}")
    finally:
        doc.close()
