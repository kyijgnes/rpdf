# -*- coding: utf-8 -*-
"""엔진(convert · tools · ocr) 시험. 작은 PDF 를 그 자리에서 만들어 쓴다(바깥 파일 없음).

COM(한글·오피스)·OCR 은 PDFAPP_FULL_TESTS=1 일 때만 돈다.
"""
from __future__ import annotations

import json
import os
import sys

import fitz
import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from app.engine import convert, ocr, tools  # noqa: E402
from app.engine.tools import EngineError  # noqa: E402

FULL = os.environ.get("PDFAPP_FULL_TESTS") == "1"
TEXTS = ["첫째 쪽 Hello World 2026", "둘째 쪽 비밀번호 1234 secret", "셋째 쪽 마지막 last page"]


def make_doc(texts=TEXTS) -> fitz.Document:
    doc = fitz.open()
    font = tools.korean_font()
    for i, t in enumerate(texts):
        page = doc.new_page(width=595, height=842)
        tw = fitz.TextWriter(page.rect)
        tw.append((72, 100), t, font=font, fontsize=14)
        tw.append((72, 140), f"줄 둘 line two {i}", font=font, fontsize=11)
        tw.write_text(page)
        page.draw_rect(fitz.Rect(72, 300, 300, 400), color=(0, 0, 1), width=2)
    return doc


@pytest.fixture
def doc():
    d = make_doc()
    yield d
    d.close()


@pytest.fixture
def pdf_path(tmp_path, doc):
    p = str(tmp_path / "src.pdf")
    doc.save(p)
    return p


# ── convert ─────────────────────────────────────────────────────────────
def test_to_images(doc, tmp_path):
    files = convert.to_images(doc, str(tmp_path), fmt="png", dpi=50, stem="t")
    assert len(files) == 3 and files[0].endswith("t_p0001.png") and os.path.getsize(files[0]) > 0
    jpgs = convert.to_images(doc, str(tmp_path), fmt="jpg", dpi=50, pages=[1])
    assert len(jpgs) == 1 and jpgs[0].endswith("page_p0002.jpg")


def test_to_text(doc, tmp_path):
    out = convert.to_text(doc, str(tmp_path / "a.txt"), layout=True)
    s = open(out, encoding="utf-8").read()
    assert "첫째 쪽" in s and "last page" in s


def test_to_html(doc, tmp_path):
    out = convert.to_html(doc, str(tmp_path / "a.html"))
    s = open(out, encoding="utf-8").read()
    assert s.startswith("<!DOCTYPE html>") and "둘째 쪽" in s and 'id="p3"' in s


def test_to_word_and_rtf(doc, tmp_path):
    from docx import Document

    calls = []
    out = convert.to_word(doc, str(tmp_path / "a.docx"), progress=lambda d, t, m: calls.append((d, t)))
    dx = Document(out)
    text = "\n".join(p.text for p in dx.paragraphs)
    assert "첫째 쪽 Hello World 2026" in text and "셋째 쪽" in text
    assert calls and calls[-1] == (3, 3)
    rtf = convert.to_rtf(doc, str(tmp_path / "a.rtf"))
    s = open(rtf, encoding="ascii").read()
    assert s.startswith("{\\rtf1") and "\\u" in s and "Hello" in s


def test_to_excel(doc, tmp_path):
    from openpyxl import load_workbook

    out = convert.to_excel(doc, str(tmp_path / "a.xlsx"))
    wb = load_workbook(out)
    assert "p1" in wb.sheetnames and any("첫째" in str(c.value) for row in wb["p1"].iter_rows() for c in row)


def test_to_pptx(doc, tmp_path):
    from pptx import Presentation

    out = convert.to_pptx(doc, str(tmp_path / "a.pptx"), dpi=40)
    prs = Presentation(out)
    assert len(prs.slides) == 3
    assert "첫째 쪽" in prs.slides[0].notes_slide.notes_text_frame.text


def test_to_epub(doc, tmp_path):
    import zipfile

    out = convert.to_epub(doc, str(tmp_path / "a.epub"))
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert names[0] == "mimetype" and z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        assert "META-INF/container.xml" in names and "OEBPS/content.opf" in names and "OEBPS/nav.xhtml" in names
        assert sum(n.startswith("OEBPS/page") for n in names) == 3
        assert "둘째" in z.read("OEBPS/page0002.xhtml").decode("utf-8")


def test_to_pdfa(doc, tmp_path):
    out = convert.to_pdfa(doc, str(tmp_path / "a_pdfa.pdf"))
    d = fitz.open(out)
    try:
        assert d.page_count == 3
        assert "pdfaid:part>2" in (d.get_xml_metadata() or "")
    finally:
        d.close()


def test_from_files_and_blank(doc, tmp_path):
    p1 = str(tmp_path / "1.pdf")
    p2 = str(tmp_path / "2.pdf")
    doc.save(p1)
    doc.save(p2)
    png = str(tmp_path / "img.png")
    fitz.open(p1)[0].get_pixmap(dpi=30).save(png)
    out = convert.from_files([p1, png, p2], str(tmp_path / "merged.pdf"))
    d = fitz.open(out)
    try:
        assert d.page_count == 7
    finally:
        d.close()
    b = convert.blank_pdf(str(tmp_path / "blank.pdf"), size="A5", orientation="landscape", count=2)
    d = fitz.open(b)
    try:
        assert d.page_count == 2 and d[0].rect.width > d[0].rect.height
        assert abs(d[0].rect.width - 210 * tools.MM) < 1
    finally:
        d.close()
    with pytest.raises(EngineError):
        convert.from_files([str(tmp_path / "없음.xyz")], str(tmp_path / "x.pdf"))


def test_text_to_pdf_bytes():
    data = convert.text_to_pdf_bytes("한글 줄\n" * 80)
    d = fitz.open("pdf", data)
    try:
        assert d.page_count >= 2 and "한글 줄" in d[0].get_text()
    finally:
        d.close()


# ── tools ───────────────────────────────────────────────────────────────
def test_compress(doc, tmp_path):
    from PIL import Image
    import io

    im = Image.radial_gradient("L").resize((1600, 1600)).convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    doc[0].insert_image(fitz.Rect(72, 420, 400, 750), stream=buf.getvalue())
    doc.set_metadata({"title": "제목", "author": "누구"})
    r = tools.compress(doc, str(tmp_path / "c.pdf"), level="high")
    assert r["before"] > 0 and r["after"] > 0 and r["after"] < r["before"]
    c = fitz.open(str(tmp_path / "c.pdf"))
    assert c.page_count == 3 and not c.metadata["title"] and not c.metadata["author"]  # 메타데이터 삭제가 기본
    w = c.extract_image(c[0].get_images()[0][0])["width"]
    assert w <= 400  # 높음 = 72dpi 쪽으로(MuPDF 는 정수배로 솎아 1600 → 400)


def test_compress_bitonal_and_not_smaller(tmp_path):
    from PIL import Image
    import io

    d = fitz.open()
    p = d.new_page()
    im = Image.effect_noise((2400, 2400), 80).point(lambda v: 255 if v > 128 else 0).convert("1")  # 흑백 300dpi 넘는 스캔
    buf = io.BytesIO(); im.save(buf, "PNG")
    p.insert_image(fitz.Rect(0, 0, 400, 400), stream=buf.getvalue())
    r = tools.compress(d, str(tmp_path / "b.pdf"), level="medium")
    assert r["after"] < r["before"]
    b = fitz.open(str(tmp_path / "b.pdf"))
    assert b.extract_image(b[0].get_images()[0][0])["width"] < 2400  # 흑백 그림도 줄었다(전에는 bitonal=False 라 그대로)
    # 이미 다 줄인 문서는 파일을 만들지 않고 알린다
    small = str(tmp_path / "s.pdf")
    with pytest.raises(tools.NotSmaller):
        tools.compress(b, small, level="low", before=os.path.getsize(str(tmp_path / "b.pdf")))
    assert not os.path.exists(small) and not os.path.exists(small + ".rpdf-tmp")


def test_flatten(doc, tmp_path):
    if not hasattr(fitz.Document, "bake"):
        pytest.skip("bake 없음")
    doc[0].add_highlight_annot(fitz.Rect(72, 88, 200, 104))
    out = tools.flatten(doc, str(tmp_path / "f.pdf"))
    d = fitz.open(out)
    try:
        assert len(list(d[0].annots())) == 0
    finally:
        d.close()


def test_watermark_text_and_image(doc, tmp_path):
    tools.watermark(doc, text="대외비 CONFIDENTIAL", opacity=0.3, rotate=45, pages=[0, 2])
    assert "CONFIDENTIAL" in doc[0].get_text() and "CONFIDENTIAL" not in doc[1].get_text()
    png = str(tmp_path / "wm.png")
    doc[0].get_pixmap(dpi=20).save(png)
    n_before = len(doc[1].get_images())
    tools.watermark(doc, image_path=png, opacity=0.5, position="bottom_right", pages=[1])
    assert len(doc[1].get_images()) == n_before + 1
    with pytest.raises(EngineError):
        tools.watermark(doc)


def test_header_footer_and_bates(doc):
    tools.header_footer(doc, left="{file}", center="{date}", right="{page}/{total}",
                        footer_center="바닥 {page}", file_name="시험", start_number=1)
    t = doc[1].get_text()
    assert "2/3" in t and "바닥 2" in t and "시험" in t and "{" not in t
    tools.bates_numbers(doc, prefix="KY-", start=10, digits=4, suffix="-A")
    assert "KY-0011-A" in doc[1].get_text()


def test_background(doc):
    tools.background(doc, color="#ffeeaa", pages=[0])
    pix = doc[0].get_pixmap(dpi=20, clip=fitz.Rect(500, 800, 590, 840))
    assert pix.pixel(2, 2) == (255, 238, 170)
    assert "첫째 쪽" in doc[0].get_text()
    with pytest.raises(EngineError):
        tools.background(doc)


def test_delete_blank_pages(doc):
    doc.new_page(pno=1)
    assert doc.page_count == 4
    deleted = tools.delete_blank_pages(doc)
    assert deleted == [1] and doc.page_count == 3 and "둘째 쪽" in doc[1].get_text()


def test_deskew_skips_text_pages(doc):
    assert tools.deskew(doc, pages=[0]) == [0.0]


def test_deskew_image_page():
    from PIL import Image, ImageDraw
    import io

    im = Image.new("L", (800, 1100), 255)
    dr = ImageDraw.Draw(im)
    for y in range(150, 1000, 40):
        dr.rectangle((100, y, 700, y + 12), fill=0)
    im = im.rotate(3, expand=False, fillcolor=255)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    d = fitz.open()
    p = d.new_page(width=400, height=550)
    p.insert_image(p.rect, stream=buf.getvalue())
    angles = tools.deskew(d)
    assert len(angles) == 1 and abs(abs(angles[0]) - 3) <= 0.6
    assert d.page_count == 1 and abs(d[0].rect.width - 400) < 0.5
    d.close()


def test_crop_resize_boxes_labels(doc):
    tools.crop_pages(doc, fitz.Rect(50, 50, 500, 700), pages=[0])
    assert abs(doc[0].rect.width - 450) < 0.5
    tools.set_page_boxes(doc, "trim", fitz.Rect(60, 60, 400, 600), pages=[1])
    assert abs(doc[1].trimbox.width - 340) < 0.5
    with pytest.raises(EngineError):
        tools.set_page_boxes(doc, "foo", fitz.Rect(0, 0, 1, 1))
    tools.resize_pages(doc, size="A5", pages=[2])
    assert abs(doc[2].rect.width - 148 * tools.MM) < 0.5 and "셋째 쪽" in doc[2].get_text()
    assert doc.page_count == 3
    tools.set_page_labels(doc, 0, style="r", prefix="", first_number=1)
    assert doc[1].get_label() == "ii"


def test_encrypt_decrypt(doc, tmp_path):
    enc = str(tmp_path / "e.pdf")
    tools.encrypt(doc, enc, user_pw="user1", owner_pw="owner1", allow_copy=False)
    assert tools.is_encrypted(enc)
    d = fitz.open(enc)
    try:
        assert d.needs_pass and not tools.authenticate(d, "wrong") and tools.authenticate(d, "user1")
        assert d.page_count == 3
    finally:
        d.close()
    with pytest.raises(EngineError):
        tools.decrypt(enc, "wrong", str(tmp_path / "x.pdf"))
    dec = tools.decrypt(enc, "user1", str(tmp_path / "d.pdf"))
    assert not tools.is_encrypted(dec)
    with pytest.raises(EngineError):
        tools.encrypt(doc, str(tmp_path / "e2.pdf"))


def test_redact(doc):
    n = tools.redact_search(doc, "비밀번호")
    assert n == 1
    n2 = tools.redact_search(doc, r"\d{4}", regex=True)
    assert n2 >= 2
    tools.apply_redactions(doc)
    assert "비밀번호" not in doc[1].get_text() and "1234" not in doc[1].get_text() and "secret" in doc[1].get_text()
    with pytest.raises(EngineError):
        tools.redact_search(doc, "")


def test_remove_hidden_data(doc):
    doc.set_metadata({"title": "제목", "author": "글쓴이"})
    doc[0].add_highlight_annot(fitz.Rect(72, 88, 200, 104))
    doc.set_toc([[1, "첫째", 1]])
    tools.remove_hidden_data(doc, annotations=True, bookmarks=True)
    assert not doc.metadata.get("title") and len(list(doc[0].annots())) == 0 and doc.get_toc() == []


def test_compare(doc):
    same = fitz.open("pdf", doc.tobytes())
    res = tools.compare(doc, same)
    assert len(res) == 3 and all(r["text_diff"] == [] and r["changed_rects_a"] == [] for r in res)
    mod = make_doc(["첫째 쪽 Hello World 2026", "둘째 쪽 바뀐 글자 9999 secret"])
    res = tools.compare(doc, mod)
    assert len(res) == 3
    assert res[0]["text_diff"] == []
    assert any(op[0] == "replace" for op in res[1]["text_diff"]) and res[1]["changed_rects_a"] and res[1]["changed_rects_b"]
    assert res[2]["only_in"] == "a" and res[2]["changed_rects_a"]
    # 글자 없는 쪽: 그림 차이
    a = fitz.open()
    a.new_page().draw_rect(fitz.Rect(100, 100, 200, 200), fill=(0, 0, 0))
    b = fitz.open()
    b.new_page()
    r = tools.compare(a, b)[0]
    assert r["changed_rects_a"] and r["changed_rects_a"][0][0] < 110 and r["text_diff"] == []
    same.close(); mod.close(); a.close(); b.close()


def test_properties_roundtrip(doc, tmp_path):
    tools.set_properties(doc, title="문서 제목", author="홍길동", keywords="시험, pdf")
    p = str(tmp_path / "m.pdf")
    doc.save(p)
    d = fitz.open(p)
    try:
        info = tools.properties(d)
        assert info["title"] == "문서 제목" and info["author"] == "홍길동" and info["page_count"] == 3
        assert info["page_size_mm"] == (209.9, 297.0) and info["encrypted"] is False
        assert info["pdf_version"].startswith("1.") and info["file_size"] > 0
    finally:
        d.close()
    with pytest.raises(EngineError):
        tools.set_properties(doc, nope="x")


def test_form_fields_export_import(doc, tmp_path):
    w = fitz.Widget()
    w.field_name = "이름"
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.rect = fitz.Rect(72, 500, 300, 530)
    w.field_value = "김영"
    doc[0].add_widget(w)
    fields = tools.form_fields(doc)
    assert len(fields) == 1 and fields[0]["name"] == "이름" and fields[0]["value"] == "김영" and fields[0]["page"] == 0
    js = tools.export_form_data(doc, str(tmp_path / "f.json"))
    assert json.load(open(js, encoding="utf-8"))["fields"] == {"이름": "김영"}
    cs = tools.export_form_data(doc, str(tmp_path / "f.csv"))
    assert "이름" in open(cs, encoding="utf-8-sig").read()
    json.dump({"이름": "박영"}, open(str(tmp_path / "in.json"), "w", encoding="utf-8"), ensure_ascii=False)
    assert tools.import_form_data(doc, str(tmp_path / "in.json")) == 1
    assert tools.form_fields(doc)[0]["value"] == "박영"
    tools.reset_form(doc)
    assert not tools.form_fields(doc)[0]["value"]
    p = str(tmp_path / "form.pdf")
    doc.save(p)
    from openpyxl import load_workbook
    x = tools.extract_data([p, p], str(tmp_path / "data.xlsx"))
    ws = load_workbook(x).active
    assert ws.max_row == 3 and ws.cell(1, 2).value == "이름"


def test_batch(doc, tmp_path):
    p = str(tmp_path / "in.pdf")
    doc.save(p)
    out_dir = str(tmp_path / "out")
    res = tools.batch([
        {"op": "convert", "inputs": [p], "out_dir": out_dir, "options": {"target": "text"}},
        {"op": "compress", "inputs": [p], "out_dir": out_dir, "options": {"level": "low"}},
        {"op": "watermark", "inputs": [p], "out_dir": out_dir, "options": {"text": "시험"}},
        {"op": "encrypt", "inputs": [p], "out_dir": out_dir, "options": {"user_pw": "a"}},
        {"op": "create", "inputs": [p, p], "out_dir": out_dir, "options": {"name": "둘.pdf"}},
        {"op": "nope", "inputs": [p], "out_dir": out_dir},
    ])
    assert [r["ok"] for r in res] == [True, True, True, True, True, False]
    assert res[0]["output"].endswith("in.txt") and os.path.exists(res[3]["output"])
    assert fitz.open(res[4]["output"]).page_count == 6


def test_list_printers():
    assert isinstance(tools.list_printers(), list)


def test_page_size_points():
    assert tools.page_size_points("a4") == pytest.approx((210 * tools.MM, 297 * tools.MM))
    assert tools.page_size_points((100, 50), None)[0] == pytest.approx(100 * tools.MM)
    with pytest.raises(EngineError):
        tools.page_size_points("A9")


# ── 10-19 2차 점검에서 고친 것 ───────────────────────────────────────────
def _app_doc(texts=TEXTS):
    from app.engine.document import Document
    d = Document()
    d.fitz.close()
    d.fitz = make_doc(texts)
    return d


def test_radio_group_export_import(tmp_path):
    from app.engine import annots as A
    d = _app_doc()
    xs = [A.add_field(d, 0, fitz.Rect(50, 500 + 30 * k, 64, 514 + 30 * k), "radio") for k in range(3)]
    assert {w.field_name for w in d.fitz[0].widgets()} == {"라디오 버튼"}  # 차례로 찍으면 한 묶음
    assert [w.on_state() for w in d.fitz[0].widgets()] == ["Choice1", "Choice2", "Choice3"]
    A.set_field_value(d, 0, xs[1], True)
    assert [w.field_value for w in d.fitz[0].widgets()] == ["Off", "Choice2", "Off"]
    A.align_fields(d, [(0, x) for x in xs], "left")  # update 가 꺼진 것을 켜던 PyMuPDF 함정
    assert [w.field_value for w in d.fitz[0].widgets()] == ["Off", "Choice2", "Off"]
    out = str(tmp_path / "f.json")
    tools.export_form_data(d.fitz, out)
    assert json.load(open(out, encoding="utf-8"))["fields"] == {"라디오 버튼": "Choice2"}
    json.dump({"라디오 버튼": "Choice3"}, open(out, "w", encoding="utf-8"), ensure_ascii=False)
    tools.import_form_data(d.fitz, out)
    assert [w.field_value for w in d.fitz[0].widgets()] == ["Off", "Off", "Choice3"]
    A.reset_fields(d)
    assert [w.field_value for w in d.fitz[0].widgets()] == ["Off", "Off", "Off"]
    A.rename_field(d, 0, xs[2], "다른 묶음")
    assert [w.field_name for w in d.fitz[0].widgets()] == ["라디오 버튼", "라디오 버튼", "다른 묶음"]


def test_resize_keeps_annots_links_fields():
    from app.engine import annots as A
    from app.engine import pages as P
    d = _app_doc()
    p = d.fitz[0]
    p.add_highlight_annot(p.search_for("Hello")[0])
    p.insert_link({"kind": fitz.LINK_GOTO, "from": fitz.Rect(72, 300, 200, 320), "page": 1})
    del p
    A.add_field(d, 0, fitz.Rect(300, 500, 450, 520), "text", value="값")
    P.resize(d, [0], 842, 595, True)
    p = d.fitz[0]
    assert (p.rect.width, p.rect.height) == (842, 595)
    hl = next(fitz.Rect(a.rect) for a in p.annots() if a.type[1] == "Highlight")
    assert hl.intersects(p.search_for("Hello")[0])
    assert len(p.get_links()) == 1 and [w.field_value for w in p.widgets()] == ["값"]


def test_search_across_lines_and_words():
    d = _app_doc(["첫 줄 끝 PDF", "다음"])
    p = d.fitz[0]
    tw = fitz.TextWriter(p.rect)
    tw.append((72, 200), "끝에 PDF", font=tools.korean_font(), fontsize=12)
    tw.append((72, 216), "편집기 word, Word wordy", font=tools.korean_font(), fontsize=12)
    tw.write_text(p)
    del p
    hits = d.search("PDF 편집기")
    assert len(hits) == 1 and len(hits[0].all_rects()) == 2
    assert [h.text for h in d.search("word", whole_word=True, case=True)] == ["word"]
    assert d.replace_text(hits, "새 글") == 1 and not d.search("편집기") and len(d.search("새 글")) == 1


def test_redact_search_spaced_phone(doc):
    page = doc[0]
    tw = fitz.TextWriter(page.rect)
    tw.append((72, 500), "전화 010 1234 5678 끝", font=tools.korean_font(), fontsize=12)
    tw.write_text(page)
    n = tools.redact_search(doc, r"(?<!\d)01[016789][-. ]?\d{3,4}[-. ]?\d{4}(?!\d)", regex=True, pages=[0])
    assert n == 1


def test_korean_page_labels_and_attachments(tmp_path):
    from app.engine import pages as P
    d = _app_doc()
    P.set_labels(d, 0, style="", prefix="표지")
    P.set_labels(d, 1, style="D", prefix="부록-", first=3)
    assert P.labels(d) == ["표지", "부록-3", "부록-4"]
    f = tmp_path / "한글 첨부.txt"
    f.write_text("내용", encoding="utf-8")
    d.add_attachment(str(f), desc="설명")
    a = d.attachments()[0]
    assert (a["filename"], a["desc"]) == ("한글 첨부.txt", "설명")
    assert d.get_attachment(a["name"]) == "내용".encode("utf-8")


def test_bookmark_rename_keeps_position(tmp_path):
    from app.engine import pages as P
    d = _app_doc()
    d.fitz.set_toc([[1, "하나", 1], [1, "둘", 2, {"kind": fitz.LINK_GOTO, "to": fitz.Point(0, 500), "zoom": 0}]])
    bms = d.bookmarks()
    bms[0].title = "하나 고침"
    d.set_bookmarks(bms)
    assert d.fitz.get_toc(simple=False)[1][3]["to"].y == pytest.approx(500)
    a, b = str(tmp_path / "a.pdf"), str(tmp_path / "b.pdf")
    d.fitz.save(a); make_doc(["B 문서"]).save(b)
    out = P.merge([a, b], str(tmp_path / "m.pdf"))
    assert [t[1] for t in fitz.open(out).get_toc()] == ["하나 고침", "둘"]


def test_layer_toggle_is_view_setting(tmp_path):
    d = _app_doc()
    oc = d.fitz.add_ocg("빨강", on=True)
    d.fitz[0].draw_rect(fitz.Rect(400, 600, 500, 700), color=(1, 0, 0), fill=(1, 0, 0), oc=oc)
    red = lambda: d.fitz[0].get_pixmap(dpi=20, clip=fitz.Rect(410, 610, 490, 690)).pixel(5, 5)  # noqa: E731
    d.set_layer(d.layers()[0]["number"], False)
    assert red() == (255, 255, 255)
    d.snapshot(); d.undo()
    assert red() == (255, 255, 255)  # 되돌리기로 다시 열어도 꺼진 채
    out = str(tmp_path / "l.pdf"); d.save(out)
    assert [v["on"] for v in fitz.open(out).get_ocgs().values()] == [True]  # 파일은 그대로(저장하면 xref 가 바뀐다)


def test_encrypted_doc_survives_exports(tmp_path):
    """암호 문서를 연 채 PDF/A·압축·암호 풀기를 해도 열린 문서의 다음 저장이 멀쩡하다(전에는 암호·글자가 다 빠졌다)."""
    from app.engine.document import Document
    src = str(tmp_path / "enc.pdf")
    make_doc().save(src, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o", permissions=fitz.PDF_PERM_PRINT)
    d = Document(src, password="u")
    c = d.copy(); convert.to_pdfa(c, str(tmp_path / "a.pdf")); c.close()
    tools.compress(d.copy(), str(tmp_path / "c.pdf"), password="u")
    tools.encrypt(d.fitz, str(tmp_path / "e.pdf"), user_pw="x", owner_pw="y")
    out = str(tmp_path / "after.pdf"); d.save(out)
    for path, pw in ((out, "u"), (str(tmp_path / "c.pdf"), "u")):
        g = fitz.open(path)
        assert g.needs_pass and g.authenticate(pw) and "Hello" in g[0].get_text()


def test_compare_aligns_inserted_page():
    a = make_doc(["A 하나 one", "B 둘 two", "C 셋 three"])
    b = make_doc(["A 하나 one", "새 쪽 new page", "B 둘 two", "C 셋 three"])
    assert tools.match_pages(a, b) == [(0, 0), (None, 1), (1, 2), (2, 3)]


def test_decorations_add_find_update_remove(doc, tmp_path):
    doc[1].set_rotation(90)
    tools.add_decoration(doc, "page_number", style="{page} / {total}", start=3)
    tools.add_decoration(doc, "watermark", text="대외비")
    tools.add_decoration(doc, "bates", prefix="A-", start=1, digits=4, reverse=True, pages=[0, 2])
    assert "3 / 5" in doc[0].get_text() and "5 / 5" in doc[2].get_text()   # n = 쪽 수 + 첫 번호 - 1
    assert "A-0002" in doc[0].get_text() and "A-0001" in doc[2].get_text()   # 거꾸로
    # 돌린 쪽에서도 바닥글은 보이는 쪽 아래쪽에 선다
    r = doc[1].search_for("4 / 5")[0] * doc[1].rotation_matrix  # search_for 는 회전 전 좌표
    assert r.y0 > doc[1].rect.height * 0.8 and abs((r.x0 + r.x1) / 2 - doc[1].rect.width / 2) < 20
    p = str(tmp_path / "d.pdf"); doc.save(p)
    d = fitz.open(p)
    got = {g["kind"]: g for g in tools.decorations(d)}
    assert set(got) == {"page_number", "watermark", "bates"} and got["bates"]["pages"] == [0, 2]
    assert got["page_number"]["settings"]["start"] == 3
    g = got["page_number"]; g["settings"]["style"] = "- {page} -"
    tools.replace_decorations(d, "page_number", [g])
    assert "- 3 -" in d[0].get_text() and "3 / 5" not in d[0].get_text()
    assert tools.remove_decorations(d, "watermark") == 3
    assert "대외비" not in d[0].get_text() and "Hello" in d[0].get_text() and "A-0002" in d[0].get_text()
    assert {g["kind"] for g in tools.decorations(d)} == {"bates", "page_number"}


def test_background_under_body(doc):
    tools.add_decoration(doc, "background", color=(0, 0, 1), pages=[0])
    pix = doc[0].get_pixmap(dpi=20)
    assert pix.pixel(5, 5) == (0, 0, 255)
    assert "Hello" in doc[0].get_text()  # 본문은 그대로 위에


def test_watermark_tile(doc):
    tools.watermark(doc, text="대외비", position="tile", pages=[0])
    assert doc[0].get_text().count("대외비") > 4


def test_txt_cp949(tmp_path):
    from app.engine.document import Document
    p = tmp_path / "a.txt"
    p.write_bytes("메모장 한글 ANSI".encode("cp949"))
    assert "메모장 한글" in Document(str(p)).fitz[0].get_text()


def test_annots_json_roundtrip():
    """주석 내보내기·가져오기: 두 줄 강조는 줄마다, 선 두께·화살표 끝·글상자 글자 크기·네모 크기가 그대로."""
    from app.engine import annots as A
    from app.engine.document import Document
    d = Document()
    d.fitz[0].insert_text((72, 100), "Hello world line one", fontsize=14)
    d.fitz[0].insert_text((72, 120), "second line text here", fontsize=14)
    A.text_markup(d, 0, [fitz.Rect(72, 88, 230, 104), fitz.Rect(72, 108, 230, 124)], "Highlight")
    A.rect(d, 0, fitz.Rect(100, 200, 200, 300), width=6)
    A.oval(d, 0, fitz.Rect(220, 200, 300, 260), width=1.5, fill=(0, 1, 0))
    A.line(d, 0, fitz.Point(100, 400), fitz.Point(300, 450), width=3, arrow=True)
    A.ink(d, 0, [[fitz.Point(50, 600), fitz.Point(150, 650)]], width=4)
    A.polygon(d, 0, [fitz.Point(400, 400), fitz.Point(500, 420), fitz.Point(450, 500)], width=2, cloud=True)
    A.free_text(d, 0, fitz.Rect(300, 500, 450, 560), "리치 글상자", 18, (0, 0, 0), (1, 1, 0.8), (1, 0, 0), 2)
    A.free_text(d, 0, fitz.Rect(300, 600, 450, 660), "보통 글", 15, (0, 0, 1), None, (0, 0, 1), 1)
    A.free_text(d, 0, fitz.Rect(300, 700, 450, 760), "말풍선", 13, (0, 0, 0), (1, 1, 1), (0, 0, 0), 1, callout=fitz.Point(260, 800))
    A.note(d, 0, fitz.Point(20, 20), "메모", icon="Comment")
    data = json.loads(json.dumps(A.export_annots(d), ensure_ascii=False))
    e = Document()
    assert A.import_annots(e, data) == (len(data), 0)

    def sig(doc):
        out = []
        for a in doc.fitz[0].annots():
            k = a.type[1]
            if k == "Popup":
                continue
            out.append((k, [round(x, 1) for x in a.rect], (a.border or {}).get("width"), (a.border or {}).get("clouds"),
                        a.line_ends if k == "Line" else None, A._freetext_style(doc, a) if k == "FreeText" else None,
                        len(a.vertices or []) if k in A.MARKUP_KINDS else None, A.annot_colors(doc, a)))
        return out
    s1, s2 = sig(d), sig(e)
    assert s1 == s2
    hl = [x for x in s2 if x[0] == "Highlight"]
    assert len(hl) == 1 and hl[0][6] > 4 and hl[0][1][3] - hl[0][1][1] > 30  # 낱말마다 사각형, 두 줄에 걸침
    assert (18.0, (0.0, 0.0, 0.0)) in [x[5] for x in s2]
    # 10-19 2차까지 내보낸 파일(글상자 colors 의 stroke 가 바탕색, 글자 크기 없음)도 읽는다
    old = [dict(x, colors={"stroke": x["colors"]["fill"], "fill": []}) for x in data if x["type"] == "FreeText"]
    for x in old:
        x.pop("fontsize")
    assert A.import_annots(Document(), old) == (len(old), 0)


def test_replace_shared_image_only_that_page(tmp_path):
    """여러 쪽이 같이 쓰는 그림(같은 객체·폼 안 그림)을 바꾸면 그 쪽만 바뀐다."""
    from PIL import Image
    from app.engine import annots as A
    from app.engine.document import Document
    red, blue = str(tmp_path / "r.png"), str(tmp_path / "b.png")
    Image.new("RGB", (20, 20), (255, 0, 0)).save(red); Image.new("RGB", (20, 20), (0, 0, 255)).save(blue)
    R = fitz.Rect(100, 100, 200, 200)

    def colors(f):
        return [(lambda px: px.pixel(px.width // 2, px.height // 2))(f[i].get_pixmap(clip=R, dpi=36)) for i in range(f.page_count)]
    d = Document(); d.fitz.new_page(); d.fitz.new_page()
    for i in range(3):
        d.fitz[i].insert_image(R, filename=red)
    assert len({A.images_on_page(d, i)[0]["xref"] for i in range(3)}) == 1  # 같은 그림 객체
    A.replace_image(d, 1, A.images_on_page(d, 1)[0]["xref"], blue)
    assert colors(d.fitz) == [(255, 0, 0), (0, 0, 255), (255, 0, 0)]
    assert colors(fitz.open("pdf", d.to_bytes())) == [(255, 0, 0), (0, 0, 255), (255, 0, 0)]
    src = fitz.open(); src.new_page().insert_image(R, filename=red)
    d = Document(); d.fitz.new_page(); d.fitz.new_page()
    for i in range(3):
        d.fitz[i].show_pdf_page(d.fitz[i].rect, src, 0)  # 그림이 폼 XObject 안에
    A.replace_image(d, 2, A.images_on_page(d, 2)[0]["xref"], blue)
    assert colors(d.fitz) == [(255, 0, 0), (255, 0, 0), (0, 0, 255)]
    d = Document(); d.fitz[0].insert_image(R, filename=red)  # 혼자 쓰는 그림은 그대로 바꿔 끼움
    A.replace_image(d, 0, A.images_on_page(d, 0)[0]["xref"], blue)
    assert colors(d.fitz) == [(0, 0, 255)]


def test_eraser_hits_segment_between_vertices():
    """지우개: 꼭짓점에서 먼 선분 가운데를 지나가도 지운다. 지나간 길(prev → point)이 획을 가로질러도 지운다."""
    from app.engine import annots as A
    from app.engine.document import Document
    d = Document()
    A.ink(d, 0, [[fitz.Point(100, 100), fitz.Point(400, 100)]], width=2)
    assert not A.erase_ink_at(d, 0, fitz.Point(250, 130), radius=6)
    assert A.erase_ink_at(d, 0, fitz.Point(250, 104), radius=6)
    assert not list(d.fitz[0].annots())
    A.ink(d, 0, [[fitz.Point(100, 100), fitz.Point(400, 100)]], width=2)
    A.ink(d, 0, [[fitz.Point(100, 300), fitz.Point(400, 300)]], width=2)
    assert A.erase_ink_at(d, 0, fitz.Point(250, 150), radius=6, prev=fitz.Point(250, 50))  # 점 둘은 멀고 길이 가로지름
    assert [a.rect.y0 > 200 for a in d.fitz[0].annots()] == [True]


# ── 전체 시험(COM · OCR) ─────────────────────────────────────────────────
@pytest.mark.skipif(not FULL, reason="PDFAPP_FULL_TESTS=1 일 때만")
def test_ocr_pages(tmp_path):
    langs = ocr.available_languages()
    if not any(l.lower().startswith("ko") for l in langs):
        pytest.skip("한국어 문자 인식 없음")
    src = make_doc()
    pix = src[0].get_pixmap(dpi=200)
    d = fitz.open()
    p = d.new_page(width=595, height=842)
    p.insert_image(p.rect, pixmap=pix)
    assert not d[0].get_text().strip()
    n = ocr.ocr_pages(d, dpi=300)
    assert n > 0 and "2026" in d[0].get_text()
    assert "2026" in ocr.ocr_area(d, 0, fitz.Rect(60, 80, 500, 120))
    src.close(); d.close()


@pytest.mark.skipif(not FULL, reason="PDFAPP_FULL_TESTS=1 일 때만")
def test_to_hwpx(doc, tmp_path):
    out = convert.to_hwpx(doc, str(tmp_path / "a.hwpx"))
    assert os.path.getsize(out) > 0


@pytest.mark.skipif(not FULL, reason="PDFAPP_FULL_TESTS=1 일 때만")
def test_office_from_files(doc, tmp_path):
    docx = str(tmp_path / "w.docx")
    convert.to_word(doc, docx)
    out = convert.from_files([docx], str(tmp_path / "w.pdf"))
    assert fitz.open(out).page_count >= 3
